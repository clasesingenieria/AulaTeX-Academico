from pathlib import Path
from copy import deepcopy
import json,sys,hashlib
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.table import WD_TABLE_ALIGNMENT,WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
sys.stdout.reconfigure(encoding='utf-8')
ROOT=Path(__file__).resolve().parents[5];COURSE=ROOT/'ITESCA/maestria-en-gestion-administrativa/plan-de-negocios-mga'
ASSET=COURSE/'assets-plan-de-negocios/actividad-11-2026-10-01'
SOURCE=COURSE/'Entregas/reporte-plan-de-negocios-Actividad-10-AM-Taller.docx'
baseline=hashlib.sha256(SOURCE.read_bytes()).hexdigest()
data=json.loads((ASSET/'proceso.json').read_text(encoding='utf-8'))
d=Document(SOURCE)
def replace(p,text):
 props=deepcopy(p.runs[0]._r.rPr) if p.runs and p.runs[0]._r.rPr is not None else None
 p.clear();r=p.add_run(text)
 if props is not None:r._r.insert(0,props)
replace(d.paragraphs[4],'Proceso de prestación del servicio')
d.paragraphs[4].style='Title'
for r in d.paragraphs[4].runs:r.font.size=Pt(20);r.font.color.rgb=RGBColor(0,0,0)
replace(d.paragraphs[5],'AM Taller Autocentro\nCambio de aceite y filtro en vehículo ligero\nActividad 11 · Versiones gráfica y redactada')
replace(d.paragraphs[7],'Monterrey, Nuevo León\n1 de octubre de 2026')
keep={p._p for p in d.paragraphs[:9]}
for e in list(d._element.body):
 if e not in keep and e.tag!=qn('w:sectPr'):d._element.body.remove(e)
for name in ['Normal','Title','Heading 1','Heading 2']:
 d.styles[name].font.name='Arial';d.styles[name].font.color.rgb=RGBColor(0,0,0)
d.styles['Normal'].font.size=Pt(11)
d.styles['Normal'].paragraph_format.line_spacing=1.2
d.styles['Normal'].paragraph_format.space_after=Pt(8)
d.styles['Heading 1'].font.size=Pt(16)
d.styles['Heading 2'].font.size=Pt(12)
def p(t):return d.add_paragraph(t,style='Normal')
def h(t,page=False):
 x=d.add_paragraph(t,style='Heading 1');x.paragraph_format.page_break_before=page;x.paragraph_format.keep_with_next=True;return x
h('Alcance y criterio de diseño')
p('El proceso propuesto organiza un servicio concreto de AM Taller Autocentro: el cambio de aceite y filtro de un vehículo ligero, conforme a las especificaciones de su fabricante. Comienza con la solicitud y concluye con el registro de entrega y seguimiento. Los trabajos distintos se evalúan y autorizan por separado; no se incorporan automáticamente al mantenimiento solicitado.')
p('La actividad presenta una versión gráfica y otra redactada con los mismos 16 identificadores. La tabla detalla actividad, tiempo, materiales y equipo, y personal, conforme a la consigna del aula (Instituto Tecnológico Superior de Cajeme [ITESCA], s. f.). Las preguntas dentro de los rectángulos representan decisiones; sus salidas indican la ruta que corresponde seguir.')
p('La elección de este servicio permite explicar el proceso con recursos definidos; no se atribuye a resultados del sondeo. La retroalimentación de mercados requiere observar servicio, frecuencia, gasto y satisfacción mediante respuestas reales. En este diseño se prevé registrar esos datos al operar: la orden identificaría el servicio y el gasto, y el seguimiento recogería satisfacción. El sondeo de demanda continúa siendo un levantamiento independiente.')
p('Los tiempos de la tabla son estimaciones académicas de dedicación activa. No incluyen espera de cita, respuesta del cliente o suministro. La reautorización es condicional y una corrección agrega tiempo. El seguimiento se propone a las 24–72 horas posteriores, con consentimiento; ese intervalo no equivale a horas de trabajo. Estos rangos se validarían con registros reales antes de fijar promesas de entrega o capacidad diaria.')
p('Los puestos expresan funciones requeridas para el diseño, no contrataciones realizadas. Una misma persona podría cubrir funciones compatibles según la operación, manteniendo la responsabilidad por cada paso. El glosario del curso dedica los apartados 3.4 y 3.5 al proceso de servicio y a la maquinaria y herramientas; este diseño concreta esos componentes para el mantenimiento elegido (Propuesta de glosario de conceptos, 2024).')
d.add_paragraph('Lectura de los conectores',style='Heading 2')
p('A retorna al paso 8 para evaluar un hallazgo. B continúa en el paso 10 de la segunda hoja. C lleva al paso 15 para entregar, cerrar sin ejecución o reprogramar. Inicio y Fin son límites del proceso, no operaciones adicionales de la tabla.')
for n,label in [(1,'Solicitud y autorización'),(2,'Ejecución y cierre')]:
 h(f'Versión gráfica {n} de 2',True)
 x=p(label+' · Los números corresponden a las filas de la versión redactada.');x.paragraph_format.space_after=Pt(4)
 pic=d.add_paragraph();pic.paragraph_format.space_after=Pt(0);pic.paragraph_format.line_spacing=1
 pic.add_run().add_picture(str(ASSET/f'flujo-{n}.png'),width=Inches(6.1))
h('Versión redactada del proceso',True)
p('Los 16 pasos siguientes coinciden con el diagrama. Todos los rangos en minutos son estimados y se refieren al trabajo activo; las esperas y correcciones se registran aparte.')
t=d.add_table(rows=1,cols=4);t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
widths=[2.8,.85,1.85,1.0]
for col,w in zip(t.columns,widths):col.width=Inches(w)
headers=['Actividad detallada','Tiempo estimado','Materiales equipo y herramientas','Personal requerido']
for c,s in zip(t.rows[0].cells,headers):c.text=s
for step in data['steps']:
 cells=t.add_row().cells
 cells[0].text=f"{step['id']}. {step['title']}\n{step['detail']}"
 short=step['time'].split(' minutos')[0]+' min'
 if step['id']==12:short+='\nCorrección variable'
 if step['id']==13:short+='\nPor revisión'
 if step['id']==16:short+='\nContacto a las 24–72 h'
 cells[1].text=short;cells[2].text=step['materials'];cells[3].text=step['staff']
for ri,row in enumerate(t.rows):
 pr=row._tr.get_or_add_trPr();pr.append(OxmlElement('w:cantSplit'))
 if ri==0:pr.append(OxmlElement('w:tblHeader'))
 for ci,c in enumerate(row.cells):
  c.width=Inches(widths[ci]);c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
  cp=c._tc.get_or_add_tcPr();sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'DCE6F1' if ri==0 else 'FFFFFF');cp.append(sh)
  borders=OxmlElement('w:tcBorders')
  for side in ['top','bottom','left','right']:
   b=OxmlElement('w:'+side);b.set(qn('w:val'),'single');b.set(qn('w:sz'),'4');b.set(qn('w:color'),'D9D9D9');borders.append(b)
  cp.append(borders);mar=OxmlElement('w:tcMar')
  for side in ['top','bottom','left','right']:
   m=OxmlElement('w:'+side);m.set(qn('w:w'),'90');m.set(qn('w:type'),'dxa');mar.append(m)
  cp.append(mar)
  for par in c.paragraphs:
   par.paragraph_format.line_spacing=1.05;par.paragraph_format.space_before=Pt(2);par.paragraph_format.space_after=Pt(2);par.paragraph_format.keep_with_next=ri in [14,15]
   if ci==1:par.alignment=WD_ALIGN_PARAGRAPH.CENTER
   for r in par.runs:r.font.name='Arial';r.font.size=Pt(10.5);r.bold=ri==0
note=p('Nota. Estimaciones propias de planeación. No se presenta un tiempo total único: las rutas de negativa, reautorización, reprogramación y corrección son diferentes. Una confirmación de catálogo no sustituye la disponibilidad física del insumo.');note.paragraph_format.space_before=Pt(8)
h('Registros y aplicación al plan de negocios',True)
p('La orden de servicio propuesta debe vincular un folio con la solicitud, la identificación técnica del vehículo, el presupuesto aceptado, las autorizaciones, los insumos utilizados y el resultado de la revisión. El cierre diferenciará trabajo terminado, no ejecutado y reprogramado. Así se podrá explicar al cliente qué se hizo y qué importe se había autorizado, sin confundir una solicitud con una venta.')
p('Para atender la observación de la docente, el registro posterior de operación distinguirá cuatro datos: tipo de servicio; ocasiones de atención, precisando si provienen del historial del taller o de la respuesta anual D08 del sondeo; gasto de la ocasión, separado por servicio y alcance; y satisfacción expresada por el cliente. La falta de respuesta no se contará como satisfacción. Estos campos están previstos para recabar evidencia; no se presentan como resultados obtenidos.')
p('El presupuesto inicial y la autorización anteceden a la inspección con costo y a la intervención. Un hallazgo adicional obliga a detener la ampliación del trabajo, revisar el alcance y solicitar nueva autorización. La revisión de calidad distingue corregir el trabajo contratado de realizar otro servicio. Esa separación conecta la promesa comercial con las actividades que el taller deberá organizar.')
p('El proceso ofrece una base para la siguiente actividad de puestos y funciones: recepción coordina la relación con el cliente; el técnico ejecuta y registra; la función de insumos verifica disponibilidad y compatibilidad; coordinación y revisión mantienen el control del servicio. Los tiempos, recursos y responsables deberán ajustarse cuando exista evidencia de funcionamiento.')
d.add_paragraph('Referencias',style='Heading 2')
refs=[('Instituto Tecnológico Superior de Cajeme. (s. f.). ','Diagrama de flujo de proceso de producción o prestación del servicio',' [Consigna de Plan de Negocios]. https://cursos3.e-itesca.edu.mx/mod/assign/view.php?id=6550'),('','Propuesta de glosario de conceptos','. (2024). [Material didáctico, apartados 3.4 y 3.5]. ITESCA Virtual. https://cursos3.e-itesca.edu.mx/mod/resource/view.php?id=6527')]
for a,b,c in refs:
 q=p('');q.paragraph_format.left_indent=Inches(.3);q.paragraph_format.first_line_indent=Inches(-.3);q.add_run(a);q.add_run(b).italic=True;q.add_run(c)
p('Se utilizó asistencia de inteligencia artificial para organizar el flujo y revisar su correspondencia con la tabla. No se generaron respuestas del sondeo ni mediciones de operación.')
d.core_properties.title='Proceso de prestación del servicio de AM Taller';d.core_properties.author='Martín Jonathan de la Cruz Muñoz'
out=COURSE/'Entregas/reporte-plan-de-negocios-Actividad-11-AM-Taller.docx';d.save(out)
assert hashlib.sha256(SOURCE.read_bytes()).hexdigest()==baseline
assert len(t.rows)==17 and len(t.columns)==4
print(out)
