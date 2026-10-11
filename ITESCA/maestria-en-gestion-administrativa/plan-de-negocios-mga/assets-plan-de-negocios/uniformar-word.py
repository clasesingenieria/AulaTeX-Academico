import hashlib
import json
import re
from copy import deepcopy
from pathlib import Path
from shutil import copy2

import pymupdf
from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

ASSETS = Path(__file__).resolve().parent
SUBJECT = ASSETS.parent
OUTPUT = SUBJECT / "Entregas"
BACKUP = SUBJECT / "referencias-plan-de-negocios/notas-plan-de-negocios/materiales-generales/word-antes-portada-oficial"
COVERS = ASSETS / "portadas-word"
IMAGE = ASSETS / "actividad-12-2026-10-05/portada-oficial.png"


def digest(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def cover_image(title, subtitle, number, date, target):
    with pymupdf.open() as pdf:
        page = pdf.new_page(width=612, height=792)
        page.insert_image(page.rect, filename=str(IMAGE))
        page.insert_font(fontname="Arial", fontfile="C:/Windows/Fonts/arial.ttf")
        page.insert_font(fontname="ArialBold", fontfile="C:/Windows/Fonts/arialbd.ttf")
        block = pymupdf.Rect(56.7, 294.8, 544.3, 454)
        page.draw_rect(block, fill=(1, 1, 1), color=None, fill_opacity=0.96)
        unit = 2 if number <= 12 else 3
        metadata = ("Maestría en Gestión Administrativa\nAsignatura: Plan de Negocios\n"
                    f"Curso: GGPN01 - Unidad {unit}\nEstudiante: Martín Jonathan de la Cruz Muñoz\n"
                    "Matrícula: 26130503\nDocente: Celia Velázquez Reyna\n"
                    f"Fecha: {date}\nLugar: Monterrey, Nuevo León")
        assert page.insert_textbox(block + (12, 12, -12, -12), metadata,
                                   fontname="Arial", fontsize=11, lineheight=1.35) >= 0
        box = pymupdf.Rect(263.6, 552.7, 582, 705)
        content = f"{title}\n{subtitle}\n\nActividad {number} - Plan de Negocios"
        for size in (14, 13, 12, 11):
            shape = page.new_shape()
            result = shape.insert_textbox(box, content, fontname="ArialBold", fontsize=size,
                                          lineheight=1.25, color=(1, 1, 1))
            if result >= 0:
                shape.commit()
                break
        else:
            raise ValueError("El titulo no cabe en la portada")
        page.get_pixmap(matrix=pymupdf.Matrix(2, 2)).save(target)


def anchor_image(header, path):
    paragraph = header.paragraphs[0]
    paragraph.paragraph_format.space_before = paragraph.paragraph_format.space_after = Pt(0)
    picture = paragraph.add_run().add_picture(str(path), width=Inches(8.5), height=Inches(11))
    inline = picture._inline
    anchor = OxmlElement("wp:anchor")
    for name, value in {"distT": "0", "distB": "0", "distL": "0", "distR": "0", "simplePos": "0",
                        "relativeHeight": "0", "behindDoc": "1", "locked": "0", "layoutInCell": "1",
                        "allowOverlap": "1"}.items():
        anchor.set(name, value)
    simple = OxmlElement("wp:simplePos")
    simple.set("x", "0")
    simple.set("y", "0")
    anchor.append(simple)
    for direction in ("H", "V"):
        position = OxmlElement("wp:position" + direction)
        position.set("relativeFrom", "page")
        offset = OxmlElement("wp:posOffset")
        offset.text = "0"
        position.append(offset)
        anchor.append(position)
    anchor.append(deepcopy(inline.find(qn("wp:extent"))))
    anchor.append(OxmlElement("wp:wrapNone"))
    for name in ("wp:docPr", "wp:cNvGraphicFramePr", "a:graphic"):
        anchor.append(deepcopy(inline.find(qn(name))))
    inline.getparent().replace(inline, anchor)


def main():
    OUTPUT.mkdir(exist_ok=True)
    BACKUP.mkdir(parents=True, exist_ok=True)
    COVERS.mkdir(exist_ok=True)
    records = []
    for source in sorted(SUBJECT.glob("*.docx")):
        if source.name.startswith("~$"):
            continue
        target = OUTPUT / source.name
        backup = BACKUP / source.name
        if target.exists() or backup.exists():
            raise FileExistsError(f"Archivo ya procesado o destino existente: {source.name}")
        original_hash = digest(source)
        document = Document(source)
        paragraphs = [paragraph for paragraph in document.paragraphs if paragraph.text.strip()]
        first_texts = [paragraph.text for paragraph in paragraphs[:12]]
        match = re.search(r"Actividad-(\d+)", source.name)
        number = int(match[1]) if match else 10
        if number == 6545:
            number = 6
        titles = {5: "Preguntas para el análisis de demanda y oferta", 6: "Imagen de la empresa",
                  7: "Área geográfica de impacto", 8: "Mercadotecnia y publicidad",
                  9: "Resultados del examen de Mercadotecnia e Imagen",
                  10: "Resultados de la investigación de mercados",
                  11: "Proceso de prestación del servicio", 12: "Administración de recursos humanos: puestos y funciones"}
        title = titles[number]
        subtitle = "AM Taller Autocentro"
        if "Industrial-Revolucionaria" in source.name:
            subtitle = "Industrial Revolucionaria"
        elif "NexoTeX" in source.name:
            subtitle = "NexoTeX"
        if "Revision" in source.name:
            subtitle += " - Revisión metodológica del sondeo"
        if source.name.startswith("Sondeo"):
            title = "Sondeo sobre atención automotriz"
            subtitle += " - Instrumento de participación voluntaria"
        dates = [re.search(r"(?:Fecha[^:]*:|realizaci[oó]n:)\s*([^\n]+)", text, re.I) for text in first_texts]
        date = next((match[1].strip() for match in dates if match), "Fecha conservada en el cuerpo")
        original_final = deepcopy(document.element.body.sectPr)
        blocks = list(document.element.body)[:-1]
        cover_count = 0
        has_cover = first_texts[0].startswith("Instituto Tecnológico")
        if has_cover:
            for index, block in enumerate(blocks):
                if any(node.tag == qn("w:br") and node.get(qn("w:type")) == "page" for node in block.iter()):
                    cover_count = index + 1
                    break
            if not cover_count:
                raise ValueError("Portada sin limite confirmado: " + source.name)
        elif number == 5:
            first_heading = next((block for block in blocks if "Introducción" in "".join(block.itertext())), None)
            if first_heading is None:
                raise ValueError("Introduccion no localizada")
            cover_count = blocks.index(first_heading)
        retained = blocks[cover_count:]
        from lxml import etree
        retained_xml = [etree.tostring(block) for block in retained]
        for block in blocks[:cover_count]:
            document.element.body.remove(block)
        image = COVERS / (source.stem + ".png")
        cover_image(title, subtitle, number, date, image)
        document.add_section(WD_SECTION_START.NEW_PAGE)
        section_break = document.paragraphs[-1]._p
        document.element.body.remove(section_break)
        document.element.body.insert(0, section_break)
        final = document.element.body.sectPr
        final.getparent().replace(final, deepcopy(original_final))
        cover = document.sections[0]
        cover.page_width, cover.page_height = Inches(8.5), Inches(11)
        cover.top_margin = cover.bottom_margin = cover.left_margin = cover.right_margin = Inches(0)
        cover.header_distance = cover.footer_distance = Inches(0)
        cover.different_first_page_header_footer = False
        cover.header.is_linked_to_previous = False
        anchor_image(cover.header, image)
        cover.footer.is_linked_to_previous = False
        first_properties = section_break.get_or_add_pPr()
        for name in ("w:spacing",):
            spacing = OxmlElement(name)
            spacing.set(qn("w:before"), "0")
            spacing.set(qn("w:after"), "0")
            first_properties.append(spacing)
        body_section = document.sections[-1]
        if not original_final.findall(qn("w:headerReference")):
            body_section.header.is_linked_to_previous = False
        if not original_final.findall(qn("w:footerReference")):
            body_section.footer.is_linked_to_previous = False
        for expected, block in zip(retained_xml, retained):
            assert expected == etree.tostring(block), "Contenido previo alterado"
        copy2(source, backup)
        document.save(target)
        assert digest(source) == original_hash == digest(backup)
        records.append({"origen": source.name, "destino": target.relative_to(SUBJECT).as_posix(),
                        "respaldo": backup.relative_to(SUBJECT).as_posix(), "sha256_original": original_hash,
                        "sha256_word_nuevo": digest(target), "bloques_de_portada_reemplazados": cover_count,
                        "bloques_de_contenido_conservados": len(retained), "contenido_xml_conservado": True,
                        "actividad": number, "titulo": title, "subtitulo": subtitle, "fecha_portada": date,
                        "original_raiz_pendiente_retirada": True})
        (COVERS / "registro-word.json").write_text(json.dumps(records, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps({"word_generados": len(records), "contenido_conservado": True, "originales_respaldados": True}))


if __name__ == "__main__":
    main()