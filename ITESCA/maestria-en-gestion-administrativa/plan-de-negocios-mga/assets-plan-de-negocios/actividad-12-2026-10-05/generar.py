from copy import deepcopy
import hashlib
import json
from pathlib import Path
import re
import sys

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
import pymupdf

ASSETS = Path(__file__).resolve().parent
COURSE = ASSETS.parents[1]
ROOT = COURSE.parents[2]
sys.path.insert(0, str(ROOT))
from scripts.aulatex.seminario_objectives_renderer import (
    _restore_package_and_footnote, export_word_pdf, tex_escape,
)

STEM = "reporte-plan-de-negocios-Actividad-12-Recursos-Humanos"


def build_org_chart(chart):
    nodes = {node["id"]: node for node in chart["nodes"]}
    assert chart["root"] == "P1"
    assert set(nodes) == {f"P{number}" for number in range(1, 9)}
    placements = {"P1": (210, 12, 430, 74), "P7": (630, 82, 786, 145), "P8": (630, 166, 786, 229)}
    for index, identifier in enumerate(["P2", "P3", "P4", "P5", "P6"]):
        placements[identifier] = (14 + index * 114, 110, 124 + index * 114, 214)
    with pymupdf.open() as graph:
        page = graph.new_page(width=800, height=270)
        dark = (0.18, 0.18, 0.18)
        red = (0.56, 0.07, 0.09)
        page.draw_line((320, 74), (320, 92), color=dark, width=1.5)
        page.draw_line((69, 92), (525, 92), color=dark, width=1.5)
        for identifier in ["P2", "P3", "P4", "P5", "P6"]:
            rect = pymupdf.Rect(placements[identifier])
            page.draw_line(((rect.x0 + rect.x1) / 2, 92), ((rect.x0 + rect.x1) / 2, rect.y0), color=dark, width=1.5)
        for start, stop in [((430, 40), (604, 40)), ((604, 40), (604, 197.5)), ((604, 113.5), (630, 113.5)), ((604, 197.5), (630, 197.5))]:
            page.draw_line(start, stop, color=red, width=1.5, dashes="[5 3] 0")
        page.insert_text((636, 65), "Apoyos por encargo", fontsize=13, fontname="hebo", color=dark)
        for identifier, bounds in placements.items():
            rect = pymupdf.Rect(bounds)
            external = nodes[identifier]["relation"] == "apoyo externo"
            fill = red if identifier == chart["root"] else (0.97, 0.93, 0.91) if external else (0.95, 0.96, 0.97)
            page.draw_rect(rect, color=red if external else dark, fill=fill, width=1)
            color = (1, 1, 1) if identifier == chart["root"] else dark
            label = identifier + "\n" + nodes[identifier]["label"]
            remaining = page.insert_textbox(rect + (5, 6, -5, -5), label, fontname="hebo" if identifier == "P1" else "helv", fontsize=15, align=1, color=color)
            assert remaining >= 0, (identifier, remaining)
        page.draw_line((16, 244), (55, 244), color=dark, width=1.5)
        page.insert_text((64, 248), "Coordinación de funciones", fontsize=12, color=dark)
        page.draw_line((320, 244), (359, 244), color=red, width=1.5, dashes="[5 3] 0")
        page.insert_text((368, 248), "Servicio profesional externo", fontsize=12, color=dark)
        graph.save(ASSETS / "organigrama.pdf")
        page.get_pixmap(matrix=pymupdf.Matrix(3, 3)).save(ASSETS / "organigrama.png")
    return {"nodes": list(nodes), "placements": placements, "pdf": "organigrama.pdf", "png": "organigrama.png", "all_labels_fit": True}


def main():
    data = json.loads((ASSETS / "puestos.json").read_text(encoding="utf-8"))
    process_path = COURSE / "assets-plan-de-negocios/actividad-11-2026-10-01/proceso.json"
    process_hash = hashlib.sha256(process_path.read_bytes()).hexdigest()
    process = json.loads(process_path.read_text(encoding="utf-8"))
    steps = {step["id"]: step for step in process["steps"]}
    role_ids = {role["id"] for role in data["roles"]}
    assert len(role_ids) == 8
    assert [item["step"] for item in data["process_assignments"]] == list(steps)
    assert all(item["owner"] in role_ids and set(item["support"]) <= role_ids for item in data["process_assignments"])
    chart_receipt = build_org_chart(data["org_chart"])
    references = {item["key"]: item for item in data["references"]}

    def prose(value):
        return re.sub(r"\[@(\w+)\]", lambda match: "(" + references[match[1]]["citation"] + ")", value)

    def latex(value):
        parts = re.split(r"(\[@\w+\])", value)
        return "".join(r"\citepalias{" + part[2:-1] + "}" if part.startswith("[@") else tex_escape(part) for part in parts)

    source = COURSE / "Entregas/reporte-plan-de-negocios-Actividad-11-AM-Taller.docx"
    source_hash = hashlib.sha256(source.read_bytes()).hexdigest()
    document = Document(source)

    def replace(paragraph, text):
        properties = deepcopy(paragraph.runs[0]._r.rPr) if paragraph.runs and paragraph.runs[0]._r.rPr is not None else None
        paragraph.clear()
        run = paragraph.add_run(text)
        if properties is not None:
            run._r.insert(0, properties)

    replace(document.paragraphs[4], data["title"])
    replace(document.paragraphs[5], "AM Taller Autocentro\nActividad 12 - Puestos y funciones\nPropuesta de organización del personal")
    replace(document.paragraphs[7], data["location"] + "\n" + data["date"])
    protected = {paragraph._p for paragraph in document.paragraphs[:9]}
    for element in list(document.element.body):
        if element not in protected and element.tag != qn("w:sectPr"):
            document.element.body.remove(element)
    for name in ("Normal", "Heading 1", "Heading 2", "Heading 3"):
        document.styles[name].font.name = "Arial"
        document.styles[name].font.color.rgb = RGBColor(0, 0, 0)
    document.styles["Normal"].font.size = Pt(11)
    document.styles["Normal"].paragraph_format.line_spacing = 1.2
    document.styles["Normal"].paragraph_format.space_after = Pt(8)
    for name, size in [("Heading 1", 15), ("Heading 2", 12), ("Heading 3", 11)]:
        document.styles[name].font.size = Pt(size)
    tex = []

    def heading(title, level=1, page=False):
        paragraph = document.add_paragraph(title, style=f"Heading {level}")
        paragraph.paragraph_format.keep_with_next = True
        paragraph.paragraph_format.page_break_before = page
        for run in paragraph.runs:
            run.bold = True
            run.font.size = Pt({1: 15, 2: 12, 3: 11}[level])
        command = {1: "section", 2: "subsection", 3: "subsubsection"}[level]
        tex.append((r"\clearpage" if page else "") + "\\" + command + "{" + latex(title) + "}")

    def paragraph(value):
        result = document.add_paragraph(prose(value), style="Normal")
        tex.append(latex(value) + "\n\n")
        return result

    def table(headers, rows, widths, keep=False):
        result = document.add_table(rows=1, cols=len(headers))
        result.alignment = WD_TABLE_ALIGNMENT.CENTER
        result.autofit = False
        for column, width in zip(result.columns, widths):
            column.width = Inches(width)
        for cell, text in zip(result.rows[0].cells, headers):
            cell.text = text
        for row in rows:
            for cell, value in zip(result.add_row().cells, row):
                cell.text = value
        for row_index, row in enumerate(result.rows):
            properties = row._tr.get_or_add_trPr()
            properties.append(OxmlElement("w:cantSplit"))
            if row_index == 0:
                properties.append(OxmlElement("w:tblHeader"))
            for column_index, cell in enumerate(row.cells):
                cell.width = Inches(widths[column_index])
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                borders = OxmlElement("w:tcBorders")
                for side in ("top", "bottom", "left", "right"):
                    border = OxmlElement("w:" + side)
                    for key, value in [("val", "single"), ("sz", "4"), ("color", "D9D9D9")]:
                        border.set(qn("w:" + key), value)
                    borders.append(border)
                cell._tc.get_or_add_tcPr().append(borders)
                if row_index == 0:
                    shade = OxmlElement("w:shd")
                    shade.set(qn("w:fill"), "E7EDF2")
                    cell._tc.get_or_add_tcPr().append(shade)
                for block in cell.paragraphs:
                    block.paragraph_format.line_spacing = 1.05
                    block.paragraph_format.space_before = Pt(1)
                    block.paragraph_format.space_after = Pt(1)
                    block.paragraph_format.keep_with_next = keep and row_index < len(result.rows) - 1
                    for run in block.runs:
                        run.font.name = "Arial"
                        run.font.size = Pt(10)
                        run.bold = row_index == 0
        return result

    def longtable(headers, rows, layout):
        tex.append(r"{\small\setstretch{1.05}\renewcommand{\arraystretch}{1.08}\begin{longtable}{" + layout + "}")
        header = " & ".join(r"\textbf{" + latex(value) + "}" for value in headers) + r" \\ \midrule"
        tex.append(r"\toprule " + header + r"\endhead")
        tex.extend(" & ".join(latex(value) for value in row) + r" \\ \midrule" for row in rows)
        tex.append(r"\end{longtable}}")

    heading("Introducción")
    for value in data["introduction"]:
        paragraph(value)
    heading("Puestos, funciones y coordinación del servicio")
    heading("Estructura funcional propuesta", 2)
    for value in data["organization"]:
        paragraph(value)
    heading("Organigrama funcional", 2)
    graphic = document.add_paragraph()
    graphic.paragraph_format.keep_with_next = True
    graphic.add_run().add_picture(str(ASSETS / "organigrama.png"), width=Inches(6.5))
    tex.append(r"\par\noindent\includegraphics[width=\linewidth]{" + (ASSETS / "organigrama.pdf").relative_to(ROOT).as_posix() + "}\n")
    paragraph("Figura 1. Organigrama funcional propuesto. Elaboración propia. " + data["org_chart"]["note"])
    heading("Fichas de puestos", 2)
    paragraph("Las funciones se proponen para la etapa de diseño. Su agrupación o separación dependerá de competencias, carga de trabajo y recursos comprobados; los servicios profesionales se contratarían por un encargo específico.")
    for role in data["roles"]:
        title = role["id"] + ". " + role["title"]
        block = document.add_paragraph(title, style="Heading 3")
        block.paragraph_format.keep_with_next = True
        for run in block.runs:
            run.bold = True
        rows = [
            ["Ubicación y enlaces", role["area"] + "\n" + role["reporting"]],
            ["Función principal", role["function"]],
            ["Actividades", "\n".join(f"{number}. {value}" for number, value in enumerate(role["activities"], 1))],
            ["Límite de decisión", role["authority"]],
            ["Registros", role["evidence"]],
            ["Perfil funcional", role["profile"]],
        ]
        table(["Puesto " + role["id"], role["title"]], rows, [1.25, 5.25], keep=True)
        document.add_paragraph()
        tex.append(r"\par\medskip\noindent\begin{minipage}{\linewidth}\subsubsection*{" + latex(title) + "}")
        tex.append(r"{\small\setstretch{1.05}\renewcommand{\arraystretch}{1.08}\begin{tabular}{@{}p{0.20\linewidth}p{0.75\linewidth}@{}}\toprule")
        tex.append(r"\textbf{Componente} & \textbf{Descripción del puesto} \\ \midrule")
        for label, value in rows:
            value_tex = latex(value).replace("\n", r"\par ")
            tex.append(latex(label) + " & " + value_tex + r" \\ \midrule")
        tex.append(r"\end{tabular}}\end{minipage}\par\medskip")

    heading("Responsabilidades en los 16 pasos del servicio", 2, page=True)
    paragraph("Los números y nombres de los pasos conservan el proceso de la Actividad 11. Responsable significa asegurar la realización y el registro; apoyo no elimina esa responsabilidad. La aceptación del cliente se mantiene como decisión externa al taller. P1 asegura recursos y asignación, pero no reemplaza al técnico en decisiones de su especialidad.")
    rows = [[str(item["step"]) + ". " + steps[item["step"]]["title"], item["owner"], ", ".join(item["support"]) or "Sin apoyo fijo", item["record"]] for item in data["process_assignments"]]
    headers = ["Paso del proceso", "A cargo", "Apoyo", "Registro o control"]
    table(headers, rows, [2.25, 0.8, 0.95, 2.5])
    longtable(headers, rows, r"@{}p{0.33\linewidth}p{0.13\linewidth}p{0.12\linewidth}p{0.32\linewidth}@{}")

    heading("Actividades complementarias y apoyos", 2)
    paragraph("La cobertura no termina en el cambio de aceite: los servicios de llantas, la venta de refacciones y la administración requieren encargos propios. Los apoyos externos se definen por necesidad; no se presume un velador, un proveedor contratado o un puesto permanente para cada especialidad.")
    rows = [[item["activity"], item["owner"], item["support"]] for item in data["complementary_assignments"]]
    headers = ["Actividad", "A cargo", "Coordinación necesaria"]
    table(headers, rows, [2.3, 0.8, 3.4])
    longtable(headers, rows, r"@{}p{0.33\linewidth}p{0.13\linewidth}p{0.47\linewidth}@{}")

    heading("Asignación gradual y controles", 2)
    paragraph("Las etapas son criterios de decisión, no fechas ni contrataciones comprometidas. P1 revisará la evidencia antes de incorporar personal; una necesidad de revisión técnica o asesoría no se pospone por encontrarse en la etapa inicial.")
    rows = [[item["stage"], item["assignment"], item["trigger"], item["evidence"]] for item in data["incorporation_stages"]]
    headers = ["Etapa", "Agrupación o incorporación", "Cuándo decidir", "Evidencia necesaria"]
    table(headers, rows, [1.05, 1.95, 1.75, 1.75])
    longtable(headers, rows, r"@{}p{0.13\linewidth}p{0.28\linewidth}p{0.24\linewidth}p{0.25\linewidth}@{}")
    heading("Revisión técnica y liberación", 2)
    paragraph("Ejecutar, comprobar y autorizar la entrega son responsabilidades relacionadas, pero distintas. P3 responde técnicamente; P1 asegura la asignación del revisor; P2 controla el cierre administrativo y la comunicación con el cliente. La intervención de P1 no equivale a aprobación técnica.")
    rows = [[item["case"], item["execution"], item["verification"], item["release"]] for item in data["technical_review"]]
    headers = ["Situación", "Ejecución", "Verificación", "Condición de liberación"]
    table(headers, rows, [1.05, 1.6, 1.9, 1.95])
    longtable(headers, rows, r"@{}p{0.15\linewidth}p{0.22\linewidth}p{0.26\linewidth}p{0.27\linewidth}@{}")
    for value in data["controls"]:
        paragraph(value)
    heading("Seguimiento propuesto", 2)
    paragraph("Los indicadores siguientes son criterios de revisión, no resultados ni metas alcanzadas. Las razones pueden expresarse como porcentaje; se anotarán numerador, denominador y periodo. Con denominador cero se registrará no aplicable. La revisión inicial se propone al cierre de cada orden y mediante consolidación periódica de P1, ajustable a la operación real.")
    rows = [[item["name"], item["definition"], item["owner"] + ". " + item["review"]] for item in data["indicators"]]
    headers = ["Criterio", "Cálculo o evidencia", "Responsable y uso"]
    table(headers, rows, [1.25, 3.0, 2.25])
    longtable(headers, rows, r"@{}p{0.19\linewidth}p{0.43\linewidth}p{0.31\linewidth}@{}")

    heading("Conclusiones", page=True)
    for value in data["conclusion"]:
        last = paragraph(value)
    note = OxmlElement("w:footnoteReference")
    note.set(qn("w:id"), "424242")
    last.add_run()._r.append(note)
    tex.append(r"\footnote{" + latex(data["ai_note"]) + "}")
    heading("Referencias", page=True)
    tex.pop()
    tex.append(r"\clearpage\renewcommand{\refname}{Referencias}\begin{thebibliography}{3}")
    emphasis = {"itescaRH": "Administración de Recursos Humanos (Puestos y funciones)", "glosario": "Propuesta de glosario de conceptos", "openstax": "Principles of management"}
    for reference in sorted(data["references"], key=lambda item: item["apa"].casefold()):
        block = document.add_paragraph(style="Normal")
        block.paragraph_format.left_indent = Inches(0.5)
        block.paragraph_format.first_line_indent = Inches(-0.5)
        marked = emphasis[reference["key"]]
        before, after = reference["apa"].split(marked, 1)
        block.add_run(before)
        block.add_run(marked).italic = True
        block.add_run(after)
        prefix, url = reference["apa"].rsplit("https://", 1)
        tex.append(r"\bibitem[" + tex_escape(reference["citation"]) + "]{" + reference["key"] + "}")
        formatted = tex_escape(prefix).replace(tex_escape(marked), r"\textit{" + tex_escape(marked) + "}")
        tex.append(formatted + r"\url{https://" + url + "}")
    tex.append(r"\end{thebibliography}")
    source_tex = r"""\def\actividadnumero{12}
\def\actividadtitulo{Administración de recursos humanos: puestos y funciones}
\def\actividadsubtitulo{AM Taller Autocentro: organización del personal}
\def\actividadunidad{2}
\def\actividadproducto{Fichas de puestos y responsabilidades}
\def\actividadfecha{5 de octubre de 2026}
\def\actividadlugar{Monterrey, Nuevo León}
\def\actividadextras{\hypersetup{hidelinks}\renewcommand{\bibliography}[1]{}\setlength{\emergencystretch}{2em}
"""
    for reference in data["references"]:
        source_tex += r"\defcitealias{" + reference["key"] + "}{" + tex_escape(reference["citation"]) + "}\n"
    source_tex += "}\n\\long\\def\\actividadcontenido{\n" + "\n".join(tex) + "\n}\n"
    source_tex += r"\input{ITESCA/maestria-en-gestion-administrativa/plan-de-negocios-mga/reporte-plan-de-negocios-plantilla-actividad.tex}" + "\n"
    (COURSE / (STEM + ".tex")).write_text(source_tex, encoding="utf-8")
    document.core_properties.title = data["title"]
    document.core_properties.subject = "Actividad 12 - Plan de Negocios"
    document.core_properties.author = "Martín Jonathan de la Cruz Muñoz"
    destination = COURSE / (STEM + ".docx")
    document.save(destination)
    _restore_package_and_footnote(source, destination, data["ai_note"])
    word_pdf = COURSE / (STEM + "-Word.pdf")
    receipt = export_word_pdf(destination, word_pdf)
    assert hashlib.sha256(source.read_bytes()).hexdigest() == source_hash
    assert hashlib.sha256(process_path.read_bytes()).hexdigest() == process_hash
    receipt.update(source_word_sha256=source_hash, process_sha256=process_hash, roles=8, process_steps=16, organigrama=chart_receipt, stages=3, review_cases=3, submitted=False)
    (ASSETS / "generacion.json").write_text(json.dumps(receipt, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps(receipt, ensure_ascii=True, indent=2))


if __name__ == "__main__":
    main()