import hashlib
import json
from pathlib import Path
import re
import unicodedata
from zipfile import ZipFile

from docx import Document
import pymupdf

ASSETS = Path(__file__).resolve().parent
COURSE = ASSETS.parents[1]
ROOT = COURSE.parents[2]
STEM = "reporte-plan-de-negocios-Actividad-12-Recursos-Humanos"


def normalize(value):
    value = unicodedata.normalize("NFKC", value).replace("\xad", "")
    value = re.sub(r"(?<=\w)-\s*\n(?=\w)", "", value)
    return " ".join(value.split())


def digest(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def main():
    data_path = ASSETS / "puestos.json"
    data = json.loads(data_path.read_text(encoding="utf-8"))
    process_path = COURSE / "assets-plan-de-negocios/actividad-11-2026-10-01/proceso.json"
    process = json.loads(process_path.read_text(encoding="utf-8"))
    receipt = json.loads((ASSETS / "generacion.json").read_text(encoding="utf-8"))
    source = COURSE / "Entregas/reporte-plan-de-negocios-Actividad-11-AM-Taller.docx"
    word = COURSE / (STEM + ".docx")
    report = COURSE / (STEM + ".tex")
    document = Document(word)
    role_ids = {role["id"] for role in data["roles"]}
    checks = {
        "eight_distinct_roles": len(role_ids) == 8,
        "sixteen_steps_same_order": [item["step"] for item in data["process_assignments"]] == [step["id"] for step in process["steps"]],
        "valid_owners_and_support": all(item["owner"] in role_ids and set(item["support"]) <= role_ids for item in data["process_assignments"]),
        "complementary_owners": all(item["owner"] in role_ids for item in data["complementary_assignments"]),
        "source_word_unchanged": digest(source) == receipt["source_word_sha256"],
        "source_process_unchanged": digest(process_path) == receipt["process_sha256"],
        "eleven_word_tables": len(document.tables) == 11,
        "three_report_acts": len(re.findall(r"\\section\{", report.read_text(encoding="utf-8"))) == 3,
        "word_three_acts": len([paragraph for paragraph in document.paragraphs if paragraph.style.name == "Heading 1" and paragraph.text != "Referencias"]) == 3,
    }
    for index, role in enumerate(data["roles"]):
        table_text = "\n".join(cell.text for row in document.tables[index].rows for cell in row.cells)
        expected = [role[key] for key in ["title", "area", "function", "authority", "evidence", "profile", "reporting"]] + role["activities"]
        checks[role["id"] + "_complete_table"] = all(value in table_text for value in expected)
    matrix = document.tables[8]
    checks["matrix_seventeen_rows"] = len(matrix.rows) == 17
    checks["matrix_step_owner_parity"] = all(
        matrix.rows[index].cells[0].text == f"{step['id']}. {step['title']}"
        and matrix.rows[index].cells[1].text == item["owner"]
        for index, (step, item) in enumerate(zip(process["steps"], data["process_assignments"]), 1)
    )
    with ZipFile(word) as archive, ZipFile(source) as original:
        checks["native_ai_note"] = "word/footnotes.xml" in archive.namelist() and data["ai_note"].encode("utf-8") in archive.read("word/footnotes.xml")
        checks["institutional_banner_preserved"] = archive.read("word/media/image1.png") == original.read("word/media/image1.png")
    pdfs = []
    for kind, path, input_path in [("word", COURSE / (STEM + "-Word.pdf"), word), ("latex", COURSE / (STEM + ".pdf"), report)]:
        checks[kind + "_fresh"] = path.exists() and path.stat().st_mtime >= input_path.stat().st_mtime
        with pymupdf.open(path) as pdf:
            texts = [page.get_text() for page in pdf]
            combined = "\n".join(texts)
            checks[kind + "_metadata"] = all(value in texts[0] for value in ["26130503", "Celia", "Velázquez", "Plan de Negocios"])
            checks[kind + "_eight_functions_visible"] = all(normalize(role["function"]) in normalize(combined) for role in data["roles"])
            checks[kind + "_all_steps_visible"] = all(normalize(step["title"]) in normalize(combined) for step in process["steps"])
            checks[kind + "_no_foreign_identity"] = not re.search(r"UnADM|vTaxi|Industrial Revolucionaria|NexoTeX|Seminario I", combined)
            checks[kind + "_no_editorial_debris"] = not re.search(r"\[PENDIENTE|\[\?\]|\[@|Refuerzo Ciclo A|Módulo Moodle|Vencimiento publicado|Actividad local", combined, re.I)
            checks[kind + "_scope_caveat"] = "no ocho contrataciones" in normalize(combined)
            checks[kind + "_ai_note"] = normalize(data["ai_note"]) in normalize(combined)
            folder = ASSETS / "revision-visual" / kind
            folder.mkdir(parents=True, exist_ok=True)
            for number, page in enumerate(pdf, 1):
                page.get_pixmap(matrix=pymupdf.Matrix(1.15, 1.15)).save(folder / f"pagina-{number:02}.png")
            pdfs.append({"path": path.relative_to(ROOT).as_posix(), "pages": len(pdf)})
    log = report.with_suffix(".log")
    if not log.exists():
        log = ASSETS / "compilacion" / log.name
    checks["latex_log_clean"] = not re.search(r"Overfull|undefined citations|Citation .* undefined|^!", log.read_text(encoding="utf-8", errors="replace"), re.M)
    files = [data_path, source, process_path, word, report, report.with_suffix(".bib"), report.with_suffix(".pdf"), COURSE / (STEM + "-Word.pdf")]
    result = {
        "passed": all(checks.values()), "checks": checks, "pdfs": pdfs,
        "artifacts": [{"path": path.relative_to(ROOT).as_posix(), "sha256": digest(path)} for path in files],
        "submitted": False,
        "scope": "Validacion documental; no acredita contrataciones, sondeo aplicado, operacion real ni certificacion del ciclo multiagente.",
    }
    (ASSETS / "validacion.json").write_text(json.dumps(result, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps({"passed": result["passed"], "failures": [key for key, passed in checks.items() if not passed], "pdfs": pdfs}, ensure_ascii=True, indent=2))
    if not result["passed"]:
        raise SystemExit(1)


if __name__ == "__main__":
    main()