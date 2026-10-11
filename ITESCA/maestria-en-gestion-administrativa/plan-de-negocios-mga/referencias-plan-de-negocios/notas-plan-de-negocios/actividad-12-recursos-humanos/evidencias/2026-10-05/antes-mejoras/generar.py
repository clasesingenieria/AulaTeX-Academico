from copy import deepcopy
import hashlib
import json
from pathlib import Path
import re
import sys

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ASSETS = Path(__file__).resolve().parent
COURSE = ASSETS.parents[1]
ROOT = COURSE.parents[2]
sys.path.insert(0, str(ROOT))
from scripts.aulatex.seminario_objectives_renderer import (
    _restore_package_and_footnote, export_word_pdf, tex_escape,
)

STEM = "reporte-plan-de-negocios-Actividad-12-Recursos-Humanos"


def main():
    data = json.loads((ASSETS / "puestos.json").read_text(encoding="utf-8"))
    process_path = COURSE / "assets-plan-de-negocios/actividad-11-2026-10-01/proceso.json"
    process_hash = hashlib.sha256(process_path.read_bytes()).hexdigest()
    process = json.loads(process_path.read_text(encoding="utf-8"))
    steps = {step["id"]: step for step in process["steps"]}
    role_ids = {role["id"] for role in data["roles"]}
    assert len(role_ids) == 8
    assert [item["step"] for item in data["process_assignments"]] == list(steps)
    assert all(item["owner"] in role_ids and set(item["support"]) <= role_ids for item in data["process_assignments"])
    references = {item["key"]: item for item in data["references"]}

    def prose(value):
        return re.sub(r"\[@(\w+)\]", lambda match: "(" + references[match[1]]["citation"] + ")", value)

    def latex(value):
        parts = re.split(r"(\[@\w+\])", value)
        return "".join(r"\citepalias{" + part[2:-1] + "}" if part.startswith("[@") else tex_escape(part) for part in parts)

    source = COURSE / "Entregas/reporte-plan-de-negocios-Actividad-11-AM-Taller.docx"
    source_hash = hashlib.sha256(source.read_bytes()).hexdigest()
    document = Document(source)

    def replace(paragraph, text):
        properties = deepcopy(paragraph.runs[0]._r.rPr) if paragraph.runs and paragraph.runs[0]._r.rPr is not None else None
        paragraph.clear()
        run = paragraph.add_run(text)
        if properties is not None:
            run._r.insert(0, properties)

    replace(document.paragraphs[4], data["title"])
    replace(document.paragraphs[5], "AM Taller Autocentro\nActividad 12 - Puestos y funciones\nPropuesta de organización del personal")
    replace(document.paragraphs[7], data["location"] + "\n" + data["date"])
    protected = {paragraph._p for paragraph in document.paragraphs[:9]}
    for element in list(document.element.body):
        if element not in protected and element.tag != qn("w:sectPr"):
            document.element.body.remove(element)
    for name in ("Normal", "Heading 1", "Heading 2", "Heading 3"):
        document.styles[name].font.name = "Arial"
        document.styles[name].font.color.rgb = RGBColor(0, 0, 0)
    document.styles["Normal"].font.size = Pt(11)
    document.styles["Normal"].paragraph_format.line_spacing = 1.2
    document.styles["Normal"].paragraph_format.space_after = Pt(8)
    for name, size in [("Heading 1", 15), ("Heading 2", 12), ("Heading 3", 11)]:
        document.styles[name].font.size = Pt(size)
    tex = []

    def heading(title, level=1, page=False):
        paragraph = document.add_paragraph(title, style=f"Heading {level}")
        paragraph.paragraph_format.keep_with_next = True
        paragraph.paragraph_format.page_break_before = page
        for run in paragraph.runs:
            run.bold = True
            run.font.size = Pt({1: 15, 2: 12, 3: 11}[level])
        command = {1: "section", 2: "subsection", 3: "subsubsection"}[level]
        tex.append((r"\clearpage" if page else "") + "\\" + command + "{" + latex(title) + "}")

    def paragraph(value):
        result = document.add_paragraph(prose(value), style="Normal")
        tex.append(latex(value) + "\n\n")
        return result

    def table(headers, rows, widths, keep=False):
        result = document.add_table(rows=1, cols=len(headers))
        result.alignment = WD_TABLE_ALIGNMENT.CENTER
        result.autofit = False
        for column, width in zip(result.columns, widths):
            column.width = Inches(width)
        for cell, text in zip(result.rows[0].cells, headers):
            cell.text = text
        for row in rows:
            for cell, value in zip(result.add_row().cells, row):
                cell.text = value
        for row_index, row in enumerate(result.rows):
            properties = row._tr.get_or_add_trPr()
            properties.append(OxmlElement("w:cantSplit"))
            if row_index == 0:
                properties.append(OxmlElement("w:tblHeader"))
            for column_index, cell in enumerate(row.cells):
                cell.width = Inches(widths[column_index])
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                borders = OxmlElement("w:tcBorders")
                for side in ("top", "bottom", "left", "right"):
                    border = OxmlElement("w:" + side)
                    for key, value in [("val", "single"), ("sz", "4"), ("color", "D9D9D9")]:
                        border.set(qn("w:" + key), value)
                    borders.append(border)
                cell._tc.get_or_add_tcPr().append(borders)
                if row_index == 0:
                    shade = OxmlElement("w:shd")
                    shade.set(qn("w:fill"), "E7EDF2")
                    cell._tc.get_or_add_tcPr().append(shade)
                for block in cell.paragraphs:
                    block.paragraph_format.line_spacing = 1.05
                    block.paragraph_format.space_before = Pt(2)
                    block.paragraph_format.space_after = Pt(2)
                    block.paragraph_format.keep_with_next = keep and row_index < len(result.rows) - 1
                    for run in block.runs:
                        run.font.name = "Arial"
                        run.font.size = Pt(10)
                        run.bold = row_index == 0
        return result

    def longtable(headers, rows, layout):
        tex.append(r"{\small\renewcommand{\arraystretch}{1.15}\begin{longtable}{" + layout + "}")
        header = " & ".join(r"\textbf{" + latex(value) + "}" for value in headers) + r" \\ \midrule"
        tex.append(r"\toprule " + header + r"\endhead")
        tex.extend(" & ".join(latex(value) for value in row) + r" \\ \midrule" for row in rows)
        tex.append(r"\end{longtable}}")

    heading("Introducción")
    for value in data["introduction"]:
        paragraph(value)
    heading("Puestos, funciones y coordinación del servicio")
    heading("Estructura funcional propuesta", 2)
    for value in data["organization"]:
        paragraph(value)
    heading("Fichas de puestos", 2, page=True)
    paragraph("Las funciones se proponen para la etapa de diseño. Su agrupación o separación dependerá de competencias, carga de trabajo y recursos comprobados; los servicios profesionales se contratarían por un encargo específico.")
    for role in data["roles"]:
        title = role["id"] + ". " + role["title"]
        block = document.add_paragraph(title, style="Heading 3")
        block.paragraph_format.keep_with_next = True
        for run in block.runs:
            run.bold = True
        rows = [
            ["Área", role["area"]],
            ["Función principal", role["function"]],
            ["Actividades", "\n".join(f"{number}. {value}" for number, value in enumerate(role["activities"], 1))],
            ["Límite de decisión", role["authority"]],
            ["Registros", role["evidence"]],
            ["Perfil funcional", role["profile"]],
            ["Coordinación", role["reporting"]],
        ]
        table(["Puesto " + role["id"], role["title"]], rows, [1.25, 5.25], keep=True)
        document.add_paragraph()
        tex.append(r"\par\medskip\noindent\begin{minipage}{\linewidth}\subsubsection*{" + latex(title) + "}")
        tex.append(r"{\small\renewcommand{\arraystretch}{1.15}\begin{tabular}{@{}p{0.20\linewidth}p{0.75\linewidth}@{}}\toprule")
        tex.append(r"\textbf{Componente} & \textbf{Descripción del puesto} \\ \midrule")
        for label, value in rows:
            value_tex = latex(value).replace("\n", r"\par ")
            tex.append(latex(label) + " & " + value_tex + r" \\ \midrule")
        tex.append(r"\end{tabular}}\end{minipage}\par\medskip")

    heading("Responsabilidades en los 16 pasos del servicio", 2, page=True)
    paragraph("Los números y nombres de los pasos conservan el proceso de la Actividad 11. Responsable significa asegurar la realización y el registro; apoyo no elimina esa responsabilidad. La aceptación del cliente se mantiene como decisión externa al taller. P1 asegura recursos y asignación, pero no reemplaza al técnico en decisiones de su especialidad.")
    rows = [[str(item["step"]) + ". " + steps[item["step"]]["title"], item["owner"], ", ".join(item["support"]) or "Sin apoyo fijo", item["record"]] for item in data["process_assignments"]]
    headers = ["Paso del proceso", "A cargo", "Apoyo", "Registro o control"]
    table(headers, rows, [2.25, 0.8, 0.95, 2.5])
    longtable(headers, rows, r"@{}p{0.33\linewidth}p{0.13\linewidth}p{0.12\linewidth}p{0.32\linewidth}@{}")

    heading("Actividades complementarias y apoyos", 2)
    paragraph("La cobertura no termina en el cambio de aceite: los servicios de llantas, la venta de refacciones y la administración requieren encargos propios. Los apoyos externos se definen por necesidad; no se presume un velador, un proveedor contratado o un puesto permanente para cada especialidad.")
    rows = [[item["activity"], item["owner"], item["support"]] for item in data["complementary_assignments"]]
    headers = ["Actividad", "A cargo", "Coordinación necesaria"]
    table(headers, rows, [2.3, 0.8, 3.4])
    longtable(headers, rows, r"@{}p{0.33\linewidth}p{0.13\linewidth}p{0.47\linewidth}@{}")

    heading("Asignación gradual y controles", 2)
    for value in data["controls"]:
        paragraph(value)
    heading("Seguimiento propuesto", 2)
    paragraph("Los indicadores siguientes son criterios de revisión, no resultados ni metas alcanzadas. Las razones pueden expresarse como porcentaje; se anotarán numerador, denominador y periodo. Con denominador cero se registrará no aplicable. La revisión inicial se propone al cierre de cada orden y mediante consolidación periódica de P1, ajustable a la operación real.")
    rows = [[item["name"], item["definition"], item["owner"] + ". " + item["review"]] for item in data["indicators"]]
    headers = ["Criterio", "Cálculo o evidencia", "Responsable y uso"]
    table(headers, rows, [1.25, 3.0, 2.25])
    longtable(headers, rows, r"@{}p{0.19\linewidth}p{0.43\linewidth}p{0.31\linewidth}@{}")

    heading("Conclusiones", page=True)
    for value in data["conclusion"]:
        last = paragraph(value)
    note = OxmlElement("w:footnoteReference")
    note.set(qn("w:id"), "424242")
    last.add_run()._r.append(note)
    tex.append(r"\footnote{" + latex(data["ai_note"]) + "}")
    heading("Referencias", page=True)
    tex.pop()
    tex.append(r"\clearpage\renewcommand{\refname}{Referencias}\begin{thebibliography}{3}")
    bib = []
    emphasis = {"itescaRH": "Administración de Recursos Humanos (Puestos y funciones)", "glosario": "Propuesta de glosario de conceptos", "openstax": "Principles of management"}
    for reference in sorted(data["references"], key=lambda item: item["apa"].casefold()):
        block = document.add_paragraph(style="Normal")
        block.paragraph_format.left_indent = Inches(0.5)
        block.paragraph_format.first_line_indent = Inches(-0.5)
        marked = emphasis[reference["key"]]
        before, after = reference["apa"].split(marked, 1)
        block.add_run(before)
        block.add_run(marked).italic = True
        block.add_run(after)
        prefix, url = reference["apa"].rsplit("https://", 1)
        tex.append(r"\bibitem[" + tex_escape(reference["citation"]) + "]{" + reference["key"] + "}")
        formatted = tex_escape(prefix).replace(tex_escape(marked), r"\textit{" + tex_escape(marked) + "}")
        tex.append(formatted + r"\url{https://" + url + "}")
        author, year = reference["citation"].rsplit(", ", 1)
        title = "10.1 Organizational structures and design" if reference["key"] == "openstax" else marked
        author_field = "key" if reference["key"] == "glosario" else "author"
        bib.append("@misc{" + reference["key"] + ",\n  " + author_field + " = {{" + tex_escape(author) + "}},\n  year = {" + year + "},\n  title = {{" + tex_escape(title) + "}},\n  howpublished = {\\url{https://" + url + "}}\n}")
    tex.append(r"\end{thebibliography}")
    source_tex = r"""\def\actividadnumero{12}
\def\actividadtitulo{Administración de recursos humanos: puestos y funciones}
\def\actividadsubtitulo{AM Taller Autocentro: organización del personal}
\def\actividadunidad{2}
\def\actividadproducto{Fichas de puestos y responsabilidades}
\def\actividadfecha{5 de octubre de 2026}
\def\actividadlugar{Monterrey, Nuevo León}
\def\actividadextras{\hypersetup{hidelinks}\renewcommand{\bibliography}[1]{}\setlength{\emergencystretch}{2em}
"""
    for reference in data["references"]:
        source_tex += r"\defcitealias{" + reference["key"] + "}{" + tex_escape(reference["citation"]) + "}\n"
    source_tex += "}\n\\long\\def\\actividadcontenido{\n" + "\n".join(tex) + "\n}\n"
    source_tex += r"\input{ITESCA/maestria-en-gestion-administrativa/plan-de-negocios-mga/reporte-plan-de-negocios-plantilla-actividad.tex}" + "\n"
    (COURSE / (STEM + ".tex")).write_text(source_tex, encoding="utf-8")
    (COURSE / (STEM + ".bib")).write_text("\n\n".join(bib) + "\n", encoding="utf-8")
    document.core_properties.title = data["title"]
    document.core_properties.subject = "Actividad 12 - Plan de Negocios"
    document.core_properties.author = "Martín Jonathan de la Cruz Muñoz"
    destination = COURSE / (STEM + ".docx")
    document.save(destination)
    _restore_package_and_footnote(source, destination, data["ai_note"])
    word_pdf = COURSE / (STEM + "-Word.pdf")
    receipt = export_word_pdf(destination, word_pdf)
    assert hashlib.sha256(source.read_bytes()).hexdigest() == source_hash
    assert hashlib.sha256(process_path.read_bytes()).hexdigest() == process_hash
    receipt.update(source_word_sha256=source_hash, process_sha256=process_hash, roles=8, process_steps=16, submitted=False)
    (ASSETS / "generacion.json").write_text(json.dumps(receipt, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps(receipt, ensure_ascii=True, indent=2))


if __name__ == "__main__":
    main()