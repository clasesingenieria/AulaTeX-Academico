import hashlib
import json
import shutil
from copy import deepcopy
from pathlib import Path

import pymupdf
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


HERE = Path(__file__).resolve().parent
SUBJECT = HERE.parents[1]
ROOT = SUBJECT.parents[2]


def preserve_pdf_body():
    destination = SUBJECT / 'Entregas/Tarea12_DeLaCruzMunoz.pdf'
    original = SUBJECT / 'referencias-seminario-i/notas-seminario-i/actividad-12-hipotesis/antes-portada-plan-negocios-2026-10-13/Tarea12_DeLaCruzMunoz.pdf'
    temporary = HERE / 'Tarea12-portada-cuerpo-verificado.pdf'
    with pymupdf.open(destination) as current, pymupdf.open(original) as previous, pymupdf.open() as output:
        output.insert_pdf(current, from_page=0, to_page=0)
        output.insert_pdf(previous, from_page=1)
        output.set_metadata(current.metadata)
        output.save(temporary, garbage=4, deflate=True)
    with pymupdf.open(temporary) as output, pymupdf.open(original) as previous:
        assert len(output) == len(previous)
        assert all(output[i].get_text() == previous[i].get_text() for i in range(1, len(output)))
        assert all(output[i].read_contents() == previous[i].read_contents() for i in range(1, len(output)))
        output[0].get_pixmap().save(HERE / 'portada-word-verificada.png')
    shutil.move(temporary, destination)
    print('OK: solo portada reemplazada; paginas 2-31 conservan texto y flujos originales.')


def main():
    source = ROOT / 'ITESCA/maestria-en-gestion-administrativa/plan-de-negocios-mga/assets-plan-de-negocios/actividad-12-2026-10-05/portada-oficial.png'
    image = HERE / 'portada-oficial.png'
    if image.exists():
        assert image.read_bytes() == source.read_bytes()
    else:
        shutil.copy2(source, image)
    word = SUBJECT / 'Entregas/Tarea12_DeLaCruzMunoz.docx'
    notes = SUBJECT / 'referencias-seminario-i/notas-seminario-i/actividad-12-hipotesis'
    backup = notes / 'antes-portada-plan-negocios-2026-10-13'
    if backup.exists():
        raise FileExistsError('La portada ya tiene respaldo; revisar antes de repetir')
    document = Document(word)
    paragraphs = document.paragraphs
    index = next(p for p in paragraphs if p.text == 'Índice')
    first = next(i for i, p in enumerate(paragraphs) if p._p is index._p)
    before = [element.xml for element in document.element.body if element is index._p or element in list(document.element.body)[list(document.element.body).index(index._p):]]
    tables = [table._tbl.xml for table in document.tables]
    metadata = 'Maestría en Gestión Administrativa\nAsignatura: Seminario I\nCurso: GGS102 - Unidad 3\nEstudiante: Martín Jonathan de la Cruz Muñoz\nMatrícula: 26130503\nDocente: Dra. Carla Olimpya Zapuche Moreno\nFecha: 12 de octubre de 2026\nLugar: Ciudad Obregón, Sonora'
    title = 'Gestión del cumplimiento normativo para formalizar vTaxi como plataforma de taxis en Nuevo León\n\nActividad 12 - Hipótesis de trabajo\nAnteproyecto acumulativo'
    with pymupdf.open() as cover:
        page = cover.new_page(width=612, height=792)
        page.insert_image(page.rect, filename=str(image))
        page.insert_font(fontname='Arial', fontfile='C:/Windows/Fonts/arial.ttf')
        page.insert_font(fontname='ArialBold', fontfile='C:/Windows/Fonts/arialbd.ttf')
        block = pymupdf.Rect(56.7, 294.8, 544.3, 454)
        page.draw_rect(block, fill=(1, 1, 1), color=None, fill_opacity=.96)
        assert page.insert_textbox(block + (12, 12, -12, -12), metadata, fontname='Arial', fontsize=11, lineheight=1.25) >= 0
        assert page.insert_textbox(pymupdf.Rect(263.6, 552.7, 582, 704), title, fontname='ArialBold', fontsize=12, lineheight=1.2, color=(1, 1, 1)) >= 0
        cover.save(HERE / 'portada-t12.pdf')
        page.get_pixmap(matrix=pymupdf.Matrix(2, 2)).save(HERE / 'portada-t12.png')
    cover_paragraph = paragraphs[0]
    for paragraph in paragraphs[:first]:
        if paragraph is not cover_paragraph:
            paragraph._p.getparent().remove(paragraph._p)
    cover_paragraph.clear()
    cover_paragraph.paragraph_format.space_before = 0
    cover_paragraph.paragraph_format.space_after = 0
    shape = cover_paragraph.add_run().add_picture(str(HERE / 'portada-t12.png'), width=document.sections[0].page_width, height=document.sections[0].page_height)
    inline = shape._inline
    anchor = OxmlElement('wp:anchor')
    for key, value in {'distT':'0','distB':'0','distL':'0','distR':'0','simplePos':'0','relativeHeight':'0','behindDoc':'1','locked':'1','layoutInCell':'1','allowOverlap':'1'}.items():
        anchor.set(key, value)
    simple = OxmlElement('wp:simplePos')
    simple.set('x','0'); simple.set('y','0')
    anchor.append(simple)
    for axis in ('H','V'):
        position = OxmlElement('wp:position' + axis)
        position.set('relativeFrom','page')
        offset = OxmlElement('wp:posOffset')
        offset.text = '0'
        position.append(offset); anchor.append(position)
    anchor.append(deepcopy(inline.find(qn('wp:extent'))))
    anchor.append(OxmlElement('wp:wrapNone'))
    for name in ('wp:docPr','wp:cNvGraphicFramePr','a:graphic'):
        element = inline.find(qn(name))
        if element is not None:
            anchor.append(deepcopy(element))
    inline.getparent().replace(inline, anchor)
    assert [table._tbl.xml for table in document.tables] == tables
    elements = list(document.element.body)
    assert [element.xml for element in elements[elements.index(index._p):]] == before
    backup.mkdir()
    shutil.copy2(word, backup / word.name)
    shutil.copy2(word.with_suffix('.pdf'), backup / word.with_suffix('.pdf').name)
    document.save(word)
    result = {'cover_only':True,'body_and_tables_xml_unchanged':True,'background_sha256':hashlib.sha256(image.read_bytes()).hexdigest(),'backup':str(backup.relative_to(SUBJECT)),'word':str(word.relative_to(SUBJECT)),'word_sha256_before_refresh':hashlib.sha256(word.read_bytes()).hexdigest(),'cover_metadata_is_image':True,'submission_performed':False}
    (HERE / 'verificacion-portada.json').write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding='utf-8')
    print(json.dumps(result, ensure_ascii=True))


if __name__ == '__main__':
    main()