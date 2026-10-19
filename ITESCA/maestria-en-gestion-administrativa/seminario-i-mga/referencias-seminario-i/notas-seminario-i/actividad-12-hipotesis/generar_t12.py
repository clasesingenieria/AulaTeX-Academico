import hashlib
import json
import shutil
import sys
from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.text.paragraph import Paragraph


HERE = Path(__file__).resolve().parent
SUBJECT = HERE.parents[2]
SOURCE = SUBJECT / 'entregas-seminario-i/Tarea11_DeLaCruzMunoz.docx'
TARGET = SUBJECT / 'entregas-seminario-i/Tarea12_DeLaCruzMunoz.docx'
GENERAL = ('En los componentes documentales aplicables al proyecto vTaxi como plataforma de taxis concesionados en Nuevo León durante septiembre a noviembre de 2026, un mayor nivel de gestión del cumplimiento normativo se relaciona con un mayor grado de preparación administrativa del expediente, expresado en documentos completos y coherentes con los requisitos aplicables y en observaciones internas atendidas con respaldo verificable.')
SPECIFIC = ('En los componentes del expediente de vTaxi, la asignación documentada de responsables y el seguimiento de observaciones se relacionan con una mayor proporción de observaciones internas cerradas mediante evidencia documental aceptable en la revisión del caso.')
REFERENCE = ('Price, P. C., Jhangiani, R. S., & Chiang, I.-C. A. (2015). Research methods in psychology (2nd Canadian ed.). BCcampus. https://opentextbc.ca/researchmethods/chapter/correlational-research/')
BOOK_REFERENCES = [
    'Hernández Sampieri, R., Fernández Collado, C., & Baptista Lucio, M. del P. (2014). Metodología de la investigación (6.ª ed.). McGraw-Hill Education.',
    'OECD. (2018). OECD regulatory enforcement and inspections toolkit. OECD Publishing. https://doi.org/10.1787/9789264303959-en',
]
FOUNDATION = [
    'Fundamento metodológico. Hernández Sampieri et al. (2014) definen las hipótesis como “explicaciones tentativas del fenómeno investigado” (p. 104). Por ello, HG es una respuesta provisional a la pregunta general, no una conclusión sobre vTaxi. Los autores vinculan su formulación con problema, antecedentes y marco teórico y exigen variables observables, relaciones precisas y técnicas disponibles para contrastarlas (pp. 105–107). La formulación conserva el caso, periodo y objetos documentales delimitados en los capítulos anteriores.',
    'Fundamento administrativo. El criterio 8.1 del toolkit de OECD (2018, p. 39) considera relevantes los datos actualizados, la gestión de casos, los registros y los flujos de trabajo para planear y dar seguimiento a la inspección y ejecución regulatoria. Se toma como orientación conceptual sobre trazabilidad, no como norma exigible a vTaxi: su destinatario son estructuras de inspección y supervisión, no plataformas privadas. Su adaptación al caso consiste en examinar responsables identificables, revisión de aplicabilidad y seguimiento registrado frente a documentos verificables. Esos controles pueden facilitar la localización de brechas y su revisión, pero un expediente puede permanecer incompleto pese a ellos; de ahí que se proponga una relación contrastable y no un efecto garantizado.',
    'Enlace con los antecedentes. Parker y Nielsen (2009), incorporados en el capítulo 1 con el alcance limitado de la revisión disponible, plantean el examen de sistemas de cumplimiento y su posible diferencia en la práctica. Ese antecedente y la orientación de OECD justifican investigar controles y evidencia, no dar por acreditada su eficacia. La inferencia de trabajo para vTaxi es propia: donde existen más controles documentados se espera encontrar mayor preparación del expediente, manteniendo separados los registros de gestión y la evaluación del documento resultante. Las divergencias también se registrarán y podrán debilitar la hipótesis.',
]
METHOD = [
    'HG y HE1 son hipótesis de investigación correlacionales por su contenido relacional, direccionales positivas y no causales. Esto clasifica la proposición, no anuncia un estudio de correlación estadística ya diseñado o ejecutado. El procedimiento previsto para este caso es contraste documental descriptivo por patrones de correspondencia y casos contrarios. No se propone experimento, manipulación, comparación de plataformas ni generalización estadística. Hernández Sampieri et al. (2014) señalan que “Correlación y causalidad son conceptos asociados, pero distintos” (p. 111); encontrar correspondencia no demuestra que la gestión haya producido la preparación del expediente.',
    'Convención de variables. Para responder al campo VI–VD de la consigna se identifica X: gestión del cumplimiento, y Y: preparación administrativa del expediente. En la ficha de esta actividad VI se emplea sólo como rótulo de X y VD como rótulo de Y; no se atribuyen causa y efecto. Hernández Sampieri et al. (2014, pp. 109–111) reservan la interpretación independiente/causal y dependiente/efecto para hipótesis causales; Price et al. (2015) también distinguen medición no experimental y causalidad. En el análisis se usarán X e Y y no se interpretará X como variable manipulada.',
    'La unidad de análisis es el proyecto vTaxi; los componentes del expediente y las observaciones son unidades de registro internas, no empresas independientes. Se compararán controles y calidad documental en cortes compatibles, sin tratar versiones del mismo documento como nuevos casos. No se calcularán coeficientes, valores p ni niveles de significación: no se ha previsto una prueba inferencial. La ausencia de variación o de datos evaluables impide contrastar, no confirma HG. Cualquier incorporación posterior de un análisis estadístico requeriría justificar previamente diseño, registros suficientes y supuestos; no se presupone en esta tarea.',
]
MEASUREMENT = 'Reglas observables de X e Y. Por componente, X se describirá mediante tres controles comprobables: responsable efectivamente designado, revisión de aplicabilidad fechada y seguimiento documentado; el perfil contará de cero a tres presencias sólo si los tres controles son evaluables. No verificable no se recodifica como ausencia ni entra en ese perfil. Y se observará por separado: proporción de contenidos exigibles presentes, concordancia de cada contenido con el requisito y proporción de observaciones cerradas con evidencia aceptable. Se reportarán los indicadores individualmente, sin sumar porcentajes o inventar ponderaciones. Integridad será contenidos exigibles presentes entre contenidos exigibles evaluables; cierre será observaciones cerradas respaldadas entre observaciones evaluables. Con denominador cero se anotará no calculable. La falta de observaciones no prueba integridad. Antes del cotejo se fijarán aplicabilidad y criterios de aceptación; sólo después se examinará si perfiles con más controles coinciden con mejores resultados documentales comparables.'
SECTIONS = [
    ('5.1 Hipótesis general', [
        'HG. ' + GENERAL,
        'La variable independiente (VI), denominada aquí variable predictora y no factor manipulado, es el nivel de gestión del cumplimiento normativo. Comprende asignación documentada de responsabilidades, verificación de aplicabilidad y seguimiento de acciones. La variable dependiente (VD), como resultado documental observado, es el grado de preparación administrativa del expediente. Comprende integridad, coherencia con requisitos aplicables y cierre respaldado de observaciones. Preparación no significa autorización gubernamental, operación legal acreditada ni formalización consumada.',
        'La hipótesis deriva de la brecha descrita en antecedentes y planteamiento: los borradores no acreditan una ruta validada entre requisitos, responsables y evidencias. Responde a la pregunta general y al objetivo de analizar esa relación mediante contraste documental. Los antecedentes justifican examinar procedimientos y controles administrativos, pero no demuestran anticipadamente la relación en vTaxi ni permiten trasladar hallazgos de otros contextos como resultados locales.',
    ]),
    ('5.2 Hipótesis específica', [
        'HE1. ' + SPECIFIC,
        'Se vincula con la pregunta específica 3 y el objetivo específico 3: examinar la correspondencia entre responsables, verificación y seguimiento, por una parte, e integridad del expediente y resolución de observaciones, por otra. Los objetivos 1 y 2 identifican requisitos y diagnostican brechas; el objetivo 4 propone prioridades. No requieren hipótesis adicionales porque no todos los objetivos descriptivos o propositivos expresan una relación entre variables.',
    ]),
    ('5.3 Tipo de hipótesis', [
        'HG y HE1 son hipótesis de trabajo relacionales, direccionales positivas y no causales, formuladas para un estudio de caso documental con alcance correlacional exploratorio en el sentido de contrastar correspondencias. No se propone experimento, manipulación de la VI, comparación entre plataformas ni generalización estadística. Price et al. (2015) distinguen medición de relaciones y demostración causal; aquí VI y VD designan roles analíticos exigidos por la formulación, no causalidad comprobada.',
        'La unidad de análisis sigue siendo vTaxi y su preparación administrativa. Los componentes y observaciones documentales son unidades de registro internas, no empresas independientes ni una muestra poblacional. El contraste será inicialmente descriptivo y por patrones: buscar coincidencias, divergencias y explicaciones alternativas. No se calculará un coeficiente de correlación si no existe una base suficiente y compatible con sus supuestos. No se incluyen hipótesis nula y alternativa estadísticas, valores p o niveles de significación porque no se ha previsto una prueba inferencial.',
    ]),
    ('5.4 Criterios de calidad y contraste documental', [
        'Claridad. Se diferencian gestión (acciones y registros de control) y preparación (calidad de los documentos resultantes). Se utilizará una ficha por componente con identificador, requisito aplicable, fuente, documento revisado, fecha, versión y limitaciones de acceso. La clasificación previa de taxis concesionados, SETIAP/ERT y capacitación evita atribuir requisitos de otra modalidad a vTaxi.',
        'Verificabilidad de la VI. Registrar por separado: responsable documentado (sí/no/no verificable); revisión de aplicabilidad con fuente y fecha (sí/no/no verificable); y seguimiento de acciones u observaciones mediante registro fechado (sí/no/no verificable). Distinguir responsable efectivamente designado de responsable sólo propuesto. La cobertura de cada control será el número de componentes con presencia comprobada dividido entre componentes evaluables; reportar numerador, denominador y casos no verificables. No confundir una recomendación con control ejecutado.',
        'Verificabilidad de la VD. Una lista distinta registrará integridad de los contenidos exigibles y concordancia del documento con el requisito aplicable, sin puntuar como preparación la mera existencia de responsable o bitácora. Cada observación se clasificará abierta, cerrada con evidencia aceptable o no verificable. La proporción de cierre será observaciones con cierre respaldado entre observaciones evaluables en el mismo corte; informar los conteos. Ausencia de observaciones registradas no equivale a expediente completo ni se divide entre cero.',
        'Separación y revisión. Los criterios de aceptabilidad documental se fijarán antes de comparar VI y VD, con fundamento y alcance por requisito. Cuando sea viable, una segunda revisión documentará acuerdos y discrepancias; no se afirma validación experta ya realizada. Una fuente inaccesible se registra no verificable, no incumplida. Comparar sólo componentes con exigencias y cortes compatibles; las versiones repetidas del mismo documento no cuentan como casos independientes.',
        'Contraste de HG. En una matriz por componente se cotejarán los controles observados con integridad y coherencia documental en el mismo corte. Se buscará el patrón de mejor preparación cuando existen más controles documentados y se revisarán casos contrarios: controles presentes con expediente incompleto, o documentos adecuados sin controles registrados. Tales casos pueden debilitar o acotar la hipótesis. Si no hay variación, acceso suficiente o criterios comparables, el resultado será no contrastable con la evidencia disponible, no hipótesis confirmada.',
        'Contraste de HE1. Por observación se examinarán designación documentada, seguimiento fechado y evidencia de cierre. Se distinguirá el registro de acciones de la corrección verificable del documento. Se atenderán explicaciones alternativas: complejidad del requisito, disponibilidad de información y etapa de elaboración. La correspondencia observada no prueba que los controles hayan causado el cierre. Las relaciones se describirán sólo dentro del caso y con sus limitaciones.',
        'Pertinencia y consistencia. Las hipótesis responden al problema y objetivos sin sustituirlos ni ampliar el estudio a todas las plataformas. Su interés pertenece a la LGAC Gestión e Innovación de las Organizaciones: procedimientos, responsables, documentos, controles e indicadores. El Círculo de Covey conserva problema y objetivo general en Importante/No urgente; no se agrega la hipótesis al esquema. La clasificación establece una prioridad de planeación, no una prueba de la relación propuesta.',
        'Estado de evidencia. Se formulan hipótesis para orientar el estudio, no resultados. Los borradores y fuentes disponibles permiten preparar la revisión inicial, pero no se cuenta aquí con mediciones completas, entrevistas realizadas, controles implementados ni cierres verificados. El acceso adicional requerirá autorización; no se fabricarán registros ni se tratará la falta de datos como confirmación. Los alcances y procedimientos detallados se desarrollarán en sus apartados posteriores sin anticipar su ejecución.',
    ]),
]


def improve():
    if not TARGET.is_file():
        raise FileNotFoundError('Generar T12 antes de aplicar la revisión')
    backup = HERE / 'antes-mejora-2026-10-13' / TARGET.name
    if backup.exists():
        raise FileExistsError('La revisión ya tiene respaldo; no se repite automáticamente')
    document = Document(TARGET)
    paragraphs = document.paragraphs
    chapter = next(p for p in paragraphs if p.text == '5 Hipótesis de trabajo' and p.style.name == 'Heading 1')
    end = next(p for p in paragraphs if p.text == '6 Alcances y limitaciones' and p.style.name == 'Heading 1')
    chapter_index = next(index for index, paragraph in enumerate(paragraphs) if paragraph._p is chapter._p)
    previous = [p._p.xml for p in paragraphs[:chapter_index]]
    tables = [table._tbl.xml for table in document.tables]
    general = next(p for p in paragraphs if p.text == 'HG. ' + GENERAL)
    variables = next(p for p in paragraphs if p.text == SECTIONS[0][1][1])
    variables.text = ('X (campo VI de la consigna) es el nivel de gestión del cumplimiento normativo: asignación documentada de responsabilidades, verificación de aplicabilidad y seguimiento de acciones. Y (campo VD) es el grado de preparación administrativa del expediente: integridad, coherencia con requisitos aplicables y cierre respaldado de observaciones. Son roles analíticos no causales; X no está manipulada ni Y es un efecto demostrado. Preparación no significa autorización gubernamental, operación legal acreditada ni formalización consumada.')
    section_two = next(p for p in document.paragraphs if p.text == '5.2 Hipótesis específica' and p.style.name == 'Heading 2')
    for text in FOUNDATION:
        insert_before(section_two, text)
    method_start = next(p for p in document.paragraphs if p.text == '5.3 Tipo de hipótesis' and p.style.name == 'Heading 2')
    quality_start = next(p for p in document.paragraphs if p.text == '5.4 Criterios de calidad y contraste documental' and p.style.name == 'Heading 2')
    for text in SECTIONS[2][1]:
        paragraph = next(p for p in document.paragraphs if p.text == text)
        paragraph._p.getparent().remove(paragraph._p)
    for text in METHOD:
        insert_before(quality_start, text)
    quality_text = next(p for p in document.paragraphs if p.text == SECTIONS[3][1][0])
    insert_before(quality_text, MEASUREMENT)
    for reference_text, prefix in zip(BOOK_REFERENCES, ('Instituto de Capacitación', 'Omrani,')):
        anchor = next(p for p in document.paragraphs if p.text.startswith(prefix) and p.style.name == 'Normal')
        reference = insert_before(anchor, reference_text)
        if anchor._p.pPr is not None:
            reference._p.remove(reference._p.pPr)
            reference._p.insert(0, deepcopy(anchor._p.pPr))
    assert general.text == 'HG. ' + GENERAL
    assert next(p for p in document.paragraphs if p.text == 'HE1. ' + SPECIFIC)
    revised = document.paragraphs
    revised_index = next(index for index, paragraph in enumerate(revised) if paragraph._p is chapter._p)
    assert [p._p.xml for p in revised[:revised_index]] == previous
    assert [table._tbl.xml for table in document.tables] == tables
    backup.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(TARGET, backup)
    pdf = TARGET.with_suffix('.pdf')
    if pdf.exists():
        shutil.copy2(pdf, backup.with_suffix('.pdf'))
    document.save(TARGET)
    result = {'operation': 'mejora-fundamento-t12', 'before_sha256': hashlib.sha256(backup.read_bytes()).hexdigest(), 'target_sha256_before_word_refresh': hashlib.sha256(TARGET.read_bytes()).hexdigest(), 'chapters_1_4_xml_unchanged': True, 'matrix_and_covey_xml_unchanged': True, 'hypotheses_unchanged': True, 'short_quotes': 2, 'new_references': BOOK_REFERENCES, 'submission_performed': False}
    (HERE / 'mejora-t12-2026-10-13.json').write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding='utf-8')
    print(json.dumps(result, ensure_ascii=True))


def insert_before(anchor, text, style='Normal'):
    element = OxmlElement('w:p')
    anchor._p.addprevious(element)
    paragraph = Paragraph(element, anchor._parent)
    paragraph.style = style
    paragraph.add_run(text)
    return paragraph


def main():
    if TARGET.exists():
        raise FileExistsError('T12 ya existe; no se sobrescribe')
    source_hash = hashlib.sha256(SOURCE.read_bytes()).hexdigest()
    document = Document(SOURCE)
    original = [paragraph.text for paragraph in document.paragraphs]
    tables = [[[cell.text for cell in row.cells] for row in table.rows] for table in document.tables]
    chapter = next(paragraph for paragraph in document.paragraphs if paragraph.text == '5 Hipótesis de trabajo' and paragraph.style.name == 'Heading 1')
    anchor = next(paragraph for paragraph in document.paragraphs if paragraph.text == '6 Alcances y limitaciones' and paragraph.style.name == 'Heading 1')
    assert chapter._p.getnext() is anchor._p
    inserted = []
    for heading, texts in SECTIONS:
        insert_before(anchor, heading, 'Heading 2')
        inserted.append(heading)
        for text in texts:
            insert_before(anchor, text)
            inserted.append(text)
    reference_anchor = next(paragraph for paragraph in document.paragraphs if paragraph.text.startswith('Ramírez-Solís,') and paragraph.style.name == 'Normal')
    reference = insert_before(reference_anchor, REFERENCE)
    if reference_anchor._p.pPr is not None:
        if reference._p.pPr is not None:
            reference._p.remove(reference._p.pPr)
        reference._p.insert(0, deepcopy(reference_anchor._p.pPr))
    inserted.append(REFERENCE)
    cover_changes = {}
    for paragraph in document.paragraphs:
        if paragraph.text == 'Seminario I\nActividad 11 Justificación':
            cover_changes[paragraph.text] = 'Seminario I\nActividad 12 Hipótesis de trabajo'
        elif re_date(paragraph.text):
            cover_changes[paragraph.text] = paragraph.text.replace('10 de octubre de 2026', '12 de octubre de 2026').replace('10 DE OCTUBRE DE 2026', '12 DE OCTUBRE DE 2026')
        if paragraph.text in cover_changes:
            old = paragraph.text
            if paragraph.runs:
                paragraph.runs[0].text = cover_changes[old]
                for run in paragraph.runs[1:]:
                    run.text = ''
            else:
                paragraph.add_run(cover_changes[old])
    assert any('Actividad 12' in paragraph.text for paragraph in document.paragraphs)
    retained = [paragraph.text for paragraph in document.paragraphs if paragraph.text not in inserted]
    assert retained == [cover_changes.get(text, text) for text in original]
    matrix = document.tables[0]
    row = matrix.add_row()
    row.cells[0].merge(row.cells[-1]).text = 'Hipótesis general (HG)\n' + GENERAL
    specific_matrix = document.tables[1]
    specific_matrix.rows[3].cells[1].add_paragraph('Hipótesis específica (HE1)\n' + SPECIFIC)
    assert [[cell.text for cell in row.cells] for row in document.tables[2].rows] == tables[2]
    update = OxmlElement('w:updateFields')
    update.set('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}val', 'true')
    document.settings.element.append(update)
    document.save(TARGET)
    result = {'source_sha256': source_hash, 'target_sha256': hashlib.sha256(TARGET.read_bytes()).hexdigest(), 'chapters_1_4_retained': True, 'hypothesis_sections': [heading for heading, _ in SECTIONS], 'matrix_updated': True, 'covey_unchanged': True, 'submission_performed': False, 'word_fields_refresh_pending': True}
    assert hashlib.sha256(SOURCE.read_bytes()).hexdigest() == source_hash
    (HERE / 'generacion-t12.json').write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding='utf-8')
    print(json.dumps(result, ensure_ascii=True))


def apply_feedback():
    backup_dir = HERE / 'antes-retroalimentacion-2026-10-18'
    if backup_dir.exists():
        raise FileExistsError('La revisión ya tiene respaldo; comprobar el estado antes de repetir')
    source_hash = hashlib.sha256(SOURCE.read_bytes()).hexdigest()
    document = Document(TARGET)
    original_covey = [[cell.text for cell in row.cells] for row in document.tables[2].rows]
    replacements = {
        'plataforma de taxis en Nuevo León': 'plataforma de taxis en Monterrey, Nuevo León',
        'plataforma de taxis concesionados en Nuevo León': 'plataforma de taxis concesionados en Monterrey, Nuevo León',
        'El ámbito espacial es Nuevo León': 'El ámbito espacial es el municipio de Monterrey, Nuevo León',
    }
    paragraphs = list(document.paragraphs)
    seen = set()
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell._tc in seen:
                    continue
                seen.add(cell._tc)
                paragraphs.extend(cell.paragraphs)
    changed = 0
    for paragraph in paragraphs:
        expected = paragraph.text
        for old, new in replacements.items():
            expected = expected.replace(old, new)
        if expected != paragraph.text:
            for run in paragraph.runs:
                for old, new in replacements.items():
                    run.text = run.text.replace(old, new)
            if paragraph.text != expected:
                raise ValueError('Reemplazo territorial dividido entre runs; revisar antes de guardar')
            changed += 1
    additions = [
        ('4.2 Relevancia práctica', [
            'El enfoque conceptual distingue adopción formal, gestión práctica y preparación documental. Parker y Nielsen (2009), con el alcance de resumen editorial señalado en los antecedentes, permiten separar la presencia de un sistema de cumplimiento de su administración efectiva. OECD (2018, p. 39) orienta la atención hacia registros actualizados, gestión de casos y flujos de trabajo en organismos de supervisión; su aplicación a vTaxi es analítica, no una obligación adicional impuesta al proyecto. El mecanismo propuesto es que identificar responsables y revisar requisitos facilita detectar y atender brechas; la integridad del documento resultante se evalúa por separado. La complejidad del requisito y el acceso a información pueden explicar divergencias. Así, X describe controles observables y Y describe calidad del expediente, sin definir una variable mediante la otra ni anticipar causalidad.'
        ]),
        ('5 Hipótesis de trabajo', [
            'El acceso se gestionará mediante una solicitud escrita al responsable autorizado del proyecto, con propósito, periodo y lista de documentos requeridos. La autorización deberá precisar qué versiones pueden consultarse, qué datos deben anonimizarse, quién tendrá acceso a las copias y qué extractos pueden incorporarse al informe. Se llevará un registro con solicitud, respuesta, alcance concedido y restricciones por documento. Los materiales no autorizados o no disponibles se clasificarán como no verificables; no se interpretarán como incumplimientos ni se reemplazarán por registros ficticios.',
            'Las entrevistas, si resultan necesarias, requerirán invitación y consentimiento individual separados del permiso documental. Se informarán finalidad, participación voluntaria, posibilidad de omitir respuestas o retirarse, uso de citas y condiciones de resguardo. La grabación exigirá consentimiento específico; sin él se utilizarán únicamente notas autorizadas. El plazo de conservación y la eliminación de archivos se acordarán antes de recopilar datos. Si no se obtienen permisos dentro del periodo, se mantendrá el alcance documental autorizado y se reportarán las limitaciones sin atribuir testimonios ni verificaciones inexistentes.',
            'La delimitación geográfica del caso es el municipio de Monterrey, Nuevo León. El objeto no es el mercado de transporte de toda la entidad ni el conjunto del área metropolitana. La legislación y las competencias institucionales se examinan a escala estatal porque ese es su ámbito de aplicación; esta escala normativa no amplía la unidad de análisis territorial. Los registros de vTaxi considerados deberán referirse al caso delimitado y se identificarán por fecha y versión.'
        ]),
    ]
    for heading, texts in additions:
        anchor = next(paragraph for paragraph in document.paragraphs if paragraph.text == heading and paragraph.style.name.startswith('Heading'))
        for text in texts:
            insert_before(anchor, text)
    expected_covey = original_covey
    for row in expected_covey:
        for index, text in enumerate(row):
            for old, new in replacements.items():
                text = text.replace(old, new)
            row[index] = text
    assert [[cell.text for cell in row.cells] for row in document.tables[2].rows] == expected_covey
    assert changed >= 5
    assert len(document.tables) == 3
    assert any(paragraph.text.startswith('HG. ') for paragraph in document.paragraphs)
    assert any(paragraph.text.startswith('HE1. ') for paragraph in document.paragraphs)
    backup_dir.mkdir()
    shutil.copy2(TARGET, backup_dir / TARGET.name)
    if TARGET.with_suffix('.pdf').exists():
        shutil.copy2(TARGET.with_suffix('.pdf'), backup_dir / TARGET.with_suffix('.pdf').name)
    document.save(TARGET)
    assert hashlib.sha256(SOURCE.read_bytes()).hexdigest() == source_hash
    record = {'operation': 'retroalimentacion-t11-en-t12', 'changed_geographic_paragraphs': changed, 'added_paragraphs': sum(len(texts) for _, texts in additions), 't11_unchanged_sha256': source_hash, 'covey_only_territorial_alignment': True, 'target_sha256_before_word_refresh': hashlib.sha256(TARGET.read_bytes()).hexdigest(), 'submission_performed': False}
    (HERE / 'retroalimentacion-t12-2026-10-18.json').write_text(json.dumps(record, ensure_ascii=False, indent=2), encoding='utf-8')
    print(json.dumps(record, ensure_ascii=True))


def add_mexican_context():
    document = Document(TARGET)
    if any('Lamas Flores y Acosta Reveles (2025)' in paragraph.text for paragraph in document.paragraphs):
        raise ValueError('El antecedente mexicano ya está incorporado')
    tables = [table._tbl.xml for table in document.tables]
    anchor = next(paragraph for paragraph in document.paragraphs if paragraph.text == 'Vacíos identificados')
    insert_before(anchor, 'En México, Lamas Flores y Acosta Reveles (2025) estudian las estrategias de cabildeo de plataformas de transporte y reparto desde un enfoque cualitativo y analítico. Combinan seguimiento de etiquetas de Twitter entre 2015 y 2018, revisión hemerográfica y entrevistas semiestructuradas a un integrante de cada una de dos organizaciones de conductores y repartidores. Las autoras identifican acciones mediante medios de comunicación, redes sociales y movilización de apoyos con recursos jurisdiccionales para influir en la regulación. Este antecedente muestra la importancia del entorno institucional y de distinguir la información empresarial de la evidencia verificable; no mide controles internos de vTaxi ni permite generalizar a todas las plataformas mexicanas.')
    insert_before(anchor, 'La diferencia del presente caso es su escala organizacional: examina la correspondencia entre controles de gestión y preparación de documentos de una plataforma para taxis concesionados en Monterrey. El estudio de Lamas Flores y Acosta Reveles (2025) aporta contexto mexicano sobre formación y controversia de reglas, mientras Parker y Nielsen (2009) orientan la distinción entre sistemas formales y gestión práctica. Ninguno sustituye el diagnóstico del expediente. La aportación propuesta consiste en hacer trazable cada requisito aplicable, responsable, evidencia y observación, sin atribuir a vTaxi prácticas de las empresas estudiadas ni trasladarles sus resultados.')
    anchor = next(paragraph for paragraph in document.paragraphs if paragraph.text == '4.5 Viabilidad' and paragraph.style.name == 'Heading 2')
    insert_before(anchor, 'Frente al antecedente mexicano de Lamas Flores y Acosta Reveles (2025), centrado en actores y estrategias de influencia regulatoria, la originalidad del caso reside en la evaluación interna de controles y calidad documental antes de una eventual formalización. La comparación precisa una diferencia de objeto, escala y evidencia; no constituye una afirmación de primacía ni agota la literatura sobre movilidad en México. La revisión deberá ampliarse conforme avance el estudio, en especial con trabajos directamente referidos a Monterrey y a taxis concesionados.')
    reference_anchor = next(paragraph for paragraph in document.paragraphs if paragraph.text.startswith('Melo,') and paragraph.style.name == 'Normal')
    reference = insert_before(reference_anchor, '')
    if reference_anchor._p.pPr is not None:
        reference._p.remove(reference._p.pPr)
        reference._p.insert(0, deepcopy(reference_anchor._p.pPr))
    reference.add_run('Lamas Flores, M. S., & Acosta Reveles, I. L. (2025). El cabildeo mediático y en redes sociales como freno a la regulación de plataformas digitales de transporte y reparto en México. ')
    reference.add_run('Cuadernos del Ciesal, 2').italic = True
    reference.add_run('(24), 1–28. https://doi.org/10.35305/cc.v2i24.174')
    assert [table._tbl.xml for table in document.tables] == tables
    document.save(TARGET)
    print(json.dumps({'mexican_context_added': True, 'new_reference_doi': '10.35305/cc.v2i24.174', 'tables_unchanged': True}, ensure_ascii=True))


def re_date(text):
    return '10 de octubre de 2026' in text.casefold()


if __name__ == '__main__':
    if sys.argv[1:] == ['--contexto-mexicano']:
        add_mexican_context()
    elif sys.argv[1:] == ['--retroalimentacion']:
        apply_feedback()
    elif sys.argv[1:] == ['--mejorar']:
        improve()
    else:
        main()