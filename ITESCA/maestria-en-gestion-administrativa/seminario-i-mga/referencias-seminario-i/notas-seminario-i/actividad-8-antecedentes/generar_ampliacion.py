import hashlib
import json
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.shared import Inches
from docx.text.paragraph import Paragraph

HERE = Path(__file__).resolve().parent
SUBJECT = HERE.parents[2]


def insert_before(anchor, text, style):
    element = OxmlElement("w:p")
    anchor._p.addprevious(element)
    paragraph = Paragraph(element, anchor._parent)
    paragraph.style = style
    paragraph.add_run(text)
    return paragraph


def main():
    source = SUBJECT / "Entregas/Tarea8_DeLaCruzMunoz.docx"
    target = SUBJECT / "Entregas/Tarea8_DeLaCruzMunoz_Revision-2026-10-10.docx"
    if target.exists():
        raise FileExistsError("No se sobrescribe una revision existente")
    data = json.loads((HERE / "ampliacion-antecedentes-2026-10-10.json").read_text(encoding="utf-8"))
    original = Document(source)
    document = Document(source)
    anchor = next(paragraph for paragraph in document.paragraphs if paragraph.style.name == "Heading 1" and paragraph.text.startswith("2 "))
    inserted = [data["heading"], *data["paragraphs"], *data["references"]]
    insert_before(anchor, data["heading"], "Heading 2")
    for text in data["paragraphs"]:
        insert_before(anchor, text, "Normal")
    for text in data["references"]:
        paragraph = document.add_paragraph(text, "Normal")
        paragraph.paragraph_format.left_indent = Inches(0.5)
        paragraph.paragraph_format.first_line_indent = Inches(-0.5)
    assert [paragraph.text for paragraph in document.paragraphs if paragraph.text not in inserted] == [paragraph.text for paragraph in original.paragraphs]
    assert [table._tbl.xml for table in document.tables] == [table._tbl.xml for table in original.tables]
    document.save(target)
    result = {"source": source.relative_to(SUBJECT).as_posix(), "source_sha256": hashlib.sha256(source.read_bytes()).hexdigest(), "revision": target.relative_to(SUBJECT).as_posix(), "revision_sha256": hashlib.sha256(target.read_bytes()).hexdigest(), "previous_content_retained": True, "tables_retained": True, "new_institutional_sources": 2, "toc_pending_refresh": True, "submitted": False}
    (HERE / "generacion-ampliacion.json").write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(result, ensure_ascii=True))


if __name__ == "__main__":
    main()