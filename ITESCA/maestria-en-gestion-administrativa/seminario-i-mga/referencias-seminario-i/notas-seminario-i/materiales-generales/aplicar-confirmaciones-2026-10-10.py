import hashlib
import json
from pathlib import Path
from shutil import copy2

from docx import Document
from docx.shared import Inches

HERE = Path(__file__).resolve().parent
SUBJECT = HERE.parents[2]
NOTES = SUBJECT / "referencias-seminario-i/notas-seminario-i"

REGIONAL = (
    "La delimitación actual del caso es Monterrey, Nuevo León, por indicación del estudiante. "
    "El municipio no se equipara automáticamente con el área metropolitana ni con toda la entidad. "
    "Los Censos Económicos 2024 y el Censo de Población y Vivienda 2020 del INEGI constituyen "
    "fuentes para contextualizar establecimientos y población, pero los datos agregados del portal "
    "no se presentan como cifras municipales ni como demanda validada de vTaxi. La ley de movilidad "
    "de Nuevo León corresponde al ámbito normativo estatal. Los antecedentes de Cajeme se conservan "
    "como referencias históricas, no como evidencia empírica del caso de Monterrey."
)
REFERENCES = [
    "Instituto Nacional de Estadística y Geografía. (2024). Censos Económicos 2024. https://www.inegi.org.mx/programas/ce/2024/",
    "Instituto Nacional de Estadística y Geografía. (2020). Censo de Población y Vivienda 2020. https://www.inegi.org.mx/programas/ccpv/2020/",
]


def main():
    backup = HERE / "antes-confirmaciones-2026-10-10"
    backup.mkdir(exist_ok=True)
    results = []
    for name in ["Tarea8_DeLaCruzMunoz_Revision-2026-10-10.docx", "Tarea11_DeLaCruzMunoz.docx"]:
        path = SUBJECT / "Entregas" / name
        old_hash = hashlib.sha256(path.read_bytes()).hexdigest()
        if (backup / name).exists():
            raise FileExistsError("Confirmaciones ya aplicadas; no duplicar contenido")
        copy2(path, backup / name)
        document = Document(path)
        original_tables = [[[cell.text for cell in row.cells] for row in table.rows] for table in document.tables]
        if name.startswith("Tarea11"):
            anchor = next(paragraph for paragraph in document.paragraphs if paragraph.text == "4.3 Relevancia social e institucional")
        else:
            anchor = next(paragraph for paragraph in document.paragraphs if paragraph.text == "2 Planteamiento del problema" and paragraph.style.name == "Heading 1")
        anchor.insert_paragraph_before(REGIONAL, "Normal")
        for reference in REFERENCES:
            annex = next((paragraph for paragraph in document.paragraphs if paragraph.style.name == "Heading 1" and paragraph.text.startswith("Anexo")), None)
            paragraph = annex.insert_paragraph_before(reference, "Normal") if annex else document.add_paragraph(reference)
            paragraph.paragraph_format.left_indent = Inches(0.5)
            paragraph.paragraph_format.first_line_indent = Inches(-0.5)
        first = document.paragraphs[0]
        first.insert_paragraph_before("Semestre: primero. Delimitación del estudio: Monterrey, Nuevo León.", "Normal")
        assert original_tables == [[[cell.text for cell in row.cells] for row in table.rows] for table in document.tables]
        document.save(path)
        results.append({"archivo": name, "hash_anterior": old_hash, "hash_nuevo": hashlib.sha256(path.read_bytes()).hexdigest(), "tablas_conservadas": True, "indice_por_actualizar": True})
    result = {"fecha": "2026-10-10", "confirmaciones": {"semestre": 1, "region": "Monterrey, Nuevo Leon", "t5": "omitida por el usuario, no resuelta"}, "documentos": results, "enviado": False}
    (HERE / "confirmaciones-2026-10-10.json").write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(result, ensure_ascii=True))


if __name__ == "__main__":
    main()