from pathlib import Path
import hashlib
import json
import shutil
from docx import Document

ROOT = Path(__file__).resolve().parents[1]
SUBJECT = ROOT / "ITESCA/maestria-en-gestion-administrativa/seminario-i-mga"
WORD = SUBJECT / "Entregas/Tarea11_DeLaCruzMunoz.docx"
EVIDENCE = SUBJECT / "referencias-seminario-i/notas-seminario-i/actividad-11-justificacion/comprobantes-envio/2026-10-10"
EVIDENCE.mkdir(parents=True, exist_ok=True)
backup = EVIDENCE / "Tarea11-antes-correccion-portada.docx"
if backup.exists():
    if backup.read_bytes() != WORD.read_bytes():
        raise RuntimeError("El respaldo no coincide con el original")
else:
    shutil.copy2(WORD, backup)
document = Document(WORD)
before = [paragraph.text for paragraph in document.paragraphs]
tables = [[list(cell.text for cell in row.cells) for row in table.rows] for table in document.tables]
replacements = {
    "Actividad 10 Formulación de objetivos": "Actividad 11 Justificación",
    "4 DE OCTUBRE DE 2026": "10 DE OCTUBRE DE 2026",
}
counts = dict.fromkeys(replacements, 0)
for paragraph in document.paragraphs[:25]:
    for old, new in replacements.items():
        if old in paragraph.text:
            replacement = paragraph.text.replace(old, new)
            runs = paragraph.runs
            runs[0].text = replacement
            for run in runs[1:]:
                run.text = ""
            counts[old] += 1
if not all(count == 1 for count in counts.values()):
    raise RuntimeError("Metadatos de portada inesperados")
document.save(WORD)
updated = Document(WORD)
assert [p.text for p in updated.paragraphs][25:] == before[25:]
assert [[list(cell.text for cell in row.cells) for row in table.rows] for table in updated.tables] == tables
print(json.dumps({"cover_updated": True, "body_unchanged": True, "tables_unchanged": True, "sha256": hashlib.sha256(WORD.read_bytes()).hexdigest()}))