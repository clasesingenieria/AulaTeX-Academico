"""Genera la entrega T13 (Word acumulativo) a partir del Word T12 enviado y
sincroniza el reporte de raíz reporte-seminario-i-Actividad-13.tex.

- Fuente: entregas-seminario-i/Tarea12_DeLaCruzMunoz.docx (enviado 18/10/2026,
  con retroalimentación de T11 incorporada). No se modifica.
- Destino: entregas-seminario-i/Tarea13_DeLaCruzMunoz.docx.
- Capítulo 6 (6.1 Alcances y 6.2 Limitaciones), 300--500 palabras.
- Matriz de consistencia: nota de alineación de alcances. Covey sin elementos nuevos.
- Portada gráfica: metadatos de T13.
El generador no realiza envíos.
"""
import hashlib
import json
from copy import deepcopy
from io import BytesIO
from pathlib import Path
import runpy

import pymupdf
from docx import Document
from docx.oxml import OxmlElement
from docx.table import Table
from docx.text.paragraph import Paragraph

HERE = Path(__file__).resolve().parent
SUBJECT = HERE.parents[2]
SOURCE = SUBJECT / 'entregas-seminario-i/Tarea12_DeLaCruzMunoz.docx'
TARGET = SUBJECT / 'entregas-seminario-i/Tarea13_DeLaCruzMunoz.docx'
REPORT = SUBJECT / 'reporte-seminario-i-Actividad-13.tex'
COVER_BG = SUBJECT / 'assets-seminario-i/portada/portada-oficial.png'
FECHA = '25 de octubre de 2026'
TITULO = 'Gestión del cumplimiento normativo para formalizar vTaxi como plataforma de taxis en Monterrey, Nuevo León'

SCOPES = [
    ('Alcance temático.', 'El estudio se restringe a la gestión administrativa del cumplimiento normativo en vTaxi: procedimientos internos, asignación de responsables, verificación de aplicabilidad y seguimiento documental, contrastados con la integridad, la concordancia y el cierre respaldado de observaciones del expediente. La normativa estatal sirve para definir criterios del expediente, no para emitir un dictamen jurídico. Quedan fuera el desarrollo del software de despacho, la rentabilidad del modelo de negocio, la demanda del servicio y la defensa contenciosa ante autoridades (Parker & Nielsen, 2009; vTaxi, 2026).'),
    ('Alcance espacial y temporal.', 'La unidad territorial es el municipio de Monterrey, Nuevo León. La Ley de Movilidad se examina a escala estatal por ser su ámbito de aplicación, sin extender los resultados a toda la entidad, al área metropolitana ni a la modalidad SETIAP/ERT (H. Congreso del Estado de Nuevo León, 2026). El periodo comprende septiembre a noviembre de 2026: la documentación con corte al 27 de septiembre funciona como línea de base, octubre se dedica a la revisión de requisitos y noviembre al contraste y la priorización. Son actividades programadas, no ejecutadas.'),
    ('Alcance metodológico y resultados esperados.', 'Se trata de un estudio de caso único, no experimental, documental, descriptivo y propositivo, con contraste por patrones entre los controles de gestión (X) y la calidad del expediente (Y). Se esperan cuatro productos vinculados con los objetivos específicos: catálogo de requisitos aplicables, diagnóstico de brechas, matriz de correspondencias y ruta priorizada de atención. No se demostrará causalidad ni se generalizarán resultados estadísticamente, pues las variables y técnicas deben ser observables y disponibles (Hernández Sampieri et al., 2014, pp. 106–107).'),
]
LIMITS = [
    ('Limitación teórica.', 'Los antecedentes sobre transformación digital en MiPyMEs explican las capacidades organizacionales de forma indirecta, y Parker y Nielsen (2009) se han utilizado con el alcance de su resumen editorial. Ello reduce la fuerza explicativa del marco para un caso de movilidad en México. Para atenuarlo, se ampliará la revisión sobre compliance organizacional y formalización de plataformas, distinguiendo las orientaciones conceptuales de la evidencia local.'),
    ('Limitación metodológica y de acceso.', 'vTaxi se encuentra en planeación; el corpus autorizado se limita a borradores, ficha técnica y resúmenes regulatorios, sin contratos ni actas protocolizadas. Puede faltar variación suficiente de controles o cierres para contrastar HG y HE1, en cuyo caso el resultado será no contrastable, no hipótesis confirmada. Las versiones de un mismo documento no se contarán como casos independientes; la información faltante se registrará como no verificable y no como incumplimiento. Las entrevistas quedan condicionadas a autorización y consentimiento.'),
    ('Limitación técnica y temporal.', 'Los datos de ejemplo de los borradores, las fuentes inaccesibles y las reformas normativas pueden impedir comprobar la aplicabilidad o vigencia de un requisito en el corte. Se registrarán procedencia, fecha y versión de cada fuente. El cierre de noviembre restringe verificaciones adicionales, por lo que se priorizarán las brechas documentables conforme al Círculo de Covey. Los resultados no acreditarán concesiones, operación legal ni autorización gubernamental.'),
]
MATRIX_NOTE = ('Alcance (T13): estudio de caso documental de vTaxi en Monterrey, Nuevo León, septiembre a noviembre de 2026. '
               'OE1: catálogo de requisitos; OE2: diagnóstico de brechas; OE3: contraste por patrones entre controles y expediente; '
               'OE4: ruta priorizada. Faltantes registrados como no verificables; sin causalidad ni generalización estadística. '
               'El Círculo de Covey conserva problema y objetivo general en Importante/No urgente.')


def insert_before(anchor, label, text, ppr=None):
    element = OxmlElement('w:p')
    anchor._p.addprevious(element)
    paragraph = Paragraph(element, anchor._parent)
    paragraph.style = 'Normal'
    if ppr is not None:
        paragraph._p.insert(0, deepcopy(ppr))
    paragraph.add_run(label + ' ').bold = True
    paragraph.add_run(text)
    return paragraph


def make_cover():
    metadata = ('Maestría en Gestión Administrativa\nAsignatura: Seminario I\nCurso: GGS102 - Unidad 3\n'
                'Estudiante: Martín Jonathan de la Cruz Muñoz\nMatrícula: 26130503\n'
                f'Docente: Dra. Carla Olimpya Zapuche Moreno\nFecha: {FECHA}\nLugar: Monterrey, Nuevo León')
    title = TITULO + '\n\nActividad 13 - Alcances y limitaciones\nAnteproyecto acumulativo'
    with pymupdf.open() as pdf:
        page = pdf.new_page(width=612, height=792)
        page.insert_image(page.rect, filename=str(COVER_BG))
        page.insert_font(fontname='Arial', fontfile='C:/Windows/Fonts/arial.ttf')
        page.insert_font(fontname='ArialBold', fontfile='C:/Windows/Fonts/arialbd.ttf')
        block = pymupdf.Rect(56.7, 294.8, 544.3, 454)
        page.draw_rect(block, fill=(1, 1, 1), color=None, fill_opacity=.96)
        assert page.insert_textbox(block + (12, 12, -12, -12), metadata, fontname='Arial', fontsize=11, lineheight=1.25) >= 0
        assert page.insert_textbox(pymupdf.Rect(263.6, 552.7, 582, 704), title, fontname='ArialBold', fontsize=12, lineheight=1.2, color=(1, 1, 1)) >= 0
        page.get_pixmap(matrix=pymupdf.Matrix(2, 2)).save(HERE / 'portada-t13.png')
        return page.get_pixmap(matrix=pymupdf.Matrix(2, 2)).tobytes('png')


def build_word():
    source_hash = hashlib.sha256(SOURCE.read_bytes()).hexdigest()
    document = Document(SOURCE)
    paragraphs = document.paragraphs
    original = [p.text for p in paragraphs]
    tables_before = [t._tbl.xml for t in document.tables]
    scope = next(p for p in paragraphs if p.text == '6.1 Alcances' and p.style.name == 'Heading 2')
    limit = next(p for p in paragraphs if p.text == '6.2 Limitaciones' and p.style.name == 'Heading 2')
    chapter7 = next(p for p in paragraphs if p.text == '7 Marco teórico' and p.style.name == 'Heading 1')
    assert scope._p.getnext() is limit._p and limit._p.getnext() is chapter7._p
    model = next(p for p in paragraphs if p.text.startswith('Claridad.'))
    ppr = model._p.pPr
    for label, text in SCOPES:
        insert_before(limit, label, text, ppr)
    for label, text in LIMITS:
        insert_before(chapter7, label, text, ppr)
    inserted = [f'{l} {t}' for l, t in SCOPES + LIMITS]
    assert [p.text for p in document.paragraphs if p.text not in inserted] == original
    document.tables[0].rows[-1].cells[0].add_paragraph(MATRIX_NOTE)
    assert document.tables[1]._tbl.xml == tables_before[1]
    assert document.tables[2]._tbl.xml == tables_before[2]
    blips = document.paragraphs[0]._p.xpath('.//a:blip')
    assert len(blips) == 1
    rel = blips[0].get('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed')
    document.part.related_parts[rel]._blob = make_cover()
    update = OxmlElement('w:updateFields')
    update.set('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}val', 'true')
    document.settings.element.append(update)
    document.save(TARGET)
    assert hashlib.sha256(SOURCE.read_bytes()).hexdigest() == source_hash
    words = len(' '.join(inserted).split())
    return document, source_hash, words


def build_report(document):
    helpers = runpy.run_path(str(HERE.parent / 'actividad-12-hipotesis/sincronizar_reporte_t12.py'))
    out, paragraphs, tables = [], [], []
    active = False
    in_annex = None
    for element in document.element.body:
        if element.tag.endswith('}p'):
            p = Paragraph(element, document)
            if p.text == '1 Antecedentes' and p.style.name == 'Heading 1':
                active = True
            if not (active and p.text.strip()):
                continue
            paragraphs.append(p.text)
            if p.style.name.startswith('Heading') and p.text.startswith('Anexo '):
                if in_annex:
                    out.append('\\end{landscape}\\clearpage\\pagestyle{fancy}')
                in_annex = p.text
                out.append('\\clearpage\\pagestyle{empty}\\begin{landscape}')
                out.append(helpers['render_paragraph'](p).replace('\\clearpage', ''))
                if p.text.startswith('Anexo 2'):
                    out.append('\\noindent\\textbf{Figura 1}\\par\n\\textit{Círculo de Covey aplicado a la articulación del anteproyecto de investigación}\\par\\vspace{2mm}\n\\circulodecovey[0.64]\n\\par\\vspace{2mm}\n\\noindent{\\footnotesize\\textit{Nota.} Adaptado del esquema didáctico institucional del Instituto Tecnológico Superior de Cajeme (s. f., p. 6). Problema y objetivo general permanecen en el cuadrante Importante/No urgente; los alcances de T13 no agregan elementos al esquema.}')
                continue
            if in_annex and in_annex.startswith('Anexo 2'):
                continue  # el Covey se representa con la figura TikZ
            if p.text.startswith('6 Alcances') and p.style.name == 'Heading 1':
                out.append('\\section{Alcances y limitaciones}')
                continue
            out.append(helpers['render_paragraph'](p))
        elif element.tag.endswith('}tbl') and active:
            t = Table(element, document)
            tables.append([[c.text for c in r.cells] for r in t.rows])
            if in_annex and in_annex.startswith('Anexo 2'):
                continue
            out.append(helpers['render_table'](t))
    if in_annex:
        out.append('\\end{landscape}\\clearpage\\pagestyle{fancy}')
    tex = [
        '\\def\\seminariotitulo{' + TITULO + '}',
        '\\def\\seminariosubtitulo{Actividad 13 -- Alcances y limitaciones}',
        '\\def\\seminarioasunto{Actividad 13 -- Seminario I -- MGA}',
        '\\def\\seminariofecha{' + FECHA + '}',
        '\\def\\seminarioacumulativo{}',
        '\\def\\seminarioextras{\\usepackage{longtable,booktabs,pdflscape}\\input{ITESCA/maestria-en-gestion-administrativa/seminario-i-mga/circulo-covey-tikz.tex}}',
        '\\long\\def\\seminariocontenido{',
        '\n'.join(out),
        '}',
        '\\input{ITESCA/maestria-en-gestion-administrativa/seminario-i-mga/reporte-seminario-i-plantilla-actividad.tex}',
    ]
    REPORT.write_text('\n'.join(tex) + '\n', encoding='utf-8')
    return paragraphs, tables


def main():
    backup = HERE / 'antes-sincronizacion-2026-10-25'
    backup.mkdir(exist_ok=True)
    for f in (REPORT, REPORT.with_suffix('.pdf')):
        if f.exists() and not (backup / f.name).exists():
            (backup / f.name).write_bytes(f.read_bytes())
    document, source_hash, words = build_word()
    assert 300 <= words <= 500, words
    paragraphs, tables = build_report(document)
    record = {'source': SOURCE.name, 'source_sha256': source_hash, 'target': TARGET.name,
              'target_sha256_before_word_refresh': hashlib.sha256(TARGET.read_bytes()).hexdigest(),
              'chapter6_words': words, 'prior_paragraphs_preserved': True,
              'specific_matrix_and_covey_unchanged': True, 'matrix_note_added': True,
              'report': REPORT.name, 'body_paragraphs': len(paragraphs), 'tables': len(tables),
              'submission_performed': False}
    (HERE / 'generacion-t13-2026-10-25.json').write_text(json.dumps(record, ensure_ascii=False, indent=2), encoding='utf-8')
    print(json.dumps({k: v for k, v in record.items()}, ensure_ascii=True))


if __name__ == '__main__':
    main()

