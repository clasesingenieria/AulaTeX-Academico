import hashlib
import json
import re
import runpy
from io import BytesIO
from pathlib import Path

import pymupdf
from docx import Document
from docx.table import Table
from docx.text.paragraph import Paragraph
from docx.oxml import OxmlElement


HERE = Path(__file__).resolve().parent
SUBJECT = HERE.parents[2]
SOURCE = SUBJECT / 'Entregas/Tarea12_DeLaCruzMunoz.docx'
TARGET = SUBJECT / 'Entregas/Tarea13_DeLaCruzMunoz.docx'
REPORT = SUBJECT / 'reporte-seminario-i-Actividad-13.tex'
SCOPES = [
    'Alcance temático. El estudio examina la gestión administrativa del cumplimiento en vTaxi: procedimientos, responsables, verificación de aplicabilidad y seguimiento documental frente a integridad, concordancia y cierre respaldado de observaciones. La normativa sirve para definir criterios del expediente, no para emitir un dictamen jurídico. No se estudian demanda, rentabilidad, desempeño del algoritmo, seguridad vial ni efectos sobre conductores.',
    'Alcance espacial y temporal. La delimitación es Monterrey, Nuevo León, para el proyecto de despacho de taxis concesionados; los requisitos estatales conservan su ámbito sin extender los resultados a toda la entidad o a SETIAP/ERT. El periodo septiembre–noviembre de 2026 comprende la documentación de septiembre como base, revisión en octubre y contraste y priorización durante noviembre. Son actividades previstas, no ejecutadas; cada fuente conservará fecha y versión.',
    'Alcance metodológico y resultados esperados. Se realizará un estudio de caso documental, descriptivo y exploratorio por patrones: vTaxi es la unidad de análisis; componentes y observaciones son registros internos. Se cotejarán fuentes oficiales y documentos autorizados mediante criterios previos, distinguiendo controles X de calidad documental Y. Se esperan catálogo de requisitos aplicables, diagnóstico de brechas, matriz de correspondencias y ruta priorizada, vinculados a los cuatro objetivos específicos. No se demostrarán causalidad ni generalización estadística. Las variables y técnicas deben ser observables y disponibles (Hernández Sampieri et al., 2014, pp. 106–107).',
]
LIMITS = [
    'Limitación teórica. Los antecedentes sobre MiPyMEs explican capacidades organizacionales indirectamente; la revisión específica de cumplimiento y movilidad aún es limitada y Parker y Nielsen se han usado con alcance de resumen editorial. Esto restringe la fuerza explicativa del marco. Se ampliará la revisión pertinente y se distinguirán orientaciones conceptuales de evidencia local, sin presentar analogías como validación de HG.',
    'Limitación metodológica y de acceso. El corpus autorizado puede ser incompleto y no contener variación suficiente de controles o cierres. Las versiones del mismo documento no son casos independientes. Si faltan evidencias o denominadores, se informará no verificable o no calculable; si no es posible contrastar, no se confirmará la hipótesis. Entrevistas y nuevos expedientes quedan condicionados a consentimiento y autorización, no son recursos asegurados.',
    'Limitación técnica y temporal. Borradores con datos de ejemplo, fuentes inaccesibles y cambios normativos pueden impedir comprobar integridad o aplicabilidad en el corte. Se registrarán procedencia, vigencia y discrepancias; una segunda revisión será condicionada a disponibilidad. El cierre de noviembre limita comprobaciones adicionales: se priorizarán brechas documentables mediante Covey, sin equiparar falta de acceso con incumplimiento. Los resultados no acreditarán operación, concesiones o autorización gubernamental.',
]
MATRIX_NOTE = ('Alcance T13: caso documental de vTaxi delimitado a Monterrey, septiembre–noviembre de 2026; requisitos estatales según aplicabilidad. OE1: catálogo; OE2: diagnóstico; OE3: contraste por patrones entre controles y expediente; OE4: ruta priorizada. Corpus autorizado, faltantes y ausencia de variación explícitos; sin causalidad ni generalización estadística. Covey conserva problema y objetivo general en Importante/No urgente.')


def insert_before(anchor, text):
    element = OxmlElement('w:p')
    anchor._p.addprevious(element)
    paragraph = Paragraph(element, anchor._parent)
    paragraph.style = 'Normal'
    paragraph.add_run(text)


def make_cover():
    image = SUBJECT / 'assets-seminario-i/portada/portada-oficial.png'
    output = BytesIO()
    with pymupdf.open() as pdf:
        page = pdf.new_page(width=612, height=792)
        page.insert_image(page.rect, filename=str(image))
        page.insert_font(fontname='Arial', fontfile='C:/Windows/Fonts/arial.ttf')
        page.insert_font(fontname='ArialBold', fontfile='C:/Windows/Fonts/arialbd.ttf')
        block = pymupdf.Rect(56.7, 294.8, 544.3, 454)
        page.draw_rect(block, fill=(1, 1, 1), color=None, fill_opacity=.96)
        metadata = 'Maestría en Gestión Administrativa\nAsignatura: Seminario I\nCurso: GGS102 - Unidad 3\nEstudiante: Martín Jonathan de la Cruz Muñoz\nMatrícula: 26130503\nDocente: Dra. Carla Olimpya Zapuche Moreno\nFecha: 14 de octubre de 2026\nLugar: Ciudad Obregón, Sonora'
        assert page.insert_textbox(block + (12, 12, -12, -12), metadata, fontname='Arial', fontsize=11, lineheight=1.25) >= 0
        title = 'Gestión del cumplimiento normativo para formalizar vTaxi como plataforma de taxis en Nuevo León\n\nActividad 13 - Alcances y limitaciones\nAnteproyecto acumulativo'
        assert page.insert_textbox(pymupdf.Rect(263.6, 552.7, 582, 704), title, fontname='ArialBold', fontsize=12, lineheight=1.2, color=(1, 1, 1)) >= 0
        output.write(page.get_pixmap(matrix=pymupdf.Matrix(2, 2)).tobytes('png'))
    return output.getvalue()


def main():
    if TARGET.exists() or REPORT.exists():
        raise FileExistsError('T13 ya existe; no se sobrescribe')
    words = len(' '.join(SCOPES + LIMITS).split())
    assert 300 <= words <= 500, words
    original_hash = hashlib.sha256(SOURCE.read_bytes()).hexdigest()
    document = Document(SOURCE)
    original = [p.text for p in document.paragraphs]
    tables_before = [table._tbl.xml for table in document.tables]
    limit = next(p for p in document.paragraphs if p.text == '6.2 Limitaciones' and p.style.name == 'Heading 2')
    next_chapter = next(p for p in document.paragraphs if p.text == '7 Marco teórico' and p.style.name == 'Heading 1')
    scope = next(p for p in document.paragraphs if p.text == '6.1 Alcances' and p.style.name == 'Heading 2')
    assert scope._p.getnext() is limit._p
    assert limit._p.getnext() is next_chapter._p
    for text in SCOPES:
        insert_before(limit, text)
    for text in LIMITS:
        insert_before(next_chapter, text)
    assert [p.text for p in document.paragraphs if p.text not in SCOPES + LIMITS] == original
    document.tables[0].rows[-1].cells[0].add_paragraph(MATRIX_NOTE)
    assert document.tables[1]._tbl.xml == tables_before[1]
    assert document.tables[2]._tbl.xml == tables_before[2]
    embedded = document.paragraphs[0]._p.xpath('.//a:blip')
    assert len(embedded) == 1
    relationship = embedded[0].get('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed')
    document.part.related_parts[relationship]._blob = make_cover()
    document.save(TARGET)
    helpers = runpy.run_path(str(HERE.parent / 'actividad-12-hipotesis/sincronizar_reporte_t12.py'))
    rendered = []
    paragraphs = []
    tables = []
    active = False
    for element in document.element.body:
        if element.tag.endswith('}p'):
            paragraph = Paragraph(element, document)
            if paragraph.text == '1 Antecedentes' and paragraph.style.name == 'Heading 1':
                active = True
            if active and paragraph.text.strip():
                rendered.append(helpers['render_paragraph'](paragraph))
                paragraphs.append(paragraph.text)
        elif element.tag.endswith('}tbl') and active:
            table = Table(element, document)
            rendered.append(helpers['render_table'](table))
            tables.append([[cell.text for cell in row.cells] for row in table.rows])
    output = [
        r'\def\seminariotitulo{Gestión del cumplimiento normativo para formalizar vTaxi como plataforma de taxis en Nuevo León}',
        r'\def\seminariosubtitulo{Actividad 13 -- Anteproyecto acumulativo}',
        r'\def\seminarioasunto{Actividad 13 -- Seminario I -- MGA}',
        r'\def\seminariofecha{14 de octubre de 2026}',
        r'\def\seminarioacumulativo{}',
        r'\def\seminarioextras{\usepackage{longtable,booktabs}}',
        r'\long\def\seminariocontenido{', '\n'.join(rendered), '}',
        r'\input{ITESCA/maestria-en-gestion-administrativa/seminario-i-mga/reporte-seminario-i-plantilla-actividad.tex}',
    ]
    REPORT.write_text('\n'.join(output) + '\n', encoding='utf-8')
    assert hashlib.sha256(SOURCE.read_bytes()).hexdigest() == original_hash
    record = {'source_sha256': original_hash, 'word_sha256_before_refresh': hashlib.sha256(TARGET.read_bytes()).hexdigest(), 'chapter6_words': words, 'prior_paragraphs_preserved': True, 'specific_matrix_and_covey_unchanged': True, 'body_paragraphs': paragraphs, 'tables': tables, 'submission_performed': False}
    (HERE / 'generacion-t13.json').write_text(json.dumps(record, ensure_ascii=False, indent=2), encoding='utf-8')
    print(json.dumps({'chapter6_words': words, 'prior_paragraphs_preserved': True, 'tables': len(tables), 'word': TARGET.name, 'report': REPORT.name}, ensure_ascii=True))


if __name__ == '__main__':
    main()