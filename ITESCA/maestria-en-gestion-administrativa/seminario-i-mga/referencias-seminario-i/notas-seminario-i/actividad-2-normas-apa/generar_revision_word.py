import hashlib
import json
import re
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

HERE = Path(__file__).resolve().parent
SUBJECT = HERE.parents[2]


def plain(text):
    text = re.sub(r"\\url\{([^}]+)\}", r"\1", text)
    text = re.sub(r"\\textquote\{([^}]+)\}", r"“\1”", text)
    text = text.replace(r"\textsuperscript{a}", "a")
    text = re.sub(r"\\emph\{([^}]+)\}", r"\1", text)
    text = text.replace(r"\&", "&").replace("``", "“").replace("''", "”")
    return " ".join(text.split())


def main():
    source = SUBJECT / "reporte-seminario-i-Actividad-2.tex"
    text = source.read_text(encoding="utf-8")
    body = text.split(r"\section{Evidencia y decisiones en la gestión administrativa}", 1)[1].split(r"\section{Conclusiones}", 1)[0]
    body = re.sub(r"\\footnote\{[^}]+\}", "", body)
    document = Document()
    section = document.sections[0]
    section.page_width, section.page_height = Inches(8.5), Inches(11)
    section.top_margin = section.bottom_margin = section.left_margin = section.right_margin = Inches(1)
    style = document.styles["Normal"]
    style.font.name, style.font.size = "Arial", Pt(11)
    style.paragraph_format.line_spacing = 2
    style.paragraph_format.first_line_indent = Inches(0.5)
    style.paragraph_format.space_after = Pt(0)
    for level in (1, 2):
        heading = document.styles[f"Heading {level}"]
        heading.font.name, heading.font.size = "Arial", Pt(11)
        heading.font.bold = True
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    section.header.paragraphs[0].alignment = 2
    section.header.paragraphs[0]._p.append(field)
    for line in ["Instituto Tecnológico Superior de Cajeme", "Maestría en Gestión Administrativa", "Mi primera práctica en formato APA", "Martín Jonathan de la Cruz Muñoz", "Número de control: 26130503", "Dra. Carla Olimpya Zapuche Moreno", "Seminario I", "Revisión: 10 de octubre de 2026"]:
        paragraph = document.add_paragraph(line)
        paragraph.alignment = 1
        paragraph.paragraph_format.first_line_indent = Inches(0)
    document.add_page_break()
    document.add_heading("Mi primera práctica en formato APA", 1)
    body = re.sub(r"\\(?:sub)?section\{[^}]+\}", "\n\n", body)
    for fragment in body.split("\n\n"):
        if not fragment.strip():
            continue
        block = r"\begin{apablockquote}" in fragment
        fragment = fragment.replace(r"\begin{apablockquote}", "").replace(r"\end{apablockquote}", "")
        rendered = plain(fragment)
        if rendered.startswith("El pasaje supera 40 palabras"):
            rendered = "El bloque anterior requiere interpretar las funciones directivas en su contexto. Zambrano Andrade (2018) estudió la toma de decisiones administrativas y Sinek (2009) expuso el papel del propósito en el liderazgo: son fuentes de distinta naturaleza. La American Psychological Association (2022) distingue citas breves y bloques de 40 palabras o más; ambas modalidades conservan atribución y localizador."
        paragraph = document.add_paragraph(rendered)
        if block:
            paragraph.paragraph_format.left_indent = Inches(0.5)
            paragraph.paragraph_format.first_line_indent = Inches(0)
    document.add_heading("Referencias", 1)
    references = text.split(r"\begin{apareferences}", 1)[1].split(r"\end{apareferences}", 1)[0]
    for reference in references.split(r"\item")[1:]:
        paragraph = document.add_paragraph()
        fragments = re.split(r"(\\emph\{[^}]+\})", reference)
        for fragment in fragments:
            value = plain(fragment)
            if fragment[:1].isspace() and value:
                value = " " + value
            if fragment[-1:].isspace() and value:
                value += " "
            run = paragraph.add_run(value)
            run.italic = fragment.startswith(r"\emph{")
        paragraph.paragraph_format.left_indent = Inches(0.5)
        paragraph.paragraph_format.first_line_indent = Inches(-0.5)
    target = SUBJECT / "Entregas/Tarea2_DeLaCruzMunoz_Revision-2026-10-10.docx"
    if target.exists():
        raise FileExistsError(target)
    document.save(target)
    result = {"source_sha256": hashlib.sha256(source.read_bytes()).hexdigest(), "archivo": target.relative_to(SUBJECT).as_posix(), "font": "Arial 11", "line_spacing": 2, "margins_inches": 1, "references": 6, "submission": False, "pagination_pending": True}
    (HERE / "revision-word.json").write_text(json.dumps(result, indent=2), encoding="utf-8")
    print(json.dumps(result))


if __name__ == "__main__":
    main()