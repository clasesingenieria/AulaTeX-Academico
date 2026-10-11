import hashlib
import json
import re
from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.text.paragraph import Paragraph

HERE = Path(__file__).resolve().parent
SUBJECT = HERE.parents[2]
SOURCE = SUBJECT / "Entregas/Tarea10_DeLaCruzMunoz_Revision-2026-10-04.docx"
TARGET = SUBJECT / "Entregas/Tarea11_DeLaCruzMunoz.docx"


def insert_before(anchor, text, style):
    element = OxmlElement("w:p")
    anchor._p.addprevious(element)
    paragraph = Paragraph(element, anchor._parent)
    paragraph.style = style
    paragraph.add_run(text)
    return paragraph


def tex_escape(text):
    replacements = {"\\": r"\textbackslash{}", "&": r"\&", "%": r"\%", "$": r"\$", "#": r"\#", "_": r"\_", "{": r"\{", "}": r"\}", "~": r"\textasciitilde{}", "^": r"\textasciicircum{}"}
    return "".join(replacements.get(character, character) for character in text)


def main():
    data = json.loads((HERE / "justificacion.json").read_text(encoding="utf-8"))
    original_hash = hashlib.sha256(SOURCE.read_bytes()).hexdigest()
    document = Document(SOURCE)
    original_tables = [table._tbl.xml for table in document.tables]
    original_paragraphs = [paragraph.text for paragraph in document.paragraphs]
    anchor = next(paragraph for paragraph in document.paragraphs if paragraph.text == "5 Hipótesis de trabajo" and paragraph.style.name == "Heading 1")
    chapter = next(paragraph for paragraph in document.paragraphs if paragraph.text == "4 Justificación" and paragraph.style.name == "Heading 1")
    assert chapter._p.getnext() is anchor._p, "La base contiene justificacion previa; revisar antes de regenerar"
    inserted = []
    for section in data["sections"]:
        insert_before(anchor, section["heading"], "Heading 2")
        inserted.append(section["heading"])
        for text in section["paragraphs"]:
            insert_before(anchor, text, "Normal")
            inserted.append(text)
    reference_anchor = next(paragraph for paragraph in document.paragraphs if paragraph.text.startswith("Kraus,") and paragraph.style.name == "Normal")
    new_ref = insert_before(reference_anchor, data["new_reference"], "Normal")
    previous = Paragraph(new_ref._p.getprevious(), new_ref._parent)
    if previous._p.pPr is not None:
        existing = new_ref._p.pPr
        if existing is not None:
            new_ref._p.remove(existing)
        new_ref._p.insert(0, deepcopy(previous._p.pPr))
    inserted.append(data["new_reference"])
    assert [paragraph.text for paragraph in document.paragraphs if paragraph.text not in inserted] == original_paragraphs
    assert [table._tbl.xml for table in document.tables] == original_tables
    settings = document.settings.element
    update = OxmlElement("w:updateFields")
    update.set("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}val", "true")
    settings.append(update)
    if TARGET.exists():
        raise FileExistsError("No se sobrescribe una T11 existente")
    document.save(TARGET)
    tex = [r"\def\seminariotitulo{" + tex_escape(data["title"]) + "}", r"\def\seminariosubtitulo{Actividad 11 - Relevancia, originalidad y viabilidad}", r"\def\seminariofecha{10 de octubre de 2026}", r"\def\seminarioresumen{Se argumenta la pertinencia del estudio de vTaxi desde la gestion administrativa, con cinco dimensiones de justificacion y limites documentales explicitos.}", r"\AtBeginDocument{\def\documentsubject{Actividad 11 - Seminario I}\def\authortable{\begin{tabular}{ll}Alumno: & Martin Jonathan de la Cruz Munoz\\Matricula: & 26130503\\Docente: & Dra. Carla Olimpya Zapuche Moreno\\Programa: & Maestria en Gestion Administrativa\\Fecha: & 10 de octubre de 2026\end{tabular}}}", r"\def\seminariocontenido{", r"\section{Introduccion}", "La justificacion vincula el problema de preparacion documental con los objetivos del estudio y la LGAC Gestion e Innovacion de las Organizaciones. Su utilidad propuesta se centra en responsabilidades, procedimientos, controles e indicadores, sin anticipar autorizaciones ni resultados.", r"\section{Justificacion de la preparacion administrativa de vTaxi}"]
    for section in data["sections"]:
        tex.append(r"\subsection{" + tex_escape(section["heading"][4:]) + "}")
        tex.extend(tex_escape(paragraph) + "\n\n" for paragraph in section["paragraphs"])
    tex.extend([r"\section{Conclusiones}", "Considero que la pertinencia del estudio reside en convertir una brecha documental en criterios de revision y decisiones trazables. La fase inicial es viable con las fuentes disponibles; ampliar el acceso exige autorizacion y consentimiento. No confundo una propuesta de gestion con la formalizacion efectiva de la plataforma.", "}", r"\def\seminarioreferencias{\clearpage\section*{Referencias}"])
    references = [paragraph.text for paragraph in document.paragraphs]
    start = references.index("Referencias")
    end = references.index("Anexo 1 Matriz de consistencia")
    for reference in references[start + 1:end]:
        if reference.strip():
            parts = re.split(r"(https?://\S+)", reference)
            rendered = "".join(r"\url{" + part + "}" if part.startswith(("https://", "http://")) else tex_escape(part) for part in parts)
            tex.append(rendered + r"\par\medskip")
    tex.extend(["}", r"\input{ITESCA/maestria-en-gestion-administrativa/seminario-i-mga/assets-seminario-i/plantilla/reporte-seminario-i-plantilla-actividad.tex}"])
    (SUBJECT / "reporte-seminario-i-Actividad-11.tex").write_text("\n".join(tex) + "\n", encoding="utf-8")
    assert hashlib.sha256(SOURCE.read_bytes()).hexdigest() == original_hash
    result = {"source_sha256": original_hash, "target": TARGET.relative_to(SUBJECT).as_posix(), "target_sha256": hashlib.sha256(TARGET.read_bytes()).hexdigest(), "previous_paragraphs_retained": True, "annex_tables_unchanged": True, "sections": [section["heading"] for section in data["sections"]], "toc_requires_refresh": True, "submission_performed": False}
    (HERE / "generacion-t11.json").write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(result, ensure_ascii=True))


if __name__ == "__main__":
    main()