import hashlib
import json
from pathlib import Path
import re
import sys
import unicodedata

import pymupdf
from docx import Document

ROOT = Path(__file__).resolve().parents[8]
sys.path.insert(0, str(ROOT))
from scripts.aulatex.seminario_objectives_renderer import validate_structure, word_text

SUBJECT = Path(__file__).resolve().parents[5]
OUTPUT = Path(__file__).resolve().parent


def normalized(text):
    text = unicodedata.normalize("NFKC", text).replace("\xad", "")
    text = re.sub(r"(?<=\w)-\s*\n(?=\w)", "", text)
    return " ".join(text.split())


def main():
    payload = SUBJECT / "referencias-seminario-i/notas-seminario-i/actividad-10-objetivos/elaboracion-vtaxi-2026-09-28/actividad-10-revisada.json"
    data = json.loads(payload.read_text(encoding="utf-8"))
    word = SUBJECT / "entregas/Tarea10_DeLaCruzMunoz_Revision-2026-10-04.docx"
    source_word = SUBJECT / "referencias-seminario-i/notas-seminario-i/actividad-10-objetivos/elaboracion-vtaxi-2026-09-28/base-T9-corregida-para-T10.docx"
    original_word = SUBJECT / "entregas/Tarea9_DeLaCruzMunoz.docx"
    report = SUBJECT / "reporte-seminario-i-Actividad-10.tex"
    slides = SUBJECT / "presentacion-seminario-i-Actividad-10.tex"
    continuity = validate_structure(source_word, word, data, after_word_refresh=True)
    checks = dict(continuity["checks"])
    additions = [word_text(text, data) for text in data["antecedent_additions"]] + [data["problem_addition"]]
    checks["only_declared_feedback_additions"] = [paragraph.text for paragraph in Document(source_word).paragraphs if paragraph.text not in additions] == [paragraph.text for paragraph in Document(original_word).paragraphs]
    checks["original_t9_unchanged"] = hashlib.sha256(original_word.read_bytes()).hexdigest() == "51a2c0405ef3d9261ee9ab81bcc6ea8655f58b0bc9eb5b34ef5d093b61bdf999"
    checks["five_support_sources"] = len(data["references"]) == 5
    checks["report_three_acts"] = len(re.findall(r"\\section\{", report.read_text(encoding="utf-8"))) == 3
    objectives = [data["content"]["general_objective"], *data["content"]["specific_objectives"]]
    checks["four_specific_objectives"] = len(objectives) == 5
    pdf_results = []
    for source, label in [(word, "word"), (report, "reporte"), (slides, "presentacion")]:
        pdf = source.with_suffix(".pdf")
        checks[label + "_pdf_fresh"] = pdf.exists() and pdf.stat().st_mtime >= source.stat().st_mtime
        with pymupdf.open(pdf) as document:
            pages = [page.get_text() for page in document]
            text = "\n".join(pages)
            checks[label + "_objectives_visible"] = all(normalized(value) in normalized(text) for value in objectives)
            if label == "word":
                checks["word_objectives_not_split"] = all(any(normalized(value) in normalized(page) for page in pages) for value in objectives)
            checks[label + "_student_id"] = "26130503" in pages[0]
            checks[label + "_no_foreign_identity"] = not re.search(r"UnADM|ES2611202040|Nombre por definir|\[PENDIENTE", text)
            checks[label + "_no_operational_blocks"] = not re.search(r"Actividad local|Módulo Moodle|Vencimiento publicado", text, re.I)
            checks[label + "_no_missing_citations"] = "[?]" not in text and "[@" not in text
            checks[label + "_scientific_support_visible"] = "Parker" in text and "Thelen" in text
            checks[label + "_administrative_indicators_visible"] = "cobertura de responsabilidades" in normalized(text).lower() and "cierre de observaciones" in normalized(text).lower()
            pdf_results.append({"file": pdf.relative_to(ROOT).as_posix(), "pages": len(document)})
            folder = OUTPUT / ("render-" + label)
            folder.mkdir(exist_ok=True)
            for number, page in enumerate(document, 1):
                page.get_pixmap(matrix=pymupdf.Matrix(1, 1)).save(folder / f"pagina-{number:02}.png")
    for source, log in [(report, OUTPUT / "compilacion" / report.with_suffix(".log").name), (slides, OUTPUT / "compilacion" / slides.with_suffix(".log").name)]:
        text = log.read_text(encoding="utf-8", errors="replace")
        checks[source.stem + "_log_clean"] = not re.search(r"Overfull|undefined citations|Citation .* undefined|^!", text, re.M)
    extraction = SUBJECT / "extractor-aulatex/conceptos-seminario-i-actividad-10"
    extraction.mkdir(exist_ok=True)
    with pymupdf.open(SUBJECT / "referencias-seminario-i/notas-seminario-i/actividad-10-objetivos/materiales/3.3 - Formulacion de objetivos.pdf") as material:
        extracted = "\n\n".join(f"PAGINA {number}\n{page.get_text()}" for number, page in enumerate(material, 1))
    (extraction / "material-3.3-extraido.txt").write_text(extracted, encoding="utf-8")
    artifacts = [payload, original_word, source_word, word, word.with_suffix(".pdf"), report, report.with_suffix(".pdf"), report.with_suffix(".bib"), slides, slides.with_suffix(".pdf")]
    result = {
        "date": "2026-10-04",
        "passed": all(checks.values()),
        "checks": checks,
        "pdfs": pdf_results,
        "files": [{"path": path.relative_to(ROOT).as_posix(), "sha256": hashlib.sha256(path.read_bytes()).hexdigest()} for path in artifacts],
        "submission_performed": False,
        "scope": "Validacion documental automatizada; no certifica ciclo multiagente, EMS, revision docente ni entrega. Revision visual registrada aparte.",
    }
    (OUTPUT / "validacion-final.json").write_text(json.dumps(result, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps({"passed": result["passed"], "failed": [key for key, value in checks.items() if not value], "pdfs": pdf_results}, ensure_ascii=True, indent=2))
    if not result["passed"]:
        raise SystemExit(1)


if __name__ == "__main__":
    main()