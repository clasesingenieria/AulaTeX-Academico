from copy import deepcopy
import hashlib
import json
from pathlib import Path
import sys

from docx import Document

ROOT = Path(__file__).resolve().parents[8]
sys.path.insert(0, str(ROOT))
from scripts.aulatex.seminario_objectives_renderer import (
    _bibliography, _normal, build_objectives_beamer, build_objectives_docx,
    export_word_pdf, tex_escape, tex_text, validate_structure, word_text,
)

SUBJECT = Path(__file__).resolve().parents[5]
OUTPUT = Path(__file__).resolve().parent


def main():
    data = json.loads((SUBJECT / "referencias-seminario-i/notas-seminario-i/actividad-10-objetivos/elaboracion-vtaxi-2026-09-28/actividad-10-revisada.json").read_text(encoding="utf-8"))
    previous = OUTPUT / "antes-retroalimentacion"
    old_data = json.loads((previous / "actividad-10-revisada.json").read_text(encoding="utf-8"))
    original = SUBJECT / "entregas/Tarea9_DeLaCruzMunoz.docx"
    original_hash = hashlib.sha256(original.read_bytes()).hexdigest()
    working = SUBJECT / "referencias-seminario-i/notas-seminario-i/actividad-10-objetivos/elaboracion-vtaxi-2026-09-28/base-T9-corregida-para-T10.docx"
    document = Document(original)
    anchor = next(paragraph for paragraph in document.paragraphs if paragraph.text == "Vacíos identificados")
    for text in data["antecedent_additions"]:
        _normal(anchor.insert_paragraph_before(word_text(text, data), style="Normal"))
    anchor = next(paragraph for paragraph in document.paragraphs if paragraph.text == "2.2 Formulación del problema")
    _normal(anchor.insert_paragraph_before(data["problem_addition"], style="Normal"))
    document.save(working)
    changed = Document(working)
    additions = [word_text(text, data) for text in data["antecedent_additions"]] + [data["problem_addition"]]
    assert [paragraph.text for paragraph in changed.paragraphs if paragraph.text not in additions] == [paragraph.text for paragraph in Document(original).paragraphs]
    render_data = deepcopy(data)
    render_data["content"]["chapter3_alignment"] += "\n\n" + data["administrative_indicators"]
    word = SUBJECT / "entregas/Tarea10_DeLaCruzMunoz_Revision-2026-10-04.docx"
    receipt = build_objectives_docx(working, render_data, word)
    formatted = Document(word)
    objective_texts = [data["content"]["general_objective"]] + [f"{number}. {text}" for number, text in enumerate(data["content"]["specific_objectives"], 1)]
    for paragraph in formatted.paragraphs:
        if paragraph.text in objective_texts:
            paragraph.paragraph_format.keep_together = True
    formatted.save(word)
    pdf_receipt = export_word_pdf(word)
    post_refresh = validate_structure(working, word, render_data, after_word_refresh=True)
    assert post_refresh["passed"]

    report = SUBJECT / "reporte-seminario-i-Actividad-10.tex"
    text = (previous / report.name).read_text(encoding="utf-8")
    for key in ("specific_objectives", "verification_products", "procedures"):
        for old, new in zip(old_data["content"][key], data["content"][key]):
            assert tex_text(old) in text
            text = text.replace(tex_text(old), tex_text(new))
    anchor = r"\Needspace{8\baselineskip}\subsection{Objetivo general}"
    assert anchor in text
    text = text.replace(anchor, r"\subsection{Sustento administrativo de los objetivos}" + "\n" + tex_text(data["scientific_basis"]) + "\n\n" + anchor)
    anchor = r"\subsection{Matriz de consistencia}"
    text = text.replace(anchor, r"\subsection{Criterios de seguimiento administrativo}" + "\n" + tex_text(data["administrative_indicators"]) + "\n\n" + anchor)
    references = r"""
\bibitem[Parker y Nielsen(2009)]{parkerNielsen2009}
Parker, C., \& Nielsen, V. L. (2009). Corporate compliance systems: Could they make any difference? \textit{Administration \& Society, 41}(1), 3--37. \url{https://doi.org/10.1177/0095399708328869}
\bibitem[Thelen(2018)]{thelen2018}
Thelen, K. (2018). Regulating Uber: The politics of the platform economy in Europe and the United States. \textit{Perspectives on Politics, 16}(4), 938--953. \url{https://doi.org/10.1017/S1537592718001081}
"""
    text = text.replace(r"\begin{thebibliography}{3}", r"\begin{thebibliography}{5}")
    text = text.replace(r"\end{thebibliography}", references + r"\end{thebibliography}")
    report.write_text(text, encoding="utf-8")
    report.with_suffix(".bib").write_text(_bibliography(data), encoding="utf-8")

    slides = SUBJECT / "presentacion-seminario-i-Actividad-10.tex"
    build_objectives_beamer(data, slides)
    body = slides.read_text(encoding="utf-8").split(r"\begin{document}", 1)[1]
    header = (previous / slides.name).read_text(encoding="utf-8").split(r"\begin{document}", 1)[0]
    frames = "\n".join(
        r"\begin{frame}{" + title + r"}\small " + tex_escape(word_text(data[key], data)) + r"\end{frame}"
        for title, key in [("Enfoque administrativo y antecedentes", "scientific_basis"), ("Indicadores propuestos de gestión", "administrative_indicators")]
    )
    body = body.replace(r"\begin{frame}{Círculo de Covey}", frames + "\n" + r"\begin{frame}{Círculo de Covey}")
    slides.write_text(header + r"\begin{document}" + body, encoding="utf-8")
    assert hashlib.sha256(original.read_bytes()).hexdigest() == original_hash
    (OUTPUT / "generacion-retroalimentacion.json").write_text(json.dumps({
        "original_t9_sha256": original_hash,
        "working_baseline": working.relative_to(ROOT).as_posix(),
        "only_declared_additions_to_t9": True,
        "word": receipt,
        "word_pdf": pdf_receipt,
        "post_refresh_qa": post_refresh,
        "submission_performed": False,
    }, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print("PASS: correcciones docentes incorporadas; T9 original intacta; PDF Word", pdf_receipt["word_pages"], "paginas.")


if __name__ == "__main__":
    main()