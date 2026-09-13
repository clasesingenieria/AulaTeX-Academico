"""Recupera adjuntos DOCX descargados desde la sesión autorizada del aula."""

import base64
import json
import sys
from pathlib import Path
from zipfile import ZipFile
from io import BytesIO

from docx import Document


def main():
    record = Path(sys.argv[1]).read_text()
    payload, _ = json.JSONDecoder().raw_decode(record[record.index("["):])
    destination = Path(__file__).parent / "materiales"
    destination.mkdir(parents=True, exist_ok=True)
    for item in payload:
        assert item["status"] == 200, item["name"]
        data = base64.b64decode(item["base64"])
        with ZipFile(BytesIO(data)) as archive:
            assert "word/document.xml" in archive.namelist()
        path = destination / Path(item["name"]).name
        if path.exists():
            assert path.read_bytes() == data, f"No se sobrescribe un original distinto: {path}"
        else:
            path.write_bytes(data)
        doc = Document(path)
        print(f"\n--- {path.name}: {len(data)} bytes ---")
        for index, paragraph in enumerate(doc.paragraphs):
            print(index, paragraph.style.name, repr(paragraph.text))
        for table in doc.tables:
            for row in table.rows:
                print("TABLE", [cell.text for cell in row.cells])


if __name__ == "__main__":
    main()