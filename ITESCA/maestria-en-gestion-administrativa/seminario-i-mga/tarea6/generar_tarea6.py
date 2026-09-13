"""Genera la práctica APA desde los dos originales oficiales del aula.

No inventa resultados de investigación ni referencias primarias del material.
El índice se pagina posteriormente con actualizar_indice.py y LibreOffice.
"""

from io import BytesIO
from pathlib import Path
import hashlib
import json
import re
from urllib.request import Request, urlopen

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from PIL import Image, ImageDraw, ImageFont

ROOT = Path(__file__).resolve().parent
MATERIALS = ROOT / "materiales"
ASSETS = ROOT / "fuentes"
BUILD = ROOT / "construccion"
NAME = "Tarea6_DeLaCruzMunoz.docx"
SOURCE = "https://openstax.org/books/introduction-sociology-3e/pages/"
APA_URL = "https://cursos3.e-itesca.edu.mx/mod/url/view.php?id=2881"
TASK_URL = "https://cursos3.e-itesca.edu.mx/mod/assign/view.php?id=2889"
LICENSE = "CC BY-NC-SA 4.0 (https://creativecommons.org/licenses/by-nc-sa/4.0/)"


def fetch(name, url):
    path = ASSETS / name
    if not path.exists():
        with urlopen(Request(url, headers={"User-Agent": "Mozilla/5.0"}), timeout=60) as response:
            path.write_bytes(response.read())
    return path


def field(paragraph, instruction, result=""):
    for kind, text in [("begin", None), ("instruction", instruction),
                       ("separate", None), ("text", result), ("end", None)]:
        run = OxmlElement("w:r")
        if kind in ("begin", "separate", "end"):
            element = OxmlElement("w:fldChar")
            element.set(qn("w:fldCharType"), kind)
        elif kind == "instruction":
            element = OxmlElement("w:instrText")
            element.set(qn("xml:space"), "preserve")
            element.text = text
        else:
            element = OxmlElement("w:t")
            element.text = text
        run.append(element)
        paragraph._p.append(run)


def clean_font(run, bold=None, italic=None):
    run.font.name = "Arial"
    run.font.size = Pt(11)
    run.font.color.rgb = RGBColor(0, 0, 0)
    run.font.highlight_color = None
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    for element in run._r.xpath("./w:rPr/w:shd"):
        element.getparent().remove(element)


def configure_styles(doc):
    for style in doc.styles:
        if style.type in (WD_STYLE_TYPE.PARAGRAPH, WD_STYLE_TYPE.CHARACTER):
            style.font.name = "Arial"
            style.font.size = Pt(11)
            style.font.color.rgb = RGBColor(0, 0, 0)
            style.font.highlight_color = None
    normal = doc.styles["Normal"]
    normal.font.bold = False
    normal.font.italic = False
    fmt = normal.paragraph_format
    fmt.alignment = WD_ALIGN_PARAGRAPH.LEFT
    fmt.first_line_indent = Inches(0.5)
    fmt.line_spacing = 2
    fmt.space_before = fmt.space_after = Pt(0)
    fmt.widow_control = True
    for level in range(1, 5):
        style = doc.styles[f"Heading {level}"]
        style.base_style = normal
        style.font.bold = True
        style.font.italic = level == 3
        pf = style.paragraph_format
        pf.alignment = WD_ALIGN_PARAGRAPH.CENTER if level == 1 else WD_ALIGN_PARAGRAPH.LEFT
        pf.first_line_indent = Inches(0.5) if level == 4 else Inches(0)
        pf.left_indent = pf.right_indent = Inches(0)
        pf.space_before = pf.space_after = Pt(0)
        pf.line_spacing = 2
        pf.keep_with_next = level != 4
        pf.keep_together = True
        pf.page_break_before = False
    for name in ["APA sin sangría", "APA referencia", "APA tabla", "APA portada", "APA título índice"]:
        style = doc.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
        style.base_style = normal
        style.paragraph_format.first_line_indent = Inches(0)
    doc.styles["APA referencia"].paragraph_format.left_indent = Inches(0.5)
    doc.styles["APA referencia"].paragraph_format.first_line_indent = Inches(-0.5)
    doc.styles["APA tabla"].paragraph_format.line_spacing = 1
    doc.styles["APA tabla"].paragraph_format.space_after = Pt(5)
    doc.styles["APA tabla"].paragraph_format.space_before = Pt(5)
    doc.styles["APA portada"].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.styles["APA título índice"].font.bold = True
    doc.styles["APA título índice"].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    lang = OxmlElement("w:lang")
    lang.set(qn("w:val"), "es-MX")
    normal.element.get_or_add_rPr().append(lang)


def cover(doc):
    # Se parte de Portada.docx; se conserva su orden y su imagen institucional.
    section = doc.sections[0]
    image = next(r.target_part.blob for r in section.header.part.rels.values()
                 if "image" in r.reltype)
    (ASSETS / "franja-portada-oficial.png").write_bytes(image)
    for part in [section.header, section.footer]:
        for child in list(part._element):
            part._element.remove(child)
        part.add_paragraph()
    section.different_first_page_header_footer = False
    replacements = {
        1: "INSTITUTO TECNOLÓGICO SUPERIOR DE CAJEME",
        2: "Estructura de un trabajo de investigación en gestión administrativa",
        3: "Seminario I · Tarea 6",
        4: "TESIS (EJEMPLO ACADÉMICO)",
        5: "PARA OBTENER EL GRADO DE",
        6: "MAESTRÍA EN GESTIÓN ADMINISTRATIVA",
        8: "PRESENTA",
        9: "MARTÍN JONATHAN DE LA CRUZ MUÑOZ",
        11: "DIRECTORA DE TESIS",
        12: "DRA. CARLA OLIMPYA ZAPUCHE MORENO",
        14: "CIUDAD OBREGÓN, SONORA · SEPTIEMBRE DE 2026",
    }
    for i, paragraph in enumerate(doc.paragraphs):
        paragraph.clear()
        ppr = paragraph._p.find(qn("w:pPr"))
        if ppr is not None:
            paragraph._p.remove(ppr)
        paragraph.style = doc.styles["APA portada"]
        paragraph.paragraph_format.keep_with_next = i < 14
        if i == 0:
            paragraph.add_run().add_picture(BytesIO(image), width=Inches(6.5))
            paragraph.paragraph_format.space_after = Pt(20)
        else:
            run = paragraph.add_run(replacements.get(i, ""))
            clean_font(run, bold=i in (1, 2, 4, 6, 9, 12))
        if i in (1, 2, 6, 9, 12):
            paragraph.paragraph_format.space_after = Pt(10)
    header = section.header.paragraphs[0]
    header.style = doc.styles["APA sin sangría"]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    field(header, " PAGE ", "1")


def scientific_figure():
    # Traducción/redibujo de la figura 2.2, no se presenta como figura original.
    canvas = Image.new("RGB", (1500, 880), "white")
    draw = ImageDraw.Draw(canvas)
    font_path = "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"
    font = ImageFont.truetype(font_path, 31)
    bold = ImageFont.truetype("/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf", 34)
    labels = ["Plantear una\npregunta", "Revisar las\nfuentes existentes", "Formular una\nhipótesis",
              "Diseñar y realizar\nel estudio", "Extraer\nconclusiones", "Comunicar\nlos resultados"]
    positions = [(40, 35), (550, 35), (1060, 35), (1060, 490), (550, 490), (40, 490)]
    for i, ((x, y), label) in enumerate(zip(positions, labels), 1):
        draw.rounded_rectangle((x, y, x + 400, y + 280), radius=20,
                               fill="#f4f4f4", outline="black", width=3)
        draw.text((x + 200, y + 48), str(i), font=bold, fill="black", anchor="mm")
        draw.multiline_text((x + 200, y + 160), label, font=font,
                            fill="black", anchor="mm", align="center", spacing=14)
    for x, y in [(445, 175), (955, 175)]:
        draw.line((x, y, x + 90, y), fill="black", width=5)
        draw.polygon([(x + 90, y), (x + 70, y - 12), (x + 70, y + 12)], fill="black")
    draw.line((1260, 320, 1260, 465), fill="black", width=5)
    draw.polygon([(1260, 465), (1248, 445), (1272, 445)], fill="black")
    for x, y in [(1045, 630), (535, 630)]:
        draw.line((x, y, x - 90, y), fill="black", width=5)
        draw.polygon([(x - 90, y), (x - 70, y - 12), (x - 70, y + 12)], fill="black")
    output = ASSETS / "figura1-metodo-cientifico-es.png"
    canvas.crop((0, 0, 1500, 790)).save(output, dpi=(300, 300))
    return output


def paragraph(doc, text, style=None):
    p = doc.add_paragraph(style=style)
    pattern = r"(Material para ejemplo|Resumen del Manual APA \(7\.ª edición\))"
    for part in re.split(pattern, text):
        run = p.add_run(part)
        if re.fullmatch(pattern, part):
            run.italic = True
    return p


def caption(doc, kind, number, title):
    p = paragraph(doc, "", "APA sin sangría")
    p.add_run(f"{kind} {number}").bold = True
    p.paragraph_format.keep_with_next = True
    p = paragraph(doc, "", "APA sin sangría")
    p.add_run(title).italic = True
    p.paragraph_format.keep_with_next = True


def note(doc, text):
    p = paragraph(doc, "", "APA sin sangría")
    p.paragraph_format.keep_together = True
    p.add_run("Nota. ").italic = True
    for part in re.split(r"(Introduction to Sociology 3e)", text):
        run = p.add_run(part)
        if part == "Introduction to Sociology 3e":
            run.italic = True


def table(doc, number, title, headers, rows, widths, source):
    caption(doc, "Tabla", number, title)
    t = doc.add_table(rows=1, cols=len(headers))
    t.autofit = False
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for col, width in zip(t.columns, widths):
        col.width = Inches(width)
    for cell, text, width in zip(t.rows[0].cells, headers, widths):
        cell.text = text
        cell.width = Inches(width)
    for values in rows:
        for cell, text, width in zip(t.add_row().cells, values, widths):
            cell.text = text
            cell.width = Inches(width)
    borders = OxmlElement("w:tblBorders")
    for edge in ["top", "bottom", "left", "right", "insideH", "insideV"]:
        element = OxmlElement(f"w:{edge}")
        element.set(qn("w:val"), "single" if edge in ("top", "bottom") else "nil")
        element.set(qn("w:sz"), "8")
        element.set(qn("w:color"), "000000")
        borders.append(element)
    t._tbl.tblPr.append(borders)
    for row_index, row in enumerate(t.rows):
        trpr = row._tr.get_or_add_trPr()
        trpr.append(OxmlElement("w:cantSplit"))
        if row_index == 0:
            trpr.append(OxmlElement("w:tblHeader"))
        for cell in row.cells:
            if row_index == 0:
                cb = OxmlElement("w:tcBorders")
                bottom = OxmlElement("w:bottom")
                bottom.set(qn("w:val"), "single")
                bottom.set(qn("w:sz"), "8")
                cb.append(bottom)
                cell._tc.get_or_add_tcPr().append(cb)
            for p in cell.paragraphs:
                p.style = doc.styles["APA tabla"]
                p.paragraph_format.keep_with_next = row_index < len(t.rows) - 1
                for run in p.runs:
                    clean_font(run, bold=row_index == 0)
    note(doc, source)


def figure(doc, number, title, path, source, description):
    caption(doc, "Figura", number, title)
    p = paragraph(doc, "", "APA sin sangría")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    shape = p.add_run().add_picture(str(path), width=Inches(5.3 if number == 2 else 6.2))
    shape._inline.docPr.set("descr", description)
    p.paragraph_format.keep_with_next = True
    note(doc, source)


def extras(doc, key):
    if key == "3.1":
        paragraph(doc, "La Tabla 1 compara métodos de investigación social y ayuda a seleccionar "
                  "una estrategia congruente con el problema, los objetivos y los recursos disponibles "
                  "(Conerly et al., 2021). Un método de recolección no debe confundirse con el alcance del estudio.")
        table(doc, 1, "Métodos de investigación social: técnicas, ventajas y limitaciones",
              ["Método", "Técnicas o fuentes", "Ventajas", "Limitaciones"], [
                  ["Encuesta", "Cuestionarios y entrevistas", "Permite obtener numerosas respuestas y organizar datos cuantitativos.", "La participación puede ser baja; lo declarado no siempre coincide con la conducta."],
                  ["Trabajo de campo", "Observación participante, etnografía y estudio de caso", "Aporta información detallada en contextos reales.", "Requiere tiempo; los datos cualitativos demandan organización e interpretación."],
                  ["Experimento", "Manipulación deliberada de condiciones", "Permite poner a prueba relaciones de causa y efecto.", "Exige precauciones éticas; saberse observado puede alterar la conducta."],
                  ["Análisis secundario", "Estadísticas oficiales y documentos históricos", "Aprovecha información ya disponible.", "Los datos pueden ser difíciles de obtener o responder a un propósito distinto."],
              ], [1.05, 1.6, 1.8, 2.05],
              "Traducción y adaptación de la tabla 2.2, “Main Sociological Research Methods”, "
              "en Introduction to Sociology 3e, por T. R. Conerly, K. Holmes y A. L. Tamang, 2021, "
              f"OpenStax ({SOURCE}2-section-summary). Copyright 2021 Rice University. "
              f"Adaptación bajo {LICENSE}.")
    elif key == "3.2":
        paragraph(doc, "La Tabla 2 ilustra la distinción entre variables dentro de una hipótesis "
                  "(Conerly et al., 2021). Son ejemplos conceptuales, no resultados comprobados. "
                  "La relación propuesta debe operacionalizarse y contrastarse con un diseño adecuado.")
        table(doc, 2, "Ejemplos de hipótesis y sus variables independientes y dependientes",
              ["Hipótesis ilustrativa", "Variable independiente", "Variable dependiente"], [
                  ["A mayor disponibilidad de vivienda asequible, menor tasa de personas sin hogar.", "Disponibilidad de vivienda asequible", "Tasa de personas sin hogar"],
                  ["A mayor disponibilidad de tutoría en matemáticas, mejores calificaciones.", "Disponibilidad de tutoría", "Calificaciones en matemáticas"],
                  ["A mayor iluminación de la fábrica, mayor productividad.", "Iluminación de la fábrica", "Productividad"],
              ], [2.9, 1.8, 1.8],
              "Selección, traducción y adaptación de tres ejemplos de la tabla 2.1, “Examples of "
              "Dependent and Independent Variables”, en Introduction to Sociology 3e, por T. R. Conerly, "
              f"K. Holmes y A. L. Tamang, 2021, OpenStax ({SOURCE}2-1-approaches-to-sociological-research). "
              f"Copyright 2021 Rice University. Adaptación bajo {LICENSE}.")
    elif key == "3.3":
        paragraph(doc, "La Figura 1 organiza las seis etapas del método científico descritas por "
                  "Conerly et al. (2021). Esta secuencia orienta estudios que contrastan hipótesis; "
                  "no implica que toda investigación cualitativa deba seguir un recorrido lineal.")
        figure(doc, 1, "Etapas del método científico en la investigación social",
               ASSETS / "figura1-metodo-cientifico-es.png",
               "Traducida y redibujada a partir de la figura 2.2, “The Scientific Method”, en "
               "Introduction to Sociology 3e, por T. R. Conerly, K. Holmes y A. L. Tamang, 2021, "
               f"OpenStax ({SOURCE}2-1-approaches-to-sociological-research). Copyright 2021 Rice University. "
               f"Adaptación bajo {LICENSE}.",
               "Seis pasos: plantear una pregunta, revisar fuentes, formular una hipótesis, "
               "diseñar y realizar el estudio, extraer conclusiones y comunicar resultados.")
    elif key == "3.5":
        paragraph(doc, "La Figura 2 muestra un cuestionario en soporte digital, un ejemplo de "
                  "instrumento para recabar información. El medio tecnológico no garantiza por sí "
                  "mismo la validez: las preguntas deben corresponder a las variables y al propósito "
                  "del estudio (Conerly et al., 2021).")
        figure(doc, 2, "Aplicación de un cuestionario mediante una tableta",
               ASSETS / "figura2-cuestionario.png",
               "Reproducida de la figura 2.3 de Introduction to Sociology 3e, por T. R. Conerly, "
               f"K. Holmes y A. L. Tamang, 2021, OpenStax ({SOURCE}2-2-research-methods). "
               "Crédito de la fotografía original: CDC Global/Flickr. Se conserva la atribución "
               "indicada en el libro; reutilización educativa no comercial conforme a su apartado "
               f"“Art attribution” y licencia {LICENSE}.",
               "Persona que sostiene una tableta con un cuestionario visible en pantalla.")


def edited_text(text):
    # Se atribuye lo efectivamente consultado, no fuentes primarias no verificadas.
    text = text.replace("De acuerdo con Eyssautier (2002),", "Según el material didáctico de Seminario I (Material para ejemplo, s. f.),")
    text = text.replace("De acuerdo con Hernández, Fernández y Baptista (2022),", "De acuerdo con el material didáctico de Seminario I (Material para ejemplo, s. f.),")
    text = text.replace("Según Sampieri et al. (2022),", "Según el material didáctico de Seminario I (Material para ejemplo, s. f.),")
    if "De acuerdo con Baena (2017)" in text:
        text = ("El material didáctico de Seminario I señala que el marco legal permite delimitar "
                "las normas que regulan los procesos, las actividades y los sujetos relacionados "
                "con el problema, así como las competencias y responsabilidades institucionales "
                "(Material para ejemplo, s. f.). Su organización parte del ámbito internacional, "
                "continúa con el nacional y concluye con la regulación estatal, institucional o sectorial aplicable.")
    if "2.2.3.3 Hipótesis derivadas del modelo c" == text:
        text = "2.2.3.3 Hipótesis derivadas del modelo conceptual"
    text = text.replace("Antecedentes o Estado del Arte", "Antecedentes o estado del arte")
    text = text.replace("Alcances y Limitaciones", "Alcances y limitaciones")
    text = text.replace("Marco Referencial", "Marco referencial")
    text = text.replace("(ONU, 1948)", "(Organización de las Naciones Unidas [ONU], 1948)")
    if "instrumentos ratificados por México" in text:
        text = text.replace("instrumentos ratificados por México", "instrumentos internacionales relacionados con México")
        text += " Las declaraciones y agendas no deben confundirse con tratados sujetos a ratificación."
    if "tales como el Programa Sectorial de" in text:
        text = text.replace("políticas públicas en vigor", "políticas públicas")
        text = text.replace("programas sectoriales, estrategias nacionales o políticas públicas en vigor",
                            "programas sectoriales, estrategias nacionales o políticas públicas")
        text += " Los periodos 2019–2024 y 2020–2024 se conservan como ejemplos históricos del material base; para una investigación actual deben verificarse los instrumentos vigentes."
    if "Ley de Ciencia y Tecnología," in text:
        text += " Las denominaciones legales se conservan como ejemplos del material; su vigencia debe corroborarse en las fuentes jurídicas oficiales antes de aplicarlas a una tesis."
    return text


def reference(doc, before, title, after):
    p = paragraph(doc, "", "APA referencia")
    p.add_run(before)
    p.add_run(title).italic = True
    p.add_run(after)


def references(doc):
    reference(doc, "Conerly, T. R., Holmes, K., & Tamang, A. L. (2021). ",
              "Introduction to sociology 3e", ". OpenStax. https://openstax.org/details/books/introduction-sociology-3e")
    reference(doc, "", "Material para ejemplo",
              f". (s. f.). [Material didáctico de Seminario I]. ITESCA Virtual. {TASK_URL}")
    reference(doc, "Organización de las Naciones Unidas. (1948). ",
              "Declaración Universal de los Derechos Humanos",
              ". https://www.un.org/es/about-us/universal-declaration-of-human-rights")
    reference(doc, "Organización de las Naciones Unidas. (2015). ",
              "Transformar nuestro mundo: La Agenda 2030 para el Desarrollo Sostenible",
              ". https://sdgs.un.org/es/2030agenda")
    reference(doc, "", "Resumen del Manual APA (7.ª edición): Aplicado al Seminario de Investigación y trabajos de titulación",
              f". (s. f.). [Recurso 2.1 de Seminario I]. ITESCA Virtual. {APA_URL}")


def content(doc):
    doc.add_page_break()
    paragraph(doc, "Índice", "APA título índice")
    p = paragraph(doc, "", "APA sin sangría")
    field(p, ' TOC \\o "1-3" \\h \\z \\u ', "Índice pendiente de actualización automática.")
    doc.add_page_break()
    doc.add_heading("Presentación del ejercicio", 1)
    paragraph(doc, "Este documento desarrolla la Tarea 6 de Seminario I mediante la aplicación "
              "de formato APA y un índice automatizado al archivo Material para ejemplo. Se conserva "
              "su función de guía para organizar un trabajo de titulación: las descripciones de "
              "resultados, conclusiones y anexos son orientaciones, no evidencias de una investigación realizada.")
    paragraph(doc, "La portada institucional se utiliza exclusivamente como ejemplo académico. "
              "La designación de la Dra. Carla Olimpya Zapuche Moreno como directora de tesis "
              "responde a la consigna y no constituye un nombramiento formal. La estructura capitular "
              "procede del material del curso (Material para ejemplo, s. f.).")
    paragraph(doc, "El formato general sigue el Resumen del Manual APA (7.ª edición) (s. f.). "
              "Se conserva la numeración institucional de capítulos y apartados; las Tablas 1 y 2 "
              "y las Figuras 1 y 2 se integran en el marco metodológico con sus fuentes y atribuciones.")

    source = Document(MATERIALS / "Material para ejemplo.docx")
    texts = [p.text.strip() for p in source.paragraphs if p.text.strip()]
    current = None
    i = 0
    seen = []
    while i < len(texts):
        text = edited_text(texts[i])
        match = re.match(r"^(\d+(?:\.\d+)+)\s+(.+)", text)
        is_heading = text.startswith("CAPÍTULO") or match or text.isupper() or text.endswith("(opcional)")
        if is_heading:
            extras(doc, current)
            current = None
            if text.startswith("CAPÍTULO"):
                chapter, title = text.split(". ", 1)
                p = doc.add_heading(chapter.replace("CAPÍTULO", "Capítulo") + ". " + title.capitalize(), 1)
                p.paragraph_format.page_break_before = True
                if text.startswith("CAPÍTULO VI."):
                    # Conserva la instrucción del original y añade la bibliografía realmente consultada.
                    i += 1
                    paragraph(doc, edited_text(texts[i]))
                    references(doc)
                seen.append(text)
            elif match:
                key, title = match.groups()
                level = min(key.count(".") + 1, 4)
                current = key
                body = None
                if ": " in title:
                    title, body = title.split(": ", 1)
                title = title[0].upper() + title[1:]
                if level == 4:
                    # APA nivel 4: encabezado con sangría, en negrita y seguido del texto.
                    p = paragraph(doc, "", "Heading 4")
                    p.add_run(f"{key} {title}. ").bold = True
                    if body is None and i + 1 < len(texts) and not re.match(r"^\d+(?:\.\d+)+\s", texts[i + 1]):
                        i += 1
                        body = edited_text(texts[i])
                    if body is None and key == "2.2.3.3":
                        body = ("Se formulan proposiciones contrastables sobre las relaciones entre variables; "
                                "su comprobación requiere un diseño congruente y evidencia empírica "
                                "(Conerly et al., 2021).")
                    run = p.add_run(body or "")
                    run.bold = False
                    run.italic = False
                else:
                    doc.add_heading(f"{key} {title}", level)
                    if body:
                        paragraph(doc, body)
                seen.append(key)
            else:
                doc.add_heading(text.capitalize(), 1)
                seen.append(text)
        else:
            paragraph(doc, text)
        i += 1
    extras(doc, current)
    return seen


def main():
    for folder in [ASSETS, BUILD]:
        folder.mkdir(parents=True, exist_ok=True)
    urls = {
        "openstax-enfoques.html": SOURCE + "2-1-approaches-to-sociological-research",
        "openstax-metodos.html": SOURCE + "2-2-research-methods",
        "openstax-resumen.html": SOURCE + "2-section-summary",
        "openstax-licencia.html": SOURCE + "preface",
        "figura1-original.webp": "https://openstax.org/apps/image-cdn/v1/f=webp/apps/archive/20260604.144757/resources/551ef5a8c1361bde7b39e03343d5baa18b8855ec",
        "figura2-original.webp": "https://openstax.org/apps/image-cdn/v1/f=webp/apps/archive/20260604.144757/resources/14b907f5c23f82c7415ffb5824d72cb5a432638e",
    }
    for name, url in urls.items():
        fetch(name, url)
    scientific_figure()
    with Image.open(ASSETS / "figura2-original.webp") as image:
        image.convert("RGB").save(ASSETS / "figura2-cuestionario.png")
    doc = Document(MATERIALS / "Portada.docx")
    configure_styles(doc)
    for section in doc.sections:
        section.page_width, section.page_height = Inches(8.5), Inches(11)
        section.top_margin = section.bottom_margin = section.left_margin = section.right_margin = Inches(1)
        section.header_distance = section.footer_distance = Inches(0.5)
    cover(doc)
    headings = content(doc)
    settings = doc.settings.element
    for element in settings.findall(qn("w:updateFields")):
        settings.remove(element)
    update = OxmlElement("w:updateFields")
    update.set(qn("w:val"), "true")
    settings.append(update)
    # Evita heredar colores o resaltados de la plantilla, incluso en estilos no usados.
    for root in [doc.element, doc.styles.element]:
        # Los atributos *Theme heredados de la portada prevalecen sobre font.name
        # en algunos motores. Fijar las cuatro familias evita que reaparezca Aptos.
        for fonts in root.xpath(".//w:rFonts"):
            fonts.attrib.clear()
            for family in ("ascii", "hAnsi", "eastAsia", "cs"):
                fonts.set(qn(f"w:{family}"), "Arial")
        for element in root.xpath(".//w:highlight | .//w:shd"):
            element.getparent().remove(element)
        for element in root.xpath(".//w:color"):
            element.attrib.clear()
            element.set(qn("w:val"), "000000")
    doc.core_properties.author = "Martín Jonathan de la Cruz Muñoz"
    doc.core_properties.title = "Tarea 6: Aplicación de formato APA e índice automatizado"
    doc.core_properties.subject = "Seminario I · Maestría en Gestión Administrativa"
    doc.core_properties.keywords = "APA 7; índice automatizado; investigación; tablas; figuras"
    output = BUILD / NAME
    doc.save(output)
    manifest = {
        "actividad": TASK_URL,
        "estado_aula_al_revisar": "Todavía no se han realizado envíos",
        "cierre_aula": "2026-09-13 23:59 (hora mostrada por el aula)",
        "fuentes": urls,
        "originales_sha256": {p.name: hashlib.sha256(p.read_bytes()).hexdigest() for p in MATERIALS.glob("*.docx")},
        "encabezados_conservados": headings,
        "tablas": len(doc.tables),
        "figuras_con_numero": 2,
        "advertencia": "La paginación y el resultado visible del TOC deben actualizarse antes de entregar.",
    }
    (BUILD / "manifiesto.json").write_text(json.dumps(manifest, ensure_ascii=False, indent=2) + "\n")
    print(f"Generado: {output}; {len(headings)} encabezados; {len(doc.tables)} tablas; 2 figuras.")


if __name__ == "__main__":
    main()