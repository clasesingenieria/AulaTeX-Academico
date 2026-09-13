"""Verifica integridad, estilos, campo TOC, contenido y geometría del PDF."""

from collections import Counter
import hashlib
import json
from pathlib import Path
import re
import subprocess
from zipfile import ZipFile

from docx import Document
from docx.oxml.ns import qn
from lxml import etree
from PIL import Image, ImageDraw

from generar_tarea6 import edited_text

ROOT = Path(__file__).resolve().parent
FINAL = ROOT.parent / "entregas" / "Tarea6_DeLaCruzMunoz.docx"
CHECK = ROOT / "validacion"
PDF = CHECK / "Tarea6_DeLaCruzMunoz_revision.pdf"


def normalized(text):
    return " ".join(text.lower().split())


def main():
    doc = Document(FINAL)
    ns = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
    with ZipFile(FINAL) as archive:
        assert archive.testzip() is None
        root = etree.fromstring(archive.read("word/document.xml"))
        styles_xml = etree.fromstring(archive.read("word/styles.xml"))
        instructions = root.xpath("//w:instrText/text()", namespaces=ns)
        toc = [i for i in instructions if "TOC" in i]
        assert len(toc) == 1 and '"1-3"' in toc[0], toc
        assert not root.xpath("//w:highlight", namespaces=ns)
        assert not root.xpath("//w:color[@w:val!='000000' and @w:val!='auto']", namespaces=ns)
        headers = [n for n in archive.namelist() if re.match(r"word/header\d+\.xml", n)]
        assert any("PAGE" in archive.read(h).decode() for h in headers)
        bookmarks = root.xpath("//w:bookmarkStart/@w:name", namespaces=ns)
        toc_links = root.xpath("//w:hyperlink/@w:anchor", namespaces=ns)
        assert len(toc_links) >= 50 and set(toc_links).issubset(set(bookmarks))
        body = " ".join("".join(p.xpath(".//w:t/text()", namespaces=ns))
                for p in root.xpath("//w:p", namespaces=ns))
    assert len(doc.tables) == 2
    assert len(doc.inline_shapes) == 3, "Dos figuras y la franja institucional"
    for section in doc.sections:
        assert all(abs(v.inches - 1) < 0.002 for v in [section.top_margin, section.bottom_margin, section.left_margin, section.right_margin])
        assert abs(section.page_width.inches - 8.5) < 0.002
        assert abs(section.page_height.inches - 11) < 0.002
    counts = Counter(p.style.name for p in doc.paragraphs)
    for n in (1, 2, 3):
        assert counts[f"Heading {n}"] > 0, counts
    assert doc.styles["Normal"].font.name == "Arial"
    assert doc.styles["Normal"].font.size.pt == 11
    assert doc.styles["Normal"].paragraph_format.line_spacing == 2
    assert doc.styles["APA referencia"].paragraph_format.first_line_indent.inches == -0.5
    for title in ["Tabla 1", "Tabla 2", "Figura 1", "Figura 2"]:
        assert sum(p.text == title for p in doc.paragraphs) == 1, title
    assert "MARTÍN JONATHAN DE LA CRUZ MUÑOZ" in body
    assert "DRA. CARLA OLIMPYA ZAPUCHE MORENO" in body
    for placeholder in ["NOMBRE DE TU TESIS", "TU NOMBRE COMPLETO", "NOVIEMBRE 2025", "pendiente de actualización"]:
        assert placeholder.lower() not in body.lower(), placeholder
    # Todo párrafo explicativo original está presente, con las correcciones documentadas.
    original = Document(ROOT / "materiales" / "Material para ejemplo.docx")
    retained = 0
    all_text = normalized(body)
    for p in original.paragraphs:
        text = p.text.strip()
        if not text or text.isupper() or text.endswith("(opcional)"):
            continue
        if re.match(r"^\d+(?:\.\d+)+\s", text):
            if ": " in text:
                text = text.split(": ", 1)[1]
            else:
                continue
        assert normalized(edited_text(text)) in all_text, f"Falta contenido: {text[:100]}"
        retained += 1
    layout = subprocess.check_output(["pdftotext", "-bbox", str(PDF), "-"], text=True)
    xml = etree.fromstring(layout.encode())
    n = {"x": "http://www.w3.org/1999/xhtml"}
    pages = xml.xpath("//x:page", namespaces=n)
    out_of_bounds = []
    for number, page in enumerate(pages, 1):
        words = page.xpath(".//x:word", namespaces=n)
        assert len(words) > 2, f"Página vacía: {number}"
        for word in words:
            if float(word.get("yMin")) < 65:
                assert word.text == str(number), f"Encabezado inesperado, página {number}"
                continue
            if float(word.get("xMin")) < 70 or float(word.get("xMax")) > 542 or float(word.get("yMax")) > 724:
                out_of_bounds.append((number, word.text))
    assert not out_of_bounds, out_of_bounds[:20]
    pagination = json.loads((CHECK / "paginacion.json").read_text())
    page_text = [normalized(" ".join(page.xpath(".//x:word/text()", namespaces=n))) for page in pages]
    for entry in pagination["indice_visible"].splitlines():
        heading, number = entry.rsplit("\t", 1)
        assert normalized(heading) in page_text[int(number) - 1], f"Página incorrecta en el índice: {entry}"
    # Portada, índice, marco teórico, tablas, figuras y referencias: vista de control.
    contact_pages = [1, 2]
    for number, page in enumerate(pages, 1):
        text = " ".join(page.xpath(".//x:word/text()", namespaces=n))
        if number > 4 and any(marker in text for marker in [
                "Capítulo II.", "Tabla 1 Métodos", "Tabla 2 Ejemplos", "Figura 1 Etapas",
                "Figura 2 Aplicación", "Capítulo VI."]):
            contact_pages.append(number)
    contact_pages = sorted(set(contact_pages))
    for page in contact_pages:
        subprocess.run(["pdftoppm", "-f", str(page), "-l", str(page), "-scale-to", "1100",
                        "-png", str(PDF), str(CHECK / "pagina")], check=True, capture_output=True)
    sheet = Image.new("RGB", (900, 1260), "#dddddd")
    draw = ImageDraw.Draw(sheet)
    for i, number in enumerate(contact_pages):
        with Image.open(CHECK / f"pagina-{number:02d}.png") as image:
            image.thumbnail((285, 380))
            x, y = (i % 3) * 300 + 7, (i // 3) * 420 + 24
            sheet.paste(image, (x, y))
            draw.text((x, y - 18), f"Página {number}", fill="black")
    sheet.save(CHECK / "vista-general.png")
    report = {
        "resultado": "Correcto", "paginas": len(pages),
        "tablas": len(doc.tables), "figuras": len(doc.inline_shapes) - 1,
        "campo_indice": toc[0].strip(), "enlaces_indice": len(toc_links),
        "paginas_indice_contrastadas_con_pdf": True,
        "estilos_encabezado": {k: v for k, v in counts.items() if k.startswith("Heading")},
        "parrafos_base_conservados_o_corregidos": retained,
        "margenes": "2.54 cm", "papel": "Carta", "fuente": "Arial 11",
        "interlineado_cuerpo": "Doble", "desbordamientos_texto_pdf": out_of_bounds,
        "sha256_docx": hashlib.sha256(FINAL.read_bytes()).hexdigest(),
        "validacion_visual": "Ver vista-general.png y las páginas individuales",
        "limite": "Paginado con LibreOffice; no se ha abierto en Microsoft Word ni enviado al aula.",
    }
    (CHECK / "verificacion.json").write_text(json.dumps(report, ensure_ascii=False, indent=2) + "\n")
    print(json.dumps(report, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()