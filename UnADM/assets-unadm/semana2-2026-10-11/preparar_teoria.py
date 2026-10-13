from __future__ import annotations

import hashlib
import json
from pathlib import Path
import shutil

from docx import Document
from docx.shared import Cm, Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
import win32com.client
import pymupdf

ROOT = Path(__file__).resolve().parents[3]
TARGET = ROOT / "UnADM/licenciatura-en-derecho-unadm/teoria-del-estado-y-constitucion-lde"
EVIDENCE = TARGET / "referencias-teoria-del-estado-y-constitucion/notas-teoria-del-estado-y-constitucion/actividad-02-linea-de-tiempo/revision-2026-10-11"
NAME = "2026_TeoriaEstado_S2_Martin Jonathan de la Cruz"
ENTRIES = [
    ("Hacia 3200 a. C. | Antiguedad: Uruk", "La concentracion de bienes y raciones requiere registros y administracion urbana; es un antecedente estatal, no un Estado nacional (Department of Ancient Near Eastern Art, 2003).", "Valoro la capacidad de organizar recursos, pero administrar no demuestra por si solo que el poder sea justo."),
    ("Siglo IV a. C. | Antiguedad: Aristoteles", "En la Politica relaciona la comunidad politica con la vida buena y distingue gobierno para el bien comun de gobierno para beneficio del gobernante (Miller, 2022, seccs. 2-3).", "Su criterio del bien comun permite discutir los fines del poder; rechazo tomar sus exclusiones de ciudadania como un modelo actual."),
    ("1215 | Edad Media: Carta Magna", "La crisis entre Juan y los barones ingleses lleva a limites pactados al monarca. La fragmentacion del poder y los privilegios estamentales no equivalen a derechos universales (UK Parliament, s. f.).", "Limitar la arbitrariedad es valioso, aunque el pacto protege principalmente intereses privilegiados."),
    ("1576 | Formacion del Estado moderno: Jean Bodin", "Los seis libros de la Republica sistematizan la soberania en el contexto de las guerras religiosas francesas. Bodin la concibe perpetua y absoluta respecto de leyes civiles, pero sometida a derecho natural y divino (Turchetti, 2018, secc. 4).", "La unidad de decision responde a la fragmentacion politica, pero no debe confundirse con una autorizacion de cualquier abuso."),
    ("1651 | Absolutismo: Thomas Hobbes", "Leviatan justifica una autoridad comun capaz de evitar la guerra civil; la seguridad y la paz fundamentan la obediencia. Su defensa de soberania indivisa no exige exclusivamente monarquia (Lloyd y Sreedhar, 2022, seccs. 1, 2 y 8).", "Reconozco el problema de la inseguridad, pero una autoridad eficaz necesita controles para que su proteccion no se vuelva dominacion."),
    ("1682 | Absolutismo: corte en Versalles", "Luis XIV instala la corte en Versalles: concentra a la nobleza y refuerza el control politico y administrativo de la monarquia francesa (Palace of Versailles, s. f.).", "La centralizacion puede organizar el gobierno sin legitimar privilegios ni repartir de manera igual el poder."),
    ("26 de agosto de 1789 | Revolucion francesa", "La Declaracion de Derechos del Hombre y del Ciudadano afirma igualdad juridica (art. 1), soberania nacional (art. 3), voluntad general (art. 6) y derechos con separacion de poderes (art. 16), frente al orden de privilegios (Asamblea Nacional Constituyente de Francia, 1789).", "Reconocer derechos cambia el criterio de legitimidad, pero una declaracion necesita garantias para producir proteccion efectiva."),
    ("5 de febrero de 1917 | Estado moderno constitucional: Mexico", "La Constitucion organiza poderes, derechos y soberania popular. Los arts. 39-41 y 49 permiten evaluar el fundamento y los limites del poder. Es un hito mexicano, no el nacimiento universal del Estado moderno (Camara de Diputados, 2026).", "Valoro que la autoridad se instituya para beneficio del pueblo; su existencia solo se justifica si respeta competencias y controles."),
    ("10 de junio de 2011 | Estado constitucional y derechos humanos", "La reforma mexicana fortalece las obligaciones de promover, respetar, proteger y garantizar derechos humanos, hoy contenidas en el art. 1 (Camara de Diputados, 2026; Garcia Ricci, 2011).", "Considero necesario valorar la reforma por la proteccion efectiva frente al abuso y no solo por su reconocimiento escrito."),
]
REFERENCES = [
    "Asamblea Nacional Constituyente de Francia. (1789, 26 de agosto). Declaration of Human and Civic Rights of 26 August 1789. Conseil constitutionnel. https://www.conseil-constitutionnel.fr/en/declaration-of-human-and-civic-rights-of-26-august-1789",
    "Camara de Diputados del H. Congreso de la Union. (2026). Constitucion Politica de los Estados Unidos Mexicanos (ultima reforma DOF 7 de octubre de 2026; arts. 1, 39-41 y 49). https://www.diputados.gob.mx/LeyesBiblio/pdf/CPEUM.pdf",
    "Department of Ancient Near Eastern Art. (2003, octubre). Uruk: The first city. The Metropolitan Museum of Art. https://www.metmuseum.org/essays/uruk-the-first-city",
    "Garcia Ricci, D. (2011). Estado de derecho y principio de legalidad. Comision Nacional de los Derechos Humanos. https://corteidh.or.cr/tablas/r28801.pdf",
    "Lloyd, S. A., y Sreedhar, S. (2022). Hobbes's moral and political philosophy. En E. N. Zalta y U. Nodelman (Eds.), The Stanford Encyclopedia of Philosophy. https://plato.stanford.edu/entries/hobbes-moral/",
    "Miller, F. (2022). Aristotle's political theory. En E. N. Zalta (Ed.), The Stanford Encyclopedia of Philosophy. https://plato.stanford.edu/entries/aristotle-politics/",
    "Palace of Versailles. (s. f.). Louis XIV. https://en.chateauversailles.fr/discover/history/great-characters/louis-xiv",
    "Turchetti, M. (2018). Jean Bodin. En E. N. Zalta (Ed.), The Stanford Encyclopedia of Philosophy. https://plato.stanford.edu/entries/bodin/",
    "UK Parliament. (s. f.). Magna Carta. https://www.parliament.uk/about/living-heritage/evolutionofparliament/originsofparliament/birthofparliament/overview/magnacarta/",
]


def paragraph_before(anchor, text, bold=False):
    paragraph = anchor.insert_paragraph_before(text)
    paragraph.paragraph_format.space_after = Pt(7)
    paragraph.paragraph_format.line_spacing = 1.15
    for run in paragraph.runs:
        run.font.name = "Montserrat"
        run.font.size = Pt(10)
        run.bold = bold
    return paragraph


def main() -> None:
    EVIDENCE.mkdir(parents=True, exist_ok=True)
    template = next(TARGET.glob("referencias-*/notas-*/materiales-generales/auditorias/semana2-2026-10-11/*Formato*.docx"))
    template_hash = hashlib.sha256(template.read_bytes()).hexdigest()
    document = Document(template)
    table = document.tables[0]
    table.cell(1, 0).text = "Nombre del Estudiante:\nMartin Jonathan de la Cruz\nMatricula: ES2611202040"
    table.cell(1, 1).text = "Semana: 2\nActividad: S2"
    table.cell(2, 0).text = "Tipo de actividad: Linea de tiempo\nHistoria y origen del Estado"
    table.cell(2, 1).text = "Fecha: 11/10/26"
    for row in table.rows:
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:
                    run.font.name = "Montserrat"
                    run.font.size = Pt(10)
    opinion = next(p for p in document.paragraphs if p.text.startswith("II."))
    bibliography = next(p for p in document.paragraphs if p.text.startswith("III."))
    paragraph_before(opinion, "Historia y origen del Estado: cronologia de hitos", True)
    paragraph_before(opinion, "La secuencia distingue antecedentes urbanos, organizacion medieval, formacion del Estado moderno y constitucionalismo. Las etapas se superponen: el Estado moderno se forma en los siglos XVI-XVII; no nace despues de 1789. Fechas e ideas se relacionan con necesidades economicas, conflictos y criterios de legitimidad, no con un progreso inevitable.")
    for heading, explanation, evaluation in ENTRIES:
        heading_paragraph = paragraph_before(opinion, heading, True)
        explanation_paragraph = paragraph_before(opinion, explanation)
        evaluation_paragraph = paragraph_before(opinion, "Opinion critica: " + evaluation)
        heading_paragraph.paragraph_format.keep_with_next = True
        explanation_paragraph.paragraph_format.keep_with_next = True
        for paragraph in (heading_paragraph, explanation_paragraph, evaluation_paragraph):
            paragraph.paragraph_format.keep_together = True
            paragraph.paragraph_format.left_indent = Cm(0.4)
            borders = OxmlElement('w:pBdr')
            left = OxmlElement('w:left')
            for attribute, value in [('val', 'single'), ('sz', '12'), ('space', '8'), ('color', '174A3A')]:
                left.set(qn('w:' + attribute), value)
            borders.append(left)
            paragraph._p.get_or_add_pPr().append(borders)
    paragraph_before(bibliography, "Considero que el fin del Estado debe distinguirse de la permanencia de sus gobernantes. Uruk muestra administracion, la Carta Magna limites pactados y Bodin y Hobbes respuestas distintas a la fragmentacion y al conflicto. La centralizacion no garantiza justicia; los derechos y la separacion de poderes exigen una justificacion adicional del mando.")
    paragraph_before(bibliography, "Desde esta comparacion, el constitucionalismo actual permite preguntar quien tiene competencia, para que ejerce el poder y que controles impiden el abuso. El art. 39 constitucional instituye el poder en beneficio del pueblo y el art. 49 divide el Supremo Poder de la Federacion (Camara de Diputados, 2026). Por eso valoro las instituciones por su proteccion efectiva y no solo por estabilidad o eficacia.")
    for reference in REFERENCES:
        paragraph = document.add_paragraph(reference)
        paragraph.paragraph_format.left_indent = Cm(0.6)
        paragraph.paragraph_format.first_line_indent = Cm(-0.6)
        paragraph.paragraph_format.space_after = Pt(8)
        for run in paragraph.runs:
            run.font.name = "Montserrat"
            run.font.size = Pt(10)
    document.add_paragraph("Asistencia de IA: se utilizo GitHub Copilot para organizar fuentes, contrastar hechos y preparar la redaccion y el formato. No sustituye el analisis propio ni acredita experiencia personal; las valoraciones se presentan como argumentos academicos.")
    for section in document.sections:
        footer = section.footer.paragraphs[0]
        footer.clear()
        footer.add_run("Pagina ")
        field = OxmlElement("w:fldSimple")
        field.set(qn("w:instr"), "PAGE")
        footer._p.append(field)
    source = EVIDENCE / f"{NAME}.docx"
    document.save(source)
    assert hashlib.sha256(template.read_bytes()).hexdigest() == template_hash
    destination = TARGET / "Entregas" / f"{NAME}.pdf"
    destination.parent.mkdir(exist_ok=True)
    if destination.exists():
        prior = EVIDENCE / 'validacion.json'
        if not prior.exists() or json.loads(prior.read_text(encoding='utf-8'))['estado'] != 'revision local; no enviado':
            raise RuntimeError("PDF de entrega existente no reconocido como borrador")
        previous = EVIDENCE / 'antes-ajuste-visual.pdf'
        if previous.exists():
            raise RuntimeError('Respaldo visual existente')
        shutil.copy2(destination, previous)
    application = win32com.client.DispatchEx("Word.Application")
    application.Visible = False
    application.DisplayAlerts = 0
    temporary = ROOT / 'tmp/unadm-teoria-S2-exportar.docx'
    temporary.parent.mkdir(exist_ok=True)
    shutil.copy2(source, temporary)
    try:
        opened = application.Documents.Open(str(temporary.resolve()), ReadOnly=True)
        if opened is None:
            raise RuntimeError('Word no abrio la copia temporal')
        try:
            opened.ExportAsFixedFormat(str(destination.resolve()), 17)
        finally:
            opened.Close(False)
    finally:
        application.Quit()
        temporary.unlink(missing_ok=True)
    with pymupdf.open(destination) as pdf:
        text = "\n".join(page.get_text() for page in pdf)
        for required in ("Uruk", "Aristoteles", "Bodin", "Hobbes", "1789", "1917", "II.", "III.", "Asistencia de IA"):
            assert required in text, required
        for number, page in enumerate(pdf, 1):
            page.get_pixmap(matrix=pymupdf.Matrix(1.2, 1.2)).save(EVIDENCE / f"revision-p{number}.png")
        result = {"archivo": destination.relative_to(TARGET).as_posix(), "sha256": hashlib.sha256(destination.read_bytes()).hexdigest(), "paginas": len(pdf), "formato_docente": template.relative_to(TARGET).as_posix(), "sha256_original": template_hash, "consigna": "214812", "estado": "revision local; no enviado", "rubrica": "3+2+2+2+1=10; sin autocalificacion"}
    (EVIDENCE / "validacion.json").write_text(json.dumps(result, indent=2), encoding="utf-8")
    print(json.dumps(result, ensure_ascii=True))


if __name__ == "__main__":
    main()