import tempfile
import unittest
from pathlib import Path

from actualizar_metadatos_itesca import eligible, update_tex, update_word


class MetadataTests(unittest.TestCase):
    def test_definitions_only_and_alias_preserved(self):
        text = r"\providecommand{\itescastudentid}{Por confirmar}" + "\n" + r"\def\itescastudentemail{correo.institucional@itesca.edu.mx}" + "\n" + r"\def\itescastudentid{\actividadmatricula}" + "\nDocente: docente@example.org"
        result = update_tex(text, "12345678", "alumno@example.org")
        self.assertIn(r"\itescastudentid}{12345678}", result)
        self.assertIn(r"\itescastudentemail{alumno@example.org}", result)
        self.assertIn(r"\itescastudentid{\actividadmatricula}", result)
        self.assertIn("Docente: docente@example.org", result)
        self.assertEqual(result, update_tex(result, "12345678", "alumno@example.org"))

    def test_excludes_original_sources_and_history(self):
        root = Path("ITESCA")
        self.assertFalse(eligible(root / "materia/referencias-materia/reporte-original.tex", root))
        self.assertFalse(eligible(root / "materia/historico/reporte-antiguo.tex", root))
        self.assertTrue(eligible(root / "materia/reporte-Actividad-1.tex", root))

    def test_itesca_templates_do_not_insert_platform_labels(self):
        root = Path(__file__).resolve().parents[1]
        subject = root / "ITESCA/maestria-en-gestion-administrativa/plan-de-negocios-mga"
        for name in ("reporte-plan-de-negocios-plantilla-actividad.tex", "formato-itesca-plan-negocios.tex", "formato-itesca-plan-negocios-generadas.tex"):
            text = (subject / name).read_text(encoding="utf-8")
            for label in ("Actividad local", "Módulo Moodle", "Vencimiento publicado"):
                self.assertNotIn(label, text)
        from aulatex.activity_contract import REALIZAR_ACTIVIDAD_PIPELINE_CONTRACT
        self.assertIn("itesca_operational_metadata", REALIZAR_ACTIVIDAD_PIPELINE_CONTRACT["visible_text_rules"])

    def test_word_keeps_academic_text_and_teacher(self):
        from docx import Document
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "reporte.docx"
            document = Document()
            table = document.add_table(rows=4, cols=2)
            for row, label, value in zip(table.rows, ["Matrícula:", "Correo institucional:", "Docente:", "Módulo Moodle:"], ["Por confirmar", "Sin correo", "Persona docente", "6548"]):
                row.cells[0].text, row.cells[1].text = label, value
            document.add_paragraph("Actividad local 5. Módulo Moodle 6548.")
            document.add_paragraph("Vencimiento publicado: 21/09/2026")
            document.add_paragraph("La lectura fue consultada en Moodle.")
            document.save(path)
            self.assertTrue(update_word(path, "12345678", "alumno@example.org", True))
            updated = Document(path)
            self.assertEqual(len(updated.tables[0].rows), 3)
            self.assertEqual(updated.tables[0].rows[2].cells[1].text, "Persona docente")
            self.assertEqual([paragraph.text for paragraph in updated.paragraphs], ["La lectura fue consultada en Moodle."])

    def test_word_does_not_fill_survey_contact(self):
        from docx import Document
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "encuesta.docx"
            document = Document()
            cells = document.add_table(rows=1, cols=2).rows[0].cells
            cells[0].text = "Correo electrónico"
            cells[1].text = "Respuesta del participante"
            document.save(path)
            self.assertFalse(update_word(path, "12345678", "alumno@example.org", True))


if __name__ == "__main__":
    unittest.main()