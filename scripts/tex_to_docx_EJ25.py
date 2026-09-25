import re
from pathlib import Path
from docx import Document
from docx.shared import Pt

tex_path = Path(r"c:\Users\Sysx\Documents\AulaTeX-Academico\UANL\ingeniero-agronomo\algebra-lineal\Segundo_Laboratorio_EJ25_solucion.tex")
docx_path = tex_path.with_suffix('.docx')
text = tex_path.read_text(encoding='utf-8')

# extract body between \begin{document} and \end{document}
m = re.search(r"\\begin\{document\}(.*)\\end\{document\}", text, flags=re.DOTALL)
body = m.group(1) if m else text

# normalize
body = body.replace('\r\n', '\n')

doc = Document()
doc.styles['Normal'].font.name = 'Times New Roman'
doc.styles['Normal'].font.size = Pt(12)

pos = 0
length = len(body)

# helper to add matrix from pmatrix text
def add_matrix_from_pmatrix(s):
    # s contains \begin{pmatrix} ... \end{pmatrix}
    mm = re.search(r"\\begin\{pmatrix\}(.*?)\\end\{pmatrix}", s, flags=re.DOTALL)
    if not mm:
        doc.add_paragraph(s.strip())
        return
    content = mm.group(1).strip()
    rows = [r.strip() for r in re.split(r"\\\\", content) if r.strip()]
    matrix = [ [c.strip() for c in row.split('&')] for row in rows]
    table = doc.add_table(rows=len(matrix), cols=max(len(r) for r in matrix))
    for i,row in enumerate(matrix):
        for j,val in enumerate(row):
            table.cell(i,j).text = val

while pos < length:
    # search next special token
    next_sec = re.search(r"\\section\*\{([^}]*)\}", body[pos:])
    next_sub = re.search(r"\\subsection\*\{([^}]*)\}", body[pos:])
    next_itemize = re.search(r"\\begin\{itemize\}", body[pos:])
    next_math = re.search(r"\\\\\[(.*?)\\\\\]", body[pos:], flags=re.DOTALL)
    # find earliest
    candidates = []
    if next_sec:
        candidates.append(('sec', next_sec.start()+pos, next_sec))
    if next_sub:
        candidates.append(('sub', next_sub.start()+pos, next_sub))
    if next_itemize:
        candidates.append(('itemize', next_itemize.start()+pos, next_itemize))
    if next_math:
        candidates.append(('math', next_math.start()+pos, next_math))
    if not candidates:
        remaining = body[pos:].strip()
        if remaining:
            # split into lines and add
            for line in remaining.splitlines():
                line = line.strip()
                if line:
                    doc.add_paragraph(line)
        break
    # pick earliest
    kind, idx, match = min(candidates, key=lambda x: x[1])
    # add text before it as normal paragraph
    before = body[pos:idx].strip()
    if before:
        for line in before.splitlines():
            line = line.strip()
            if line:
                doc.add_paragraph(line)
    if kind == 'sec':
        title = match.group(1).strip()
        doc.add_heading(title, level=1)
        pos = idx + match.end() - match.start()
    elif kind == 'sub':
        title = match.group(1).strip()
        doc.add_heading(title, level=2)
        pos = idx + match.end() - match.start()
    elif kind == 'itemize':
        # find end
        m_end = re.search(r"\\end\{itemize\}", body[idx:])
        if m_end:
            block = body[idx: idx + m_end.end()]
            # extract \item lines
            items = re.findall(r"\\item\s*(.*?)($|\\\\item|\\\\end\{itemize\})", block, flags=re.DOTALL)
            for it in items:
                text_item = it[0].strip()
                if text_item:
                    p = doc.add_paragraph(style='List Bullet')
                    p.add_run(text_item)
            pos = idx + m_end.end()
        else:
            pos = idx + len('\\begin{itemize}')
    elif kind == 'math':
        math_text = match.group(1).strip()
        # if contains pmatrix, add table
        if '\\begin{pmatrix}' in math_text:
            add_matrix_from_pmatrix(math_text)
        else:
            # add math as normal paragraph, replace LaTeX commands
            cleaned = re.sub(r'\\[a-zA-Z]+', '', math_text)
            cleaned = cleaned.replace('\\quad','    ')
            cleaned = cleaned.replace('\\,','')
            doc.add_paragraph(cleaned.strip())
        pos = idx + match.end() - match.start()

# save
doc.save(docx_path)
print('WROTE DOCX:', docx_path)
