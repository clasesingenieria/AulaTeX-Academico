from __future__ import annotations

import argparse
import json
import re
from pathlib import Path


MACROS = {
    "itescastudentid": "id",
    "actividadmatricula": "id",
    "studentcontrol": "id",
    "itescastudentemail": "email",
    "actividadcorreo": "email",
}
NAMES = "|".join(MACROS)
DEFINITIONS = re.compile(
    rf"(?P<head>\\(?:providecommand|newcommand|renewcommand)\s*\{{\\(?P<command>{NAMES})\}}\s*\{{|"
    rf"\\def\s*\\(?P<definition>{NAMES})\s*\{{)(?P<value>[^{{}}]*)(?P<end>\}})"
)


def update_tex(text: str, student_id: str, email: str) -> str:
    values = {"id": student_id, "email": email}

    def replace(match: re.Match[str]) -> str:
        if "\\" in match["value"]:
            return match[0]
        macro = match["command"] or match["definition"]
        return match["head"] + values[MACROS[macro]] + match["end"]

    return DEFINITIONS.sub(replace, text)


def eligible(path: Path, root: Path) -> bool:
    if any(part.startswith(("referencias-", "assets-", ".memoria", ".venv")) or part in
           {"historico", "validacion", "materiales", "node_modules"} for part in path.relative_to(root).parts[:-1]):
        return False
    return path.name.startswith(("reporte-", "presentacion-", "formato-", "imagen-empresa-", "Tarea"))


def update_word(source: Path, student_id: str, email: str, apply: bool) -> bool:
    from docx import Document

    document = Document(source)
    changed = False
    values = {"matricula": student_id, "matrícula": student_id, "numero de control": student_id,
              "número de control": student_id, "no. de control": student_id,
              "correo institucional": email, "correo electronico": email, "correo electrónico": email}
    for table in document.tables:
        labels = {row.cells[0].text.strip().rstrip(":").lower() for row in table.rows if row.cells}
        student_table = bool(labels.intersection({"alumno", "estudiante", "matricula", "matrícula", "número de control", "numero de control", "no. de control"}))
        for row in list(table.rows):
            if len(row.cells) < 2:
                continue
            label = row.cells[0].text.strip().rstrip(":").lower()
            if label in {"módulo", "modulo", "módulo moodle", "modulo moodle", "vencimiento publicado"}:
                table._tbl.remove(row._tr)
                changed = True
            elif student_table and label in values and row.cells[1].text != values[label]:
                row.cells[1].text = values[label]
                changed = True
    for paragraph in list(document.paragraphs):
        text = paragraph.text.strip()
        if re.match(r"^(Actividad local\s+\d+\.|M[oó]dulo Moodle\s+\d+|Vencimiento publicado:)", text):
            paragraph._p.getparent().remove(paragraph._p)
            changed = True
        elif re.fullmatch(r"Actividad \d+\s*[·-]\s*M[oó]dulo Moodle \d+", text):
            paragraph.text = re.match(r"Actividad \d+", text)[0]
            changed = True
    subject = document.core_properties.subject
    cleaned = re.sub(r"\s*[-·]\s*(?:M[oó]dulo\s+)?Moodle\s+\d+", "", subject)
    if cleaned != subject:
        document.core_properties.subject = cleaned
        changed = True
    if changed and apply:
        document.save(source)
    return changed


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--matricula", required=True)
    parser.add_argument("--correo", required=True)
    parser.add_argument("--apply", action="store_true")
    args = parser.parse_args()
    if not re.fullmatch(r"\d+", args.matricula) or not re.fullmatch(r"[\w.+-]+@[\w.-]+", args.correo):
        parser.error("Matrícula o correo no válidos.")
    root = Path(__file__).resolve().parents[1] / "ITESCA"
    changed, locked = [], []
    for path in sorted(root.rglob("*")):
        if not path.is_file() or not eligible(path, root):
            continue
        if path.suffix == ".tex":
            with path.open(encoding="utf-8-sig", newline="") as stream:
                text = stream.read()
            updated = update_tex(text, args.matricula, args.correo)
            if updated != text:
                if args.apply:
                    with path.open("w", encoding="utf-8", newline="") as stream:
                        stream.write(updated)
                changed.append(path.relative_to(root).as_posix())
        elif path.suffix == ".docx":
            try:
                if update_word(path, args.matricula, args.correo, args.apply):
                    changed.append(path.relative_to(root).as_posix())
            except PermissionError:
                locked.append(path.relative_to(root).as_posix())
    print(json.dumps({"applied": args.apply, "changed": changed, "locked": locked}, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()