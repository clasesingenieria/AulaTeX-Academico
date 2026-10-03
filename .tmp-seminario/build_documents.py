from pathlib import Path
from copy import deepcopy
import json, re, hashlib, shutil, textwrap
from zipfile import ZipFile, ZIP_DEFLATED
from lxml import etree
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from PIL import Image, ImageDraw, ImageFont

ROOT=Path(__file__).resolve().parents[1]
COURSE=ROOT/'ITESCA/maestria-en-gestion-administrativa/seminario-i-mga'
WORK=ROOT/'.tmp-seminario'
OUT=COURSE/'entregas'
REFS=COURSE/'referencias-seminario-i/vtaxi-2026-09-27'
DATA=json.loads((WORK/'content.json').read_text(encoding='utf8'))
SOURCE=COURSE/'entregas/Tarea6_DeLaCruzMunoz.docx'
OFFICIAL=COURSE/'anteproyecto/evidencias/moodle-2026-09-27/03.01 Seleccion de proyecto de titulacion 25130604.docx'
NAME='Martín Jonathan de la Cruz Muñoz'
TEACHER='Dra. Carla Olimpya Zapuche Moreno'
LGAC='Gestión e Innovación de las Organizaciones'

def text_runs(p,text):
    for i,part in enumerate(text.split('*')):
        r=p.add_run(part)
        r.italic=(i%2==1)
    return p

def style_para(p, indent=True, single=False):
    f=p.paragraph_format
    f.first_line_indent=Inches(.5) if indent else Inches(0)
    f.space_before=Pt(0); f.space_after=Pt(0)
    f.line_spacing=1 if single else 2
    f.widow_control=True
    p.alignment=WD_ALIGN_PARAGRAPH.LEFT
    return p

def body(d,text):
    p=d.add_paragraph(style='Normal');style_para(p);text_runs(p,text);return p

def h(d,text,level=1,page=False):
    style=next(s for s in d.styles if s.style_id==f'Heading{level}')
    p=d.add_paragraph(text,style=style)
    p.paragraph_format.page_break_before=page
    p.paragraph_format.keep_with_next=True
    p.paragraph_format.first_line_indent=Inches(0)
    p.paragraph_format.space_before=Pt(12 if level>1 else 6)
    p.paragraph_format.space_after=Pt(4)
    p.paragraph_format.line_spacing=2
    for r in p.runs:r.font.color.rgb=RGBColor(0,0,0)
    return p

def future(d):body(d,'Se completará en una tarea posterior.')

def add_toc(d):
    p=d.add_paragraph('Índice',style='APA título índice')
    p.paragraph_format.page_break_before=True
    p=d.add_paragraph();style_para(p,False)
    r=p.add_run();e=OxmlElement('w:fldChar');e.set(qn('w:fldCharType'),'begin');r._r.append(e)
    r=p.add_run();e=OxmlElement('w:instrText');e.set(qn('xml:space'),'preserve');e.text=' TOC \\o "1-2" \\h \\z \\u ';r._r.append(e)
    r=p.add_run();e=OxmlElement('w:fldChar');e.set(qn('w:fldCharType'),'separate');r._r.append(e)
    p.add_run('Índice automático')
    r=p.add_run();e=OxmlElement('w:fldChar');e.set(qn('w:fldCharType'),'end');r._r.append(e)

def base(task):
    normalized=WORK/'t6-normalized.docx'
    with ZipFile(SOURCE) as z,ZipFile(normalized,'w',ZIP_DEFLATED) as dest:
        for info in z.infolist():
            content=z.read(info.filename)
            if info.filename=='_rels/.rels':
                content=content.replace(b'http://schemas.openxmlformats.org/officedocument/2006/relationships/metadata/core-properties',b'http://schemas.openxmlformats.org/package/2006/relationships/metadata/core-properties')
            dest.writestr(info,content)
    d=Document(normalized)
    keep={p._p for p in d.paragraphs[:15]}
    for child in list(d.element.body):
        if child.tag!=qn('w:sectPr') and child not in keep:d.element.body.remove(child)
    changes={1:'INSTITUTO TECNOLÓGICO SUPERIOR DE CAJEME',2:DATA['title'],3:f'Seminario I\nTarea {task} '+('Antecedentes' if task==8 else 'Planteamiento del problema'),4:'MAESTRÍA EN GESTIÓN ADMINISTRATIVA',5:'Línea de Generación y Aplicación del Conocimiento',6:LGAC,7:'',8:'PRESENTA',9:NAME.upper(),10:'Matrícula 26130503',11:'DOCENTE DE SEMINARIO I',12:TEACHER.upper(),13:'',14:'CIUDAD OBREGÓN, SONORA\n27 DE SEPTIEMBRE DE 2026'}
    for i,value in changes.items():
        p=d.paragraphs[i];p.clear();p.add_run(value)
        p.alignment=WD_ALIGN_PARAGRAPH.CENTER
        f=p.paragraph_format;f.first_line_indent=Inches(0);f.left_indent=Inches(0);f.line_spacing=1.5
        f.space_before=Pt(0);f.space_after=Pt(10 if i not in [7,13] else 4)
        f.keep_with_next=False;f.page_break_before=False
        for r in p.runs:
            r.font.name='Arial';r.font.size=Pt(11);r.font.color.rgb=RGBColor(0,0,0);r.bold=i in [1,2,4,6,9,12]
    d.paragraphs[2].style=d.styles['Title']
    d.paragraphs[2].paragraph_format.space_after=Pt(14)
    for styname in ['Normal','Title','Heading 1','Heading 2','Heading 3','TOC 1','TOC 2','TOC 3']:
        s=next((x for x in d.styles if x.name.casefold()==styname.casefold()),None)
        if s is None:continue
        s.font.name='Arial';s.font.size=Pt(11);s.font.color.rgb=RGBColor(0,0,0)
        s.paragraph_format.line_spacing=2
        for e in s.element.xpath('.//w:rFonts'):
            for a in list(e.attrib):
                if 'Theme' in a:del e.attrib[a]
        if styname.startswith('TOC'):
            s.paragraph_format.line_spacing=1.15
            s.paragraph_format.space_after=Pt(2)
            s.paragraph_format.first_line_indent=Inches(0)
    # The requested template body is replaced, while headers and cover images remain.
    d.core_properties.title=DATA['title'];d.core_properties.author=NAME
    d.core_properties.subject=f'Seminario I Tarea {task}'
    d.core_properties.comments=''
    add_toc(d)
    return d

def references():
    tex=(COURSE/'reporte-seminario-i-Actividad-3.tex').read_text(encoding='utf8')
    block=tex.split('\\begin{apareferences}',1)[1].split('\\end{apareferences}',1)[0]
    arr=[]
    for s in re.split(r'\\item\s+',block)[1:]:
        s=re.sub(r'\\emph\{([^}]+)\}',r'*\1*',s)
        s=re.sub(r'\\url\{([^}]+)\}',r'\1',s)
        s=s.replace(r'\&','&').replace('--','–').strip()
        if s.startswith('Omrani,'):s=s.replace('(2022)','(2024)',1)
        if s.startswith('Troise,'):s=s.replace('(2021)','(2022)',1)
        s=s.replace('Article 35','Artículo 35')
        arr.append(s)
    arr.append('vTaxi. (2026). *Documentación interna del proyecto vTaxi* [README, arquitectura, ficha técnica y borradores regulatorios; corte documental del 27 de septiembre de 2026].')
    return arr

def reference_section(d,task):
    h(d,'Referencias',page=True)
    refs=references()
    if task==9:
        refs.extend([
          'H. Congreso del Estado de Nuevo León. (2026). *Ley de Movilidad Sostenible, de Accesibilidad y Seguridad Vial para el Estado de Nuevo León* (última reforma publicada el 22 de mayo de 2026; arts. 82, 83, 99 y 100). https://www.hcnl.gob.mx/trabajo_legislativo/leyes/leyes/ley_de_movilidad_sostenible_y_accesibilidad_para_el_estado_de_nuevo_leon/',
          'Instituto de Capacitación y Educación para el Trabajo del Estado de Nuevo León. (s. f.). *ICET*. Recuperado el 27 de septiembre de 2026, de https://icetnl.mx/',
          'Instituto Tecnológico Superior de Cajeme. (s. f.). *Introducción al uso de la matriz de consistencia y el círculo de Covey* [Material didáctico de Seminario I, pp. 4–6]. https://cursos3.e-itesca.edu.mx/mod/assign/view.php?id=2887'
        ])
    for text in sorted(refs,key=lambda x:x.casefold()):
        p=d.add_paragraph();style_para(p,False)
        p.paragraph_format.left_indent=Inches(.5);p.paragraph_format.first_line_indent=Inches(-.5)
        p.paragraph_format.keep_together=True
        text_runs(p,text)

def consistency(d):
    h(d,'Anexo 1 Matriz de consistencia',page=True)
    p=d.add_paragraph('Versión inicial',style='Normal');style_para(p,False)
    table=d.add_table(rows=1,cols=2)
    table.alignment=WD_TABLE_ALIGNMENT.CENTER;table.autofit=False
    table.columns[0].width=Inches(1.35);table.columns[1].width=Inches(5.15)
    table.rows[0].cells[0].text='Elemento';table.rows[0].cells[1].text='Contenido del anteproyecto'
    rows=[('Tema',DATA['title']),('Problema general',DATA['problem']),('Pregunta general',DATA['question']),('Preguntas específicas','\n\n'.join(f'{i+1}. {s}' for i,s in enumerate(DATA['specific_questions'])))]
    for a,b in rows:
        c=table.add_row().cells;c[0].text=a;c[1].text=b
    pr=table._tbl.tblPr
    borders=OxmlElement('w:tblBorders')
    for side in ['top','left','bottom','right','insideH','insideV']:
        e=OxmlElement('w:'+side);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');borders.append(e)
    pr.append(borders)
    for ri,row in enumerate(table.rows):
        trpr=row._tr.get_or_add_trPr();cant=OxmlElement('w:cantSplit');trpr.append(cant)
        if ri==0:trpr.append(OxmlElement('w:tblHeader'))
        for ci,c in enumerate(row.cells):
            c.width=Inches(1.35 if ci==0 else 5.15);c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tcpr=c._tc.get_or_add_tcPr();m=OxmlElement('w:tcMar')
            for side in ['top','bottom','left','right']:
                e=OxmlElement('w:'+side);e.set(qn('w:w'),'110');e.set(qn('w:type'),'dxa');m.append(e)
            tcpr.append(m)
            if ri==0:
                sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'E9EDF2');tcpr.append(sh)
            for p in c.paragraphs:
                style_para(p,False,True);p.paragraph_format.space_after=Pt(3)
                for r in p.runs:r.font.name='Arial';r.font.size=Pt(10);r.bold=(ri==0)
    p=d.add_paragraph();style_para(p,False)
    text_runs(p,'*Nota.* Elaboración propia. Se completan únicamente los cuatro campos solicitados para esta etapa.')

def covey_image():
    image=Image.new('RGB',(1950,1570),'white');dr=ImageDraw.Draw(image)
    fontpath=Path('C:/Windows/Fonts/arial.ttf');boldpath=Path('C:/Windows/Fonts/arialbd.ttf')
    def font(n,b=False):return ImageFont.truetype(str(boldpath if b else fontpath),n)
    def center(text,x,y,f,color='black',width=None):
        if width:text='\n'.join(textwrap.wrap(text,width))
        box=dr.multiline_textbbox((0,0),text,font=f,spacing=10,align='center')
        dr.multiline_text((x-(box[2]-box[0])/2,y),text,font=f,fill=color,spacing=10,align='center')
    for box in [(55,90,1895,1470),(145,175,1805,1385),(245,260,1705,1300),(355,350,1595,1210)]:
        dr.ellipse(box,outline='#465971',width=5)
    center('VARIABLE 1',975,126,font(34,True))
    center('VARIABLE 2',975,208,font(34,True))
    center('VARIABLE 3',975,292,font(34,True))
    center('ANÁLISIS Y CORRELACIÓN\nDE VARIABLES',975,410,font(35,True))
    dr.ellipse((410,540,1540,1055),fill='#F0F3F7',outline='#465971',width=5)
    center('PROBLEMA GENERAL',975,615,font(42,True))
    center(DATA['problem'],975,695,font(38),width=44)
    center('CONCLUSIONES DESCUBRIMIENTOS\nY PROPUESTAS',975,1100,font(31,True))
    center('Los demás componentes se desarrollarán en etapas posteriores',975,1510,font(29))
    p=WORK/'covey.png';image.save(p);return p

def build(task):
    d=base(task)
    h(d,'1 Antecedentes',page=True)
    for title,paras in DATA['background']:
        p=d.add_paragraph();style_para(p,False);p.add_run(title).bold=True;p.paragraph_format.keep_with_next=True
        for text in paras:body(d,text)
    h(d,'2 Planteamiento del problema',page=True)
    if task==8:
        future(d)
        h(d,'2.1 Pregunta general',2);future(d)
        h(d,'2.2 Preguntas específicas',2);future(d)
    else:
        for text in DATA['problem_intro']:body(d,text)
        h(d,'2.1 Enunciado del problema',2)
        for text in DATA['problem_statement']:body(d,text)
        p=body(d,DATA['problem']);p.runs[0].bold=True
        h(d,'2.2 Formulación del problema',2)
        body(d,DATA['question'])
        for text in DATA['variables']:body(d,text)
        h(d,'2.3 Preguntas específicas',2)
        for i,text in enumerate(DATA['specific_questions']):
            p=body(d,f'{i+1}. {text}');p.paragraph_format.left_indent=Inches(.25);p.paragraph_format.first_line_indent=Inches(-.25)
    h(d,'3 Formulación de objetivos general y específicos')
    body(d,'Los apartados 3 a 9 se completarán en tareas posteriores. Se conservan sus encabezados para continuar el anteproyecto.')
    h(d,'3.1 Objetivo general',2)
    h(d,'3.2 Objetivos específicos',2)
    h(d,'4 Justificación')
    h(d,'5 Hipótesis de trabajo')
    h(d,'6 Alcances y limitaciones')
    h(d,'6.1 Alcances',2)
    h(d,'6.2 Limitaciones',2)
    h(d,'7 Marco teórico')
    h(d,'8 Métodos o procedimientos')
    p=h(d,'9 Cronograma de actividades');p.paragraph_format.keep_with_next=False
    reference_section(d,task)
    if task==9:
        consistency(d)
        h(d,'Anexo 2 Círculo de Covey',page=True)
        p=d.add_paragraph('Figura 1');style_para(p,False);p.runs[0].bold=True
        p=d.add_paragraph('Problema general en el esquema de articulación del anteproyecto');style_para(p,False);p.runs[0].italic=True
        p=d.add_paragraph();style_para(p,False,True);p.add_run().add_picture(str(covey_image()),width=Inches(6.5))
        for obj in p._p.xpath('.//wp:docPr'):obj.set('descr','Diagrama de óvalos concéntricos adaptado del formato docente. El centro contiene el problema general de vTaxi; los anillos exteriores conservan etiquetas para desarrollo posterior.')
        p=d.add_paragraph();style_para(p,False)
        text_runs(p,'*Nota.* Adaptado del esquema docente de articulación de ideas (Instituto Tecnológico Superior de Cajeme, s. f., p. 6). El problema general se ubica en el núcleo central; los otros componentes no se desarrollan en esta entrega.')
    name=f'Tarea{task}_DeLaCruzMunoz.docx';d.save(OUT/name)
    return name

def registration():
    ns={'w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
    with ZipFile(OFFICIAL) as z:
        tree=etree.fromstring(z.read('word/document.xml'))
        tables=tree.xpath('/w:document/w:body/w:tbl',namespaces=ns)
        def replace_cell(ti,ci,value):
            cell=tables[ti].xpath('./w:tr/w:tc',namespaces=ns)[ci]
            texts=cell.xpath('.//w:t',namespaces=ns)
            if texts:
                texts[0].text=value
                for text in texts[1:]:text.text=''
            # Remove hyperlinks from edited fields while retaining the runs.
            for link in list(cell.xpath('.//w:hyperlink',namespaces=ns)):
                parent=link.getparent();idx=parent.index(link)
                for r in list(link):parent.insert(idx,r);idx+=1
                parent.remove(link)
            for r in cell.xpath('.//w:r',namespaces=ns):
                pr=r.find(qn('w:rPr'))
                if pr is None:pr=etree.SubElement(r,qn('w:rPr'))
                for e in list(pr):
                    if e.tag in [qn('w:color'),qn('w:highlight'),qn('w:u')]:pr.remove(e)
                color=etree.SubElement(pr,qn('w:color'));color.set(qn('w:val'),'000000')
        replace_cell(0,1,'27 de septiembre de 2026')
        replace_cell(1,1,'DE LA CRUZ MUÑOZ MARTÍN JONATHAN')
        replace_cell(2,1,'26130503')
        cvu_file=WORK/'cvu.txt'
        replace_cell(2,3,cvu_file.read_text().strip() if cvu_file.exists() else '')
        replace_cell(4,1,TEACHER)
        replace_cell(5,0,DATA['title'])
        replace_cell(6,0,LGAC.upper())
        # Preserve the signature box labels and blank signing area.
        for text in tables[7].xpath('.//w:t',namespaces=ns):
            if text.text and ('Debes poner' in text.text or 'digital' in text.text):text.text=''
        with ZipFile(OUT/'RegistroTema_DeLaCruzMunoz.docx','w',ZIP_DEFLATED) as dest:
            for info in z.infolist():
                dest.writestr(info,etree.tostring(tree,xml_declaration=True,encoding='UTF-8',standalone=True) if info.filename=='word/document.xml' else z.read(info.filename))

def archive():
    REFS.mkdir(parents=True,exist_ok=True)
    src=Path('C:/Users/Sysx/Documents/vTaxi')
    records=[]
    for rel in ['README.md','docs/ARCHITECTURE.md','docs/REGULATORY-NL.md','docs/IMA-CHECKLIST-TAXIS.md','docs/FICHA-TECNICA-TAXIS.md']:
        target=REFS/('vtaxi-'+Path(rel).name)
        shutil.copyfile(src/rel,target)
        records.append({'source':str(src/rel),'copy':target.name,'sha256':hashlib.sha256(target.read_bytes()).hexdigest()})
    (REFS/'documentos-internos.json').write_text(json.dumps(records,ensure_ascii=False,indent=2),encoding='utf8')

if __name__=='__main__':
    OUT.mkdir(parents=True,exist_ok=True)
    archive()
    print(build(8));print(build(9));registration();print('RegistroTema_DeLaCruzMunoz.docx')
