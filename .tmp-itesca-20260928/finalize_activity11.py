from pathlib import Path
import json,hashlib,shutil
from docx import Document
import fitz

ROOT=Path(__file__).resolve().parents[1]
COURSE=ROOT/'ITESCA/maestria-en-gestion-administrativa/plan-de-negocios-mga'
TMP=ROOT/'.tmp-itesca-20260928'
EV=COURSE/'referencias-plan-de-negocios/actividad-11-2026-10-01'
ASSET=COURSE/'assets-plan-de-negocios/actividad-11-2026-10-01'
EV.mkdir(parents=True,exist_ok=True)
receipt=json.loads((TMP/'6550-receipt.json').read_text(encoding='utf-8'))
assert receipt['final_submission_observed'] and all(f['download_matches'] for f in receipt['files'])
for f in receipt['files']:
 assert hashlib.sha256((COURSE/f['filename']).read_bytes()).hexdigest()==f['sha256']
for name in ['6550-receipt.json','6550-after-save.txt','6550-after-save.png']:
 shutil.copy2(TMP/name,EV/name)

doc=Document(COURSE/'reporte-plan-de-negocios-Actividad-11-AM-Taller.docx')
assert len(doc.tables)==1 and len(doc.tables[0].rows)==17 and len(doc.tables[0].columns)==4
assert all(doc.tables[0].rows[i].cells[0].text.startswith(f'{i}. ') for i in range(1,17))
files={}
for name,count in [('reporte-plan-de-negocios-Actividad-11-AM-Taller.pdf',9),('Actividad-11-AM-Taller-Version-Grafica.pdf',3),('Actividad-11-AM-Taller-Version-Redactada.pdf',7)]:
 with fitz.open(COURSE/name) as pdf:
  assert len(pdf)==count
 files[name]={'pages':count,'sha256':hashlib.sha256((COURSE/name).read_bytes()).hexdigest()}
name='reporte-plan-de-negocios-Actividad-11-AM-Taller.docx'
files[name]={'sha256':hashlib.sha256((COURSE/name).read_bytes()).hexdigest()}
qa={'date':'2026-10-01','module':6550,'case':'AM Taller Autocentro: cambio de aceite y filtro','table':{'steps':16,'columns':4,'id_sequence_verified':True},'visual_review':'9 páginas revisadas; diagramas legibles, encabezado de tabla repetido y filas sin dividir. Se revisaron nuevamente las páginas 2, 7, 8 y 9 después del último ajuste.','rendering':'Exportación con Word COM y rasterización con Poppler. render_docx.py no pudo utilizarse por ausencia de soffice.exe.','measurements':'Los tiempos son estimaciones académicas, no mediciones. No se presentan respuestas de sondeo inventadas.','files':files,'moodle_verification':'Enviado para calificar; los dos PDF descargados coinciden por SHA-256 con los originales locales.'}
(EV/'verificacion.json').write_text(json.dumps(qa,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')

# Preserve reproducible authoring sources with paths relative to this activity.
di=(TMP/'build_activity11_diagram.py').read_text(encoding='utf-8')
di=di.replace("ROOT=Path(__file__).resolve().parents[1]", "ROOT=Path(__file__).resolve().parents[5]")
di=di.replace("ROOT/'.tmp-itesca-20260928/actividad11-proceso.json'", "ASSET/'proceso.json'")
(ASSET/'generar_diagrama.py').write_text(di,encoding='utf-8')
dc=(TMP/'build_activity11_docx.py').read_text(encoding='utf-8').replace("ROOT=Path(__file__).resolve().parents[1]", "ROOT=Path(__file__).resolve().parents[5]")
(ASSET/'generar_documento.py').write_text(dc,encoding='utf-8')
shutil.copy2(TMP/'export_word.ps1',ASSET/'exportar_word.ps1')
(ASSET/'README.md').write_text('''# Fuentes de la Actividad 11 de AM Taller

Fecha: 1 de octubre de 2026. El documento combina diagrama y tabla de un proceso propuesto de cambio de aceite y filtro. `proceso.json` contiene los 16 pasos y sus transiciones.

## Regeneración

1. Ejecutar `generar_diagrama.py` con Python, ReportLab y las fuentes Arial de Windows. Produce `flujo-vectorial.pdf` (dos páginas).
2. Rasterizar ese PDF con `pdftoppm -png -r 240 flujo-vectorial.pdf flujo` para obtener `flujo-1.png` y `flujo-2.png`.
3. Ejecutar `generar_documento.py` con python-docx. Conserva la portada de la actividad 10 y genera el nuevo Word de la actividad 11 sin modificar el original.
4. Ejecutar `exportar_word.ps1 -Paths <ruta-al-Word>` mediante PowerShell con Microsoft Word instalado. Actualiza campos, guarda el Word y exporta el PDF completo.
5. Tras revisar el PDF completo, extraer las páginas 1, 3 y 4 para la versión gráfica; las páginas 1, 2, 5, 6, 7, 8 y 9 para la redactada. Esas selecciones corresponden a la revisión de nueve páginas del 01/10/2026; deben revisarse si cambia la paginación.

Los PDF finales de ambas versiones fueron enviados a Moodle. No regenerarlos después del envío sin registrar una revisión nueva, porque cambiaría la correspondencia con los comprobantes de entrega. [Evidencia](../../referencias-plan-de-negocios/actividad-11-2026-10-01/README.md).
''',encoding='utf-8')

(EV/'README.md').write_text('''# Actividad 11: proceso del servicio de AM Taller

El 1 de octubre de 2026 se completó y entregó la actividad **Diagrama de flujo del proceso de producción/prestación del servicio**, módulo [6550](https://cursos3.e-itesca.edu.mx/mod/assign/view.php?id=6550), Plan de Negocios GGPN01. Vencimiento indicado en el aula: 5 de octubre de 2026, 23:59. La numeración 11 es local.

## Entrega verificada

Moodle muestra **Enviado para calificar**, intento 1, sin calificar. Última modificación mostrada por el aula: 01/10/2026, 00:04; se conserva la hora publicada sin atribuirle una zona horaria. Registro técnico UTC: 2026-10-01T07:04:59Z.

- [Versión gráfica, 3 páginas](../../Actividad-11-AM-Taller-Version-Grafica.pdf).
- [Versión redactada, 7 páginas](../../Actividad-11-AM-Taller-Version-Redactada.pdf).
- [Word editable completo](../../reporte-plan-de-negocios-Actividad-11-AM-Taller.docx) y [PDF integrado, 9 páginas](../../reporte-plan-de-negocios-Actividad-11-AM-Taller.pdf), conservados localmente.
- [Comprobante y huellas SHA-256](6550-receipt.json), [estado textual](6550-after-save.txt), [captura del aula](6550-after-save.png) y [verificación documental](verificacion.json).

Se descargaron los dos PDF desde sus enlaces de Moodle y sus huellas SHA-256 coinciden con los originales locales. No se acredita calificación ni aprobación docente.

## Correspondencia con la consigna y la retroalimentación

La [consigna consultada el 30/09](../actividad-11-2026-09-30/6550-consigna-estado.txt) exige ambas versiones y una tabla de cuatro columnas. El caso elegido desarrolla 16 pasos para cambio de aceite y filtro en vehículo ligero: solicitud, autorización, ejecución, control de calidad, entrega y seguimiento. Ambas versiones conservan los mismos identificadores; incluyen decisiones y conectores entre hojas. La tabla especifica actividad detallada, tiempo, materiales/equipo/herramientas y personal.

Los tiempos son estimaciones de dedicación activa. El diseño no afirma demanda comprobada, contrataciones, insumos disponibles ni mediciones reales. Incorpora registros propuestos de servicio, frecuencia, gasto y satisfacción, pero no sustituye el sondeo por conveniencia solicitado por la docente en la actividad 10. Las respuestas reales de ese sondeo siguen pendientes.

Los originales y versiones anteriores se conservan. [Fuentes editables y regeneración](../../assets-plan-de-negocios/actividad-11-2026-10-01/README.md).

## Incidencia de acceso resuelta

La resolución DNS ordinaria del aula falló. Se verificó la dirección publicada por DNS sobre HTTPS y el servidor respondió con HTTPS válido. La dirección se aplicó únicamente al navegador de esta ejecución, sin alterar la configuración de red del equipo. La entrega quedó verificada con autenticación y sin guardar credenciales ni cookies en esta carpeta.
''',encoding='utf-8')

section='''## Actividad 11 entregada el 01/10/2026

**AM Taller: diagrama de flujo del servicio.** Se enviaron al módulo 6550 la [versión gráfica](Actividad-11-AM-Taller-Version-Grafica.pdf) y la [versión redactada](Actividad-11-AM-Taller-Version-Redactada.pdf). Moodle muestra **Enviado para calificar** y los dos archivos descargados coinciden con los originales locales. [Word editable](reporte-plan-de-negocios-Actividad-11-AM-Taller.docx), [PDF integrado](reporte-plan-de-negocios-Actividad-11-AM-Taller.pdf) y [comprobante con verificación](referencias-plan-de-negocios/actividad-11-2026-10-01/README.md).

El proceso contiene 16 pasos concordantes y la tabla de cuatro columnas solicitada. Los tiempos son estimaciones; los registros propuestos no sustituyen el sondeo real pendiente de la actividad 10. Se conservan las versiones anteriores y los cortes históricos que siguen.

'''
for name in ['README.md','ESTADO-ENTREGABLES.md']:
 path=COURSE/name;txt=path.read_text(encoding='utf-8')
 if '## Actividad 11 entregada el 01/10/2026' not in txt:
  title,rest=txt.split('\n',1);path.write_text(title+'\n\n'+section+rest.lstrip('\n'),encoding='utf-8')
print(json.dumps({'status':'Enviado para calificar','evidence':str(EV),'files':files},ensure_ascii=False,indent=2))
