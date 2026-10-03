from pathlib import Path
from copy import deepcopy
from zipfile import ZipFile
import json,hashlib,sys
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT,WD_CELL_VERTICAL_ALIGNMENT
sys.stdout.reconfigure(encoding='utf-8',line_buffering=True)
ROOT=Path(__file__).resolve().parents[1]
WORK=ROOT/'.tmp-itesca-20260928'
SEM=ROOT/'ITESCA/maestria-en-gestion-administrativa/seminario-i-mga'
SRC=SEM/'entregas/Tarea9_DeLaCruzMunoz.docx'
DATA=json.loads((SEM/'anteproyecto/vtaxi-2026-09-27/contenido.json').read_text(encoding='utf-8'))
OG='Analizar la relación entre el nivel de gestión del cumplimiento normativo y el grado de preparación administrativa de vTaxi para su formalización como plataforma de taxis concesionados en Nuevo León durante septiembre a noviembre de 2026, mediante el contraste documental de requisitos, responsabilidades y evidencias del proyecto.'
OES=[
'Identificar y clasificar los requisitos y las competencias institucionales aplicables al modelo de taxis concesionados de vTaxi, distinguiéndolos de los correspondientes al transporte privado mediante empresas de redes de transporte.',
'Diagnosticar las brechas entre los requisitos aplicables y las evidencias jurídicas, técnicas y operativas disponibles en los borradores del proyecto, según su integridad, coherencia y respaldo documental.',
'Examinar la correspondencia entre la asignación de responsables, la verificación de requisitos y el seguimiento documental, por una parte, y la integridad del expediente y la resolución de observaciones internas, por otra, mediante una revisión trazable de los componentes del caso.',
'Proponer una ruta priorizada de gestión para fortalecer la preparación administrativa del expediente de vTaxi, con acciones, responsables propuestos y evidencias de atención, diferenciando la capacitación de conductores de la autorización de la plataforma.'
]
PRODUCTS=['Catálogo de requisitos clasificados por modalidad, autoridad, fuente y evidencia exigible.','Diagnóstico de brechas por requisito, con documento disponible, inconsistencia y evidencia faltante.','Matriz de contraste entre controles de gestión y estado documental del expediente, con trazabilidad de cada hallazgo.','Ruta de atención con prioridad, acción, responsable propuesto y criterio documental de cierre.']
newdir=SEM/'anteproyecto/vtaxi-2026-09-28';newdir.mkdir(parents=True,exist_ok=True)
artifact=WORK/'tarea10/artifact.md'
artifact.write_text(f'''# Contrato de continuación de Tarea 9\nReferencia: {SRC}\nSHA256: {hashlib.sha256(SRC.read_bytes()).hexdigest()}\nReferencia renderizada: .tmp-itesca-20260928/tarea10/Tarea9-referencia.pdf, 14 páginas.\nUna sección carta 8.5x11, márgenes1pulgada, Arial11, cuerpo dobleespacio y sangría0.5pulgada. Portada párrafos0-14 centrados, franja institucional en0; número superior derecho.\nSlots: párrafo3 actualizaractividad;14fecha. Mantener texto de apartados1-2 (p36-74) y referencias previas. Sustituir p76 por entrada metodológica; insertar objetivo tras3.1 y cuatroobjetivos tras3.2. Apartados4-9 mantienenencabezadosfuturos.\nAnexos: reemplazar matriz2columnaspor matrizgeneral3columnas y específica2columnas; requisitos explícitosT10. Sustituir dibujoCoveyprevio por cuadrantes y problema/objetivo completos enII. Fuente3.3 y consignareferenciadas. Tipografía tablasArial10,sencillo,bordesgrises. Se permitepaginaciónnueva acordecontenido.\nPreservar imágenesinstitucionales, encabezados/pies, dimensiones y texto1-2. WordactualizaíndiceyexportaPDF. Ningúnarchivo fuente se sobrescribe.\n''',encoding='utf-8')
d=Document(SRC)
original12=[p.text for p in d.paragraphs[36:75]]
# Keep the source cover styles and individual run formatting.
def replace_text(p,text):
 rp=deepcopy(p.runs[0]._r.rPr) if p.runs and p.runs[0]._r.rPr is not None else None
 p.clear();r=p.add_run(text)
 if rp is not None:r._r.insert(0,rp)
 return p
replace_text(d.paragraphs[3],'Seminario I\nTarea 10 Formulación de objetivos')
replace_text(d.paragraphs[14],'CIUDAD OBREGÓN, SONORA\n28 DE SEPTIEMBRE DE 2026')
replace_text(d.paragraphs[76],'Los objetivos traducen el problema de investigación en metas verificables y mantienen la correspondencia entre título, pregunta general y etapas del estudio, conforme al material de formulación de objetivos del Instituto Tecnológico Superior de Cajeme (s. f.-a).')
def insert_after(anchor,text,indent=True):
 p=d.add_paragraph(text,style='Normal');f=p.paragraph_format
 f.line_spacing=2;f.space_before=Pt(0);f.space_after=Pt(0);f.first_line_indent=Inches(.5) if indent else Inches(0);f.widow_control=True
 for r in p.runs:r.font.name='Arial';r.font.size=Pt(11);r.font.color.rgb=RGBColor(0,0,0)
 anchor._p.addnext(p._p);return p
h31=next(p for p in d.paragraphs if p.text=='3.1 Objetivo general')
insert_after(h31,OG)
h32=next(p for p in d.paragraphs if p.text=='3.2 Objetivos específicos')
a=h32
for i,oe in enumerate(OES):a=insert_after(a,f'{i+1}. {oe}',False)
a=insert_after(a,'La verificación de estas metas se realizará mediante un catálogo de requisitos, un diagnóstico de brechas, una matriz de contraste y una ruta priorizada. Se trata de productos del estudio de caso; no se pretende inferir causalidad poblacional ni garantizar una autorización administrativa.')
a=insert_after(a,'Los apartados 4 a 9 se desarrollarán en las tareas posteriores; se conservan sus encabezados para continuar el anteproyecto.')
# APA author/year suffixes for the distinct institutional resources.
oldref=next(p for p in d.paragraphs if p.text.startswith('Instituto Tecnológico Superior de Cajeme. (s. f.). Introducción'))
replace_text(oldref,oldref.text.replace('(s. f.).','(s. f.-b.).'))
newref=insert_after(oldref,'Instituto Tecnológico Superior de Cajeme. (s. f.-c). Tarea 10 Formulación de objetivos general y específicos [Consigna de Seminario I]. https://cursos3.e-itesca.edu.mx/mod/assign/view.php?id=2910',False)
for p in [newref]:p.paragraph_format.left_indent=Inches(.5);p.paragraph_format.first_line_indent=Inches(-.5)
newref2=d.add_paragraph('Instituto Tecnológico Superior de Cajeme. (s. f.-a). 3.3 Formulación de objetivos general y específicos [Material didáctico de la Unidad 3 de Seminario I]. https://cursos3.e-itesca.edu.mx/mod/resource/view.php?id=2897',style='Normal')
newref2.paragraph_format.line_spacing=2;newref2.paragraph_format.left_indent=Inches(.5);newref2.paragraph_format.first_line_indent=Inches(-.5)
oldref._p.addprevious(newref2._p)
# Drop previous annexes, which T10 explicitly asks to update.
annex=next(p for p in d.paragraphs if p.text=='Anexo 1 Matriz de consistencia')
node=annex._p
while node is not None:
 nxt=node.getnext()
 if node.tag!=qn('w:sectPr'):node.getparent().remove(node)
 node=nxt

def heading(text,level=1,page=False):
 p=d.add_paragraph(text,style=d.styles[f'Heading {level}'])
 p.paragraph_format.page_break_before=page;p.paragraph_format.keep_with_next=True;return p

def paragraph(text,spacing=2):
 p=d.add_paragraph(text,style='Normal');p.paragraph_format.line_spacing=spacing;p.paragraph_format.first_line_indent=Inches(0);p.paragraph_format.space_after=Pt(6);return p

def table(headers,rows,widths):
 t=d.add_table(rows=1,cols=len(headers));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
 for col,w in zip(t.columns,widths):col.width=Inches(w)
 for i,h in enumerate(headers):t.rows[0].cells[i].text=h
 for row in rows:
  cells=t.add_row().cells
  for c,v in zip(cells,row):c.text=v
 borders=OxmlElement('w:tblBorders')
 for side in ['top','left','bottom','right','insideH','insideV']:
  el=OxmlElement('w:'+side);el.set(qn('w:val'),'single');el.set(qn('w:sz'),'4');el.set(qn('w:color'),'D9D9D9');borders.append(el)
 t._tbl.tblPr.append(borders)
 for ri,row in enumerate(t.rows):
  pr=row._tr.get_or_add_trPr();pr.append(OxmlElement('w:cantSplit'))
  if ri==0:pr.append(OxmlElement('w:tblHeader'))
  for ci,c in enumerate(row.cells):
   c.width=Inches(widths[ci]);c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
   mar=OxmlElement('w:tcMar')
   for side in ['top','left','bottom','right']:
    el=OxmlElement('w:'+side);el.set(qn('w:w'),'100');el.set(qn('w:type'),'dxa');mar.append(el)
   c._tc.get_or_add_tcPr().append(mar)
   if ri==0:
    sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'E7EDF2');c._tc.get_or_add_tcPr().append(sh)
   for p in c.paragraphs:
    p.paragraph_format.line_spacing=1.05;p.paragraph_format.space_after=Pt(3);p.paragraph_format.first_line_indent=Inches(0);p.paragraph_format.keep_with_next=False
    for r in p.runs:r.font.name='Arial';r.font.size=Pt(10);r.font.color.rgb=RGBColor(0,0,0);r.bold=ri==0
 return t
heading('Anexo 1 Matriz de consistencia',page=True)
paragraph('Correspondencia general y específica del anteproyecto',1)
table(['Tema y problema general','Pregunta general','Objetivo general'],[[DATA['title']+'\n\n'+DATA['problem'],DATA['question'],OG]],[2.2,2.15,2.15])
paragraph('La relación se examina dentro del caso vTaxi mediante evidencia documental; los productos no presuponen causalidad estadística ni una resolución favorable de la autoridad.',1)
heading('Correspondencia de preguntas y objetivos específicos',level=2,page=True)
table(['Preguntas específicas','Objetivos específicos'],[[f'{i+1}. {q}',f'{i+1}. {o}'] for i,(q,o) in enumerate(zip(DATA['specific_questions'],OES))],[3.25,3.25])
paragraph('Nota. Elaboración propia. Cada objetivo responde a una pregunta conservada de la Tarea 9.',1)
heading('Anexo 2 Círculo de Covey',page=True)
paragraph('Figura 1\nProblema general y objetivo general en el cuadrante Importante y No urgente',1)
# Table is the actual requested quadrant scheme and remains editable in Word.
cov=table(['Urgente','No urgente'],[['I Importante\n\n','II Importante\n\nProblema general\n'+DATA['problem']+'\n\nObjetivo general\n'+OG],['III No importante\n\n','IV No importante\n\n']],[2.15,4.35])
paragraph('Nota. Elaboración propia conforme a la ubicación indicada en la Tarea 10 (Instituto Tecnológico Superior de Cajeme, s. f.-c). La definición del problema y de los objetivos se sitúa en la planeación importante y no urgente del estudio.',1)
# Check exact preservation of the developed sections before saving.
allp=[p.text for p in d.paragraphs]
i=allp.index('1 Antecedentes');j=allp.index('3 Formulación de objetivos general y específicos')
assert allp[i:j]==original12,'Cambió el texto de los apartados 1 o 2'
settings=d.settings.element
if not settings.xpath('./w:updateFields'):
 e=OxmlElement('w:updateFields');e.set(qn('w:val'),'true');settings.append(e)
d.core_properties.title='Formulación de objetivos del anteproyecto vTaxi';d.core_properties.author='Martín Jonathan de la Cruz Muñoz'
out=SEM/'entregas/Tarea10_DeLaCruzMunoz.docx';d.save(out)
# Restore untouched package parts byte for byte after python-docx normalizes XML declarations.
with ZipFile(SRC) as a,ZipFile(out) as b:
 preserve=[n for n in a.namelist() if n.startswith(('word/media/','word/header','word/footer'))]
 package={n:(a.read(n) if n in preserve else b.read(n)) for n in b.namelist()}
with ZipFile(out,'w',compression=8) as z:
 for n,value in package.items():z.writestr(n,value)
with ZipFile(SRC) as a,ZipFile(out) as b:
 preserve=[n for n in a.namelist() if n.startswith(('word/media/','word/header','word/footer'))]
 assert all(a.read(n)==b.read(n) for n in preserve),'Alteración de medios institucionales o encabezados'
(newdir/'objetivos.json').write_text(json.dumps({'title':DATA['title'],'problem':DATA['problem'],'general_question':DATA['question'],'general_objective':OG,'specific_questions':DATA['specific_questions'],'specific_objectives':OES,'verification_products':PRODUCTS},ensure_ascii=False,indent=2),encoding='utf-8')
(newdir/'objetivos.md').write_text('# Formulación de objetivos de vTaxi\n\n## Objetivo general\n\n'+OG+'\n\n## Objetivos específicos\n\n'+'\n\n'.join(f'{i+1}. {o}\n\nProducto de verificación: {v}' for i,(o,v) in enumerate(zip(OES,PRODUCTS)))+'\n',encoding='utf-8')
(newdir/'3.3-Formulacion-de-objetivos.pdf').write_bytes((WORK/'3.3-Formulacion-de-objetivos.pdf').read_bytes())
print(out)
