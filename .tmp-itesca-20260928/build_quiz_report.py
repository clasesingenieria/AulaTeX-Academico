from pathlib import Path
from copy import deepcopy
import json,sys
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
sys.stdout.reconfigure(encoding='utf-8')
ROOT=Path(__file__).resolve().parents[1]
COURSE=ROOT/'ITESCA/maestria-en-gestion-administrativa/plan-de-negocios-mga'
WORK=ROOT/'.tmp-itesca-20260928'
src=COURSE/'reporte-plan-de-negocios-Actividad-10-Resultados-Mercados.docx'
d=Document(src)
def replace(p,t):
 rp=deepcopy(p.runs[0]._r.rPr) if p.runs else None
 p.clear();r=p.add_run(t)
 if rp is not None:r._r.insert(0,rp)
replace(d.paragraphs[4],'Resultados del examen de Mercadotecnia e Imagen')
d.paragraphs[4].style=d.styles['Title']
for r in d.paragraphs[4].runs:r.font.color.rgb=RGBColor(0,0,0);r.bold=True;r.font.size=Pt(19)
replace(d.paragraphs[5],'Capítulo 2 · Plan de Negocios')
replace(d.paragraphs[6],'Martín Jonathan de la Cruz Muñoz\nMatrícula 26130503\nDocente Celia Velázquez Reyna')
replace(d.paragraphs[7],'Monterrey, Nuevo León · 28 de septiembre de 2026')
for node in list(d._element.body)[9:]:
 if node.tag!=qn('w:sectPr'):d._element.body.remove(node)
for name in ['Normal','Title','Heading 1','Heading 2']:
 st=d.styles[name];st.font.name='Arial';st.font.color.rgb=RGBColor(0,0,0)
normal=d.styles['Normal'];normal.font.size=Pt(11);normal.paragraph_format.line_spacing=1.15;normal.paragraph_format.space_after=Pt(7)
d.styles['Heading 2'].paragraph_format.space_before=Pt(10)
d.styles['Heading 2'].paragraph_format.space_after=Pt(5)
def p(t):return d.add_paragraph(t)
def h(t):return d.add_paragraph(t,style='Heading 1')
h('Resultado registrado en el aula')
p('El examen del Capítulo 2 Mercadotecnia e Imagen quedó finalizado el 28 de septiembre de 2026. El aula registra 100 de 100 como calificación más alta y confirma los requisitos de tres intentos y recepción de una calificación. Este reporte reúne el resultado y una explicación de los conceptos evaluados para su aplicación al plan de negocios.')
t=d.add_table(rows=1,cols=4);t.autofit=False
for c,w in zip(t.columns,[.8,1.4,2.2,2.1]):c.width=Inches(w)
for c,x in zip(t.rows[0].cells,['Intento','Estado','Inicio y fin en el aula','Calificación']):c.text=x
for values in [('1','Finalizado','06:08 a 06:23','100 de 100'),('2','Finalizado','17:27 a 17:29','100 de 100'),('3','Finalizado','17:30 a 17:30','100 de 100')]:
 for c,x in zip(t.add_row().cells,values):c.text=x
for ri,row in enumerate(t.rows):
 row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
 for c,w in zip(row.cells,[.8,1.4,2.2,2.1]):
  c.width=Inches(w)
  if ri==0:
   sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'E7EDF2');c._tc.get_or_add_tcPr().append(sh)
  for par in c.paragraphs:
   par.paragraph_format.space_after=Pt(6);par.paragraph_format.space_before=Pt(6)
   for r in par.runs:r.font.size=Pt(10);r.bold=ri==0
b=OxmlElement('w:tblBorders')
for side in ['top','bottom','left','right','insideH','insideV']:
 el=OxmlElement('w:'+side);el.set(qn('w:val'),'single');el.set(qn('w:sz'),'4');el.set(qn('w:color'),'D9D9D9');b.append(el)
t._tbl.tblPr.append(b)
p('Nota. Los tres intentos corresponden al 28 de septiembre de 2026. Las horas se transcriben como aparecen en el aula; no se realiza conversión de zona horaria. El método de calificación es la calificación más alta.')
h('Conceptos y fundamentos de las respuestas')
p('Los enunciados siguientes sintetizan los doce conceptos abordados. La explicación distingue definiciones próximas para evitar confundir una meta de la empresa con su visión o una acción de distribución con la demanda del mercado.')
bank=json.loads((WORK/'quiz-answer-bank.json').read_text(encoding='utf-8'))
titles=['Producto','Descripción de la empresa','Nombre de la empresa','Fuentes de información','Mercadotecnia','Misión','Plan de promoción','Análisis de mercado','Objetivos y visión','Análisis FODA','Demanda y comercialización','Factores claves de éxito']
for i,(title,data) in enumerate(zip(titles,bank.values()),1):
 d.add_paragraph(f'{i} {title}',style='Heading 2')
 p('Respuesta seleccionada: '+data['answer']+'. '+data['rationale'])
hp=h('Aplicación al plan de negocios');hp.paragraph_format.page_break_before=True
p('En AM Taller, el producto comprende el servicio mecánico, la información del diagnóstico, la atención y la confianza en la reparación. En Industrial Revolucionaria, la oferta comprende mantenimiento industrial y la evidencia del trabajo realizado. En ambos casos, describir el servicio exige precisar el alcance, el destinatario y el beneficio que se propone, sin prometer resultados aún no comprobados.')
p('La mercadotecnia permite contrastar esa propuesta con las necesidades de los clientes. Para la actividad de investigación de mercados, las páginas de proveedores y los registros estadísticos ofrecen evidencia documental del entorno y de la oferta. Una encuesta aplicada a clientes potenciales genera información específica del proyecto: ambas fuentes deben identificarse y sus resultados no pueden sustituirse entre sí.')
p('La misión expresa el propósito de la empresa; la visión describe la situación futura deseada y los objetivos fijan resultados verificables. El FODA organiza condiciones internas y externas, mientras el plan de promoción coordina los mensajes y canales. Esta distinción ayuda a formular acciones congruentes con el mercado y con los recursos disponibles.')
h('Fuente del resultado')
p('Instituto Tecnológico Superior de Cajeme. (2026). Examen Capítulo 2 Mercadotecnia e Imagen [Cuestionario y resumen de intentos de Plan de Negocios]. Aula virtual ITESCA. https://cursos3.e-itesca.edu.mx/mod/quiz/view.php?id=6552')
p('Consulta del resumen personal de intentos realizada el 28 de septiembre de 2026. El comprobante del aula se conserva en el expediente de evidencias del curso. Este reporte complementa el cuestionario ya finalizado.')
d.core_properties.author='Martín Jonathan de la Cruz Muñoz';d.core_properties.title='Resultados del examen de Mercadotecnia e Imagen'
out=COURSE/'reporte-plan-de-negocios-Actividad-9-Resultado-Examen-2.docx';d.save(out);print(out)
