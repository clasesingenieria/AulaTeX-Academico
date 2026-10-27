"""Build the S1 Word evidence from the user's authentic screenshots."""
from datetime import datetime, timezone
import hashlib
from io import BytesIO
import json
from pathlib import Path
import shutil

from docx import Document
from docx.shared import Inches, Pt
from PIL import Image


ROOT = Path(__file__).resolve().parents[1]
SUBJECT = ROOT / "UAS/licenciatura-en-contaduria-uas/herramientas-digitales"
URL = "https://drive.google.com/drive/folders/11t7MGMNuMRCMJnfmcp_XAnmjFHO5qkWH?usp=sharing"
STEPS = [
    ("9A1AN2xxG3", "Acceso a Google desde el navegador."),
    ("lsIaquUguc", "Menu de aplicaciones de Google."),
    ("wsJoCholWD", "Seleccion de Google Drive."),
    ("SqVnC94oLo", "Pagina principal de Google Drive."),
    ("BnJRzCRFAh", "Menu Nuevo: opcion Carpeta nueva."),
    ("MjwSGCEult", "Nombre de carpeta con grupo, carrera y estudiante."),
    ("aEuHJMX5DV", "Vista de Mi unidad despues de crear la carpeta."),
    ("HsB6EU9P4d", "Seleccion de GPO103 LCP - Martin Jonathan de la Cruz."),
    ("5j2GG8JahP", "Menu contextual de la carpeta."),
    ("B2ZfRcALgw", "Apertura de la opcion Compartir."),
    ("X10ieLLASo", "Configuracion inicial: acceso general restringido."),
    ("8czKo5Qq8B", "Opciones de permiso para cualquier usuario con el vinculo."),
    ("UP20ZKOoBC", "Acceso por vinculo con rol Editor y confirmacion de copia del vinculo."),
]


def digest(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def main():
    source = Path.home() / "Documents/ShareX/Screenshots/2026-10"
    assets = SUBJECT / "assets-herramientas-digitales/semana-01/capturas-originales"
    notes = SUBJECT / "referencias-herramientas-digitales/notas-herramientas-digitales/evidencia-carpeta-S1"
    output = SUBJECT / "entregas-herramientas-digitales"
    for folder in (assets, notes, output):
        folder.mkdir(parents=True, exist_ok=True)
    document = Document()
    section = document.sections[0]
    section.page_width, section.page_height = Inches(8.5), Inches(11)
    section.top_margin = section.bottom_margin = Inches(.65)
    section.left_margin = section.right_margin = Inches(.7)
    document.styles["Normal"].font.name = "Arial"
    document.styles["Normal"].font.size = Pt(11)
    document.add_heading("Carpeta de trabajo en Google Drive", 0)
    document.add_paragraph("Universidad Autonoma de Sinaloa\nLicenciatura en Contaduria Publica\nHerramientas digitales - Semana 1\nMartin Jonathan de la Cruz\nGrupo 103 LCP")
    document.add_paragraph("GPO103 LCP - Martin Jonathan de la Cruz")
    document.add_paragraph("Enlace de la carpeta:\n" + URL)
    records = []
    for number, (suffix, caption) in enumerate(STEPS, 1):
        original = source / f"msedge_{suffix}.png"
        target = assets / original.name
        if target.exists():
            assert digest(target) == digest(original), "Original local diferente"
        else:
            shutil.copy2(original, target)
        document.add_page_break()
        document.add_heading(f"{number:02d}. {caption}", 2)
        with Image.open(target) as image:
            image = image.convert("RGB")
            image.thumbnail((1400, 1400))
            buffer = BytesIO()
            image.save(buffer, format="JPEG", quality=72, optimize=True)
            buffer.seek(0)
            width = min(7.1, 7.8 * image.width / image.height)
            document.add_picture(buffer, width=Inches(width))
        records.append({"order": number, "source": str(original), "local": target.relative_to(SUBJECT).as_posix(), "sha256": digest(target), "caption": caption})
    docx = output / "S1_MartinJonathanDeLaCruz.docx"
    document.save(docx)
    manifest = {"created_utc": datetime.now(timezone.utc).isoformat(), "folder_url": URL, "screenshots": records, "docx": docx.relative_to(SUBJECT).as_posix(), "docx_sha256": digest(docx), "submitted": False, "teacher_specific_access_verified": False, "observed_general_access": "Anyone with link, Editor; screenshot only, not live verification"}
    (notes / "manifest.json").write_text(json.dumps(manifest, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps({"docx": str(docx), "bytes": docx.stat().st_size, "screenshots": len(records)}, ensure_ascii=True))


if __name__ == "__main__":
    main()