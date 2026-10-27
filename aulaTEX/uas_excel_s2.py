"""Generate the workbook for the existing S2 practice example."""
from pathlib import Path
from decimal import Decimal
import json

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.workbook.properties import CalcProperties


def main():
    root = Path(__file__).resolve().parents[1]
    subject = root / "UAS/licenciatura-en-contaduria-uas/herramientas-digitales"
    target = subject / "actividades/Excel-Basico-S2-Ejemplo.xlsx"
    workbook = Workbook()
    instructions = workbook.active
    instructions.title = "Control"
    for row in [
        ["Herramientas digitales", "Semana 2: ejemplo de preparacion"],
        ["Actividad local", "8; tarea Moodle 50114"],
        ["Estado", "Ejemplo hipotetico; no ejercicio asignado por el docente"],
        ["Uso", "Editar cantidades, precios y descuento en Importes"],
        ["Recalculo", "Abrir en Excel para calcular las formulas"],
        ["Comprobacion", "Resultados esperados solo para los datos originales"],
        ["Pendiente", "Integrar ejercicios de clase y del docente en el mismo libro"],
        ["Envio", "No enviado; instrucciones completas pendientes"],
    ]:
        instructions.append(row)
    sheet = workbook.create_sheet("Importes")
    sheet.append(["Concepto", "Cantidad", "Precio unitario", "Importe", "Importe neto", "Descuento"])
    records = [("Cuadernos", 3, 40), ("Carpetas", 2, 25), ("Lapices", 5, 8)]
    for number, (concept, quantity, price) in enumerate(records, 2):
        sheet.append([concept, quantity, price, f"=B{number}*C{number}", f"=D{number}*(1-$F$2)"])
    sheet["F2"] = 0.10
    sheet["F2"].number_format = "0%"
    sheet.append(["Total", None, None, "=SUM(D2:D4)", "=SUM(E2:E4)"])
    for row in sheet.iter_rows(min_row=2, max_row=5, min_col=3, max_col=5):
        for cell in row:
            cell.number_format = '#,##0.00'
    for kind, cells, upper in [("decimal", "B2:C4", 1000000), ("decimal", "F2", 1)]:
        validation = DataValidation(type=kind, operator="between", formula1=0, formula2=upper)
        validation.errorTitle = "Valor fuera de rango"
        validation.error = "Introducir un numero no negativo dentro del rango permitido."
        validation.showErrorMessage = True
        validation.errorStyle = "stop"
        sheet.add_data_validation(validation)
        validation.add(cells)
    checks = workbook.create_sheet("Comprobacion")
    checks.append(["Operacion", "Esperado original", "Resultado Excel", "Coincide"])
    expected = []
    for number, (concept, quantity, price) in enumerate(records, 2):
        amount = Decimal(quantity) * Decimal(price)
        expected.append(amount)
        checks.append([concept, float(amount), f"=Importes!D{number}", f"=B{number}=C{number}"])
    total = sum(expected)
    checks.append(["Total bruto", float(total), "=Importes!D5", "=B5=C5"])
    checks.append(["Total neto", float(total * Decimal("0.90")), "=Importes!E5", "=B6=C6"])
    assert total == 210 and total * Decimal("0.90") == 189
    for page in workbook:
        page.freeze_panes = "A2"
        page.sheet_view.showGridLines = False
        for cell in page[1]:
            cell.fill = PatternFill("solid", fgColor="123B6D")
            cell.font = Font(name="Arial", bold=True, color="FFFFFF")
        for row in page.iter_rows(min_row=2):
            for cell in row:
                cell.font = Font(name="Arial", size=11)
                cell.alignment = Alignment(vertical="top", wrap_text=True)
        for column in ["A", "B", "C", "D", "E", "F"]:
            page.column_dimensions[column].width = 23
        page.sheet_properties.pageSetUpPr.fitToPage = True
        page.page_setup.orientation = "landscape"
        page.page_setup.paperSize = page.PAPERSIZE_A4
        page.page_setup.fitToWidth = 1
        page.page_setup.fitToHeight = 0
        page.print_title_rows = "1:1"
    instructions.column_dimensions["B"].width = 75
    for number in range(2, 9):
        instructions.row_dimensions[number].height = 34
    workbook.calculation = CalcProperties(calcId=191029, fullCalcOnLoad=True, forceFullCalc=True)
    workbook.save(target)
    workbook.close()
    saved = load_workbook(target)
    assert saved.sheetnames == ["Control", "Importes", "Comprobacion"]
    assert saved["Importes"]["D2"].value == "=B2*C2"
    assert saved["Importes"]["E4"].value == "=D4*(1-$F$2)"
    assert saved["Importes"]["D5"].value == "=SUM(D2:D4)"
    assert len(saved["Importes"].data_validations.dataValidation) == 2
    saved.close()
    print(json.dumps({"file":str(target), "bytes":target.stat().st_size, "expected_gross":210, "expected_net":189, "formulas_verified":True, "excel_recalculation_verified":False}, ensure_ascii=True))


if __name__ == "__main__":
    main()